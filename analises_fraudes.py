import sys
import json
import re
import tempfile
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pandas as pd

from config import (
    BASE_DIR,
    PYTHON_YOLO,
    PYTHON_OCR,
    DETECTOR_PY,
    OCR_PY,
)


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BATCH_SIZE = 64

OCR_CONF_MINIMA = 0.70

# Aceita OCR de 1 até 100 caracteres
MIN_CARACTERES_OCR = 1
MAX_CARACTERES_OCR = 100

MAX_IMAGENS_POR_PASSAGEM = 50


# ============================================================
# LIMPA TEXTO
# ============================================================

def limpar_texto(texto):

    if texto is None:
        return ""

    # Não aplica máscara de placa brasileira.
    # Apenas transforma em string, remove espaços
    # nas extremidades e converte para maiúsculo.
    return str(texto).strip().upper()


# ============================================================
# VALIDA TAMANHO DO OCR
# ============================================================

def ocr_valido(texto):

    texto = limpar_texto(texto)

    if not texto:
        return False

    quantidade = len(texto)

    return (
        MIN_CARACTERES_OCR
        <= quantidade
        <= MAX_CARACTERES_OCR
    )


# ============================================================
# CLASSIFICA LEGIBILIDADE
# ============================================================

def classificar_legibilidade(
    texto,
    confidence
):

    texto = limpar_texto(
        texto
    )

    try:

        confidence = float(
            confidence
            or 0
        )

    except Exception:

        confidence = 0.0

    if not texto:

        return "ILEGÍVEL"

    if (
        ocr_valido(texto)
        and
        confidence >= OCR_CONF_MINIMA
    ):

        return "LEGÍVEL"

    return "PARCIALMENTE LEGÍVEL"


# ============================================================
# JSON
# ============================================================

def enviar_json(
    processo,
    dados
):

    if processo.poll() is not None:

        raise RuntimeError(
            "Processo já encerrou. "
            f"Código: "
            f"{processo.returncode}"
        )

    processo.stdin.write(
        json.dumps(
            dados,
            ensure_ascii=False
        )
        + "\n"
    )

    processo.stdin.flush()


def ler_json(
    processo,
    nome
):

    linha = (
        processo.stdout.readline()
    )

    if not linha:

        raise RuntimeError(
            f"Processo {nome} "
            "encerrou antes de responder."
        )

    return json.loads(
        linha
    )


def encerrar_processo(
    processo,
    nome
):

    if processo is None:
        return

    if processo.poll() is not None:
        return

    try:

        enviar_json(
            processo,
            {
                "command":
                    "exit"
            }
        )

        processo.wait(
            timeout=10
        )

    except Exception:

        try:

            processo.terminate()

            processo.wait(
                timeout=5
            )

        except Exception:

            try:
                processo.kill()

            except Exception:
                pass


# ============================================================
# INICIA YOLO
# ============================================================

def iniciar_yolo():

    print(
        "[INIT] Iniciando YOLO..."
    )

    return subprocess.Popen(
        [
            str(PYTHON_YOLO),
            str(DETECTOR_PY)
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=None,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR)
    )


# ============================================================
# INICIA OCR
# ============================================================

def iniciar_ocr():

    print(
        "[INIT] Iniciando PaddleOCR..."
    )

    return subprocess.Popen(
        [
            str(PYTHON_OCR),
            str(OCR_PY)
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=None,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR)
    )


# ============================================================
# NORMALIZA NOME DE COLUNA
# ============================================================

def normalizar_coluna(
    nome
):

    if nome is None:
        return ""

    nome = str(
        nome
    ).strip().upper()

    substituicoes = {
        "Á": "A",
        "À": "A",
        "Ã": "A",
        "Â": "A",
        "Ä": "A",

        "É": "E",
        "Ê": "E",
        "Ë": "E",

        "Í": "I",
        "Î": "I",

        "Ó": "O",
        "Õ": "O",
        "Ô": "O",

        "Ú": "U",
        "Û": "U",

        "Ç": "C",
    }

    for antigo, novo in (
        substituicoes.items()
    ):

        nome = nome.replace(
            antigo,
            novo
        )

    nome = re.sub(
        r"[^A-Z0-9]+",
        "_",
        nome
    )

    return nome.strip("_")


# ============================================================
# EXTRAI BYTES DA IMAGEM
# ============================================================

def extrair_bytes_imagem(
    img
):

    try:
        return img._data()

    except Exception:
        return None


# ============================================================
# DESCOBRE EXTENSÃO
# ============================================================

def descobrir_extensao(
    img,
    dados
):

    formato = getattr(
        img,
        "format",
        None
    )

    if formato:

        formato = str(
            formato
        ).lower()

        if formato in (
            "jpg",
            "jpeg"
        ):
            return ".jpg"

        if formato == "png":
            return ".png"

        if formato == "webp":
            return ".webp"

        if formato == "bmp":
            return ".bmp"

    if dados[:3] == (
        b"\xff\xd8\xff"
    ):
        return ".jpg"

    if dados[:8] == (
        b"\x89PNG\r\n\x1a\n"
    ):
        return ".png"

    if dados[:4] == b"RIFF":
        return ".webp"

    return ".png"


# ============================================================
# CARREGA PLANILHA
# ============================================================

def carregar_planilha(
    excel_path,
    temp_dir
):

    print(
        f"\n[EXCEL] "
        f"{excel_path}"
    )

    wb = openpyxl.load_workbook(
        excel_path,
        data_only=True
    )

    if "Conferencia" in wb.sheetnames:

        sheet = wb[
            "Conferencia"
        ]

    else:

        sheet = wb.active

    # ========================================================
    # LOCALIZA COLUNAS
    # ========================================================

    colunas = {}

    for coluna in range(
        1,
        sheet.max_column + 1
    ):

        valor = sheet.cell(
            row=1,
            column=coluna
        ).value

        nome = normalizar_coluna(
            valor
        )

        if nome:

            colunas[
                nome
            ] = coluna

    coluna_id = colunas.get(
        "ID_PASSAGEM"
    )

    coluna_placa = (
        colunas.get(
            "PLACA_NO_SISTEMA"
        )
    )

    if coluna_id is None:

        raise ValueError(
            "A coluna ID_PASSAGEM "
            "não foi encontrada."
        )

    # ========================================================
    # IMAGENS POR LINHA
    # ========================================================

    imagens_por_linha = {}

    for img_index, img in enumerate(
        sheet._images
    ):

        try:

            linha = (
                img.anchor._from.row
                + 1
            )

        except Exception:

            continue

        imagens_por_linha.setdefault(
            linha,
            []
        ).append(
            (
                img_index,
                img
            )
        )

    print(
        f"[EXCEL] "
        f"Imagens encontradas: "
        f"{len(sheet._images)}"
    )

    # ========================================================
    # PASSAGENS
    # ========================================================

    passagens = []

    for linha in range(
        2,
        sheet.max_row + 1
    ):

        id_passagem = sheet.cell(
            row=linha,
            column=coluna_id
        ).value

        if (
            id_passagem is None
            or
            str(id_passagem).strip() == ""
        ):
            continue

        if coluna_placa is not None:

            placa_sistema = (
                sheet.cell(
                    row=linha,
                    column=coluna_placa
                ).value
            )

        else:

            placa_sistema = ""

        imagens = []

        for img_index, img in (
            imagens_por_linha.get(
                linha,
                []
            )
        ):

            dados = (
                extrair_bytes_imagem(
                    img
                )
            )

            if not dados:
                continue

            extensao = (
                descobrir_extensao(
                    img,
                    dados
                )
            )

            caminho = (
                temp_dir
                /
                (
                    f"linha_{linha}"
                    f"_img_{img_index}"
                    f"{extensao}"
                )
            )

            caminho.write_bytes(
                dados
            )

            imagens.append(
                {
                    "task_id":
                        (
                            f"linha{linha}"
                            f"_img{img_index}"
                        ),

                    "image_path":
                        str(caminho),

                    "indice":
                        img_index
                }
            )

        if (
            len(imagens)
            > MAX_IMAGENS_POR_PASSAGEM
        ):

            imagens = imagens[
                :MAX_IMAGENS_POR_PASSAGEM
            ]

        passagens.append(
            {
                "linha":
                    linha,

                "id_passagem":
                    id_passagem,

                "placa_sistema":
                    placa_sistema,

                "imagens":
                    imagens
            }
        )

    return (
        passagens,
        colunas
    )


# ============================================================
# YOLO — TODAS AS DETECÇÕES
# ============================================================

def executar_yolo_fraude(
    processo_yolo,
    imagens
):

    if not imagens:
        return {}

    items = []

    for imagem in imagens:

        items.append(
            {
                "task_id":
                    imagem["task_id"],

                "image_path":
                    imagem["image_path"]
            }
        )

    enviar_json(
        processo_yolo,
        {
            "command":
                "detect_batch_all",

            "images":
                items
        }
    )

    resposta = ler_json(
        processo_yolo,
        "YOLO"
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        []
    ):

        task_id = resultado.get(
            "task_id"
        )

        if task_id:

            resultados[
                task_id
            ] = resultado

    return resultados


# ============================================================
# OCR
# ============================================================

def executar_ocr(
    processo_ocr,
    crops
):

    if not crops:
        return {}

    enviar_json(
        processo_ocr,
        {
            "command":
                "ocr_batch",

            "images":
                crops
        }
    )

    resposta = ler_json(
        processo_ocr,
        "OCR"
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        []
    ):

        task_id = (
            resultado.get(
                "task_id"
            )
        )

        if task_id:

            resultados[
                task_id
            ] = resultado

    return resultados


# ============================================================
# AGRUPA LEITURAS OCR
# ============================================================

def agrupar_placas(
    leituras
):

    grupos = []

    for leitura in leituras:

        resultado = leitura[
            "resultado"
        ]

        # ====================================================
        # USA O OCR_BEST DIRETAMENTE
        # ====================================================

        texto_original = resultado.get(
            "ocr_best",
            ""
        )

        texto = limpar_texto(
            texto_original
        )

        try:

            confianca = float(
                resultado.get(
                    "confidence_best",
                    0
                )
                or 0
            )

        except Exception:

            confianca = 0.0

        # ====================================================
        # FILTROS
        # ====================================================

        if not texto:
            continue

        if not (
            MIN_CARACTERES_OCR
            <= len(texto)
            <= MAX_CARACTERES_OCR
        ):
            continue

        if (
            confianca
            < OCR_CONF_MINIMA
        ):
            continue

        # ====================================================
        # PROCURA DUPLICATA
        # ====================================================

        encontrou = False

        for grupo in grupos:

            placa_existente = grupo[
                "placa"
            ]

            if texto == placa_existente:

                encontrou = True

                # Mantém a leitura
                # com maior confiança.
                if (
                    confianca
                    > grupo[
                        "confidence"
                    ]
                ):

                    grupo.update(
                        {
                            "confidence":
                                confianca,

                            "task_id":
                                leitura[
                                    "task_id"
                                ],

                            "resultado":
                                resultado,

                            "yolo_confidence":
                                leitura.get(
                                    "yolo_confidence",
                                    0
                                )
                        }
                    )

                break

        if encontrou:
            continue

        # ====================================================
        # NOVA LEITURA
        # ====================================================

        grupos.append(
            {
                "placa":
                    texto,

                "confidence":
                    confianca,

                "task_id":
                    leitura[
                        "task_id"
                    ],

                "resultado":
                    resultado,

                "yolo_confidence":
                    leitura.get(
                        "yolo_confidence",
                        0
                    )
            }
        )

    return grupos


# ============================================================
# ANALISA PASSAGEM
# ============================================================

def analisar_passagem(
    passagem,
    yolo_results,
    ocr_results
):

    leituras = []

    # ========================================================
    # CADA IMAGEM
    # ========================================================

    for imagem in passagem[
        "imagens"
    ]:

        task_id = imagem[
            "task_id"
        ]

        yolo_result = (
            yolo_results.get(
                task_id,
                {}
            )
        )

        if not yolo_result.get(
            "detected",
            False
        ):
            continue

        detections = (
            yolo_result.get(
                "detections",
                []
            )
        )

        if not detections:

            # Compatibilidade com detector
            # antigo, caso necessário.
            crop_b64 = (
                yolo_result.get(
                    "crop_b64"
                )
            )

            if crop_b64:

                detections = [
                    {
                        "detection_index":
                            1,

                        "crop_b64":
                            crop_b64,

                        "confidence":
                            yolo_result.get(
                                "confidence",
                                0
                            )
                    }
                ]

        # ====================================================
        # TODAS AS DETECÇÕES DA IMAGEM
        # ====================================================

        for detection in detections:

            detection_index = (
                detection.get(
                    "detection_index",
                    1
                )
            )

            detection_task_id = (
                f"{task_id}"
                f"_d{detection_index}"
            )

            resultado_ocr = (
                ocr_results.get(
                    detection_task_id
                )
            )

            if not resultado_ocr:
                continue

            leituras.append(
                {
                    "task_id":
                        detection_task_id,

                    "resultado":
                        resultado_ocr,

                    "yolo_confidence":
                        detection.get(
                            "confidence",
                            0
                        ),

                    "detection_index":
                        detection_index,

                    "imagem_index":
                        imagem["indice"]
                }
            )

    # ========================================================
    # AGRUPA LEITURAS
    # ========================================================

    placas = agrupar_placas(
        leituras
    )

    # ========================================================
    # ESCOLHE A PRINCIPAL
    # ========================================================

    melhor = None

    if placas:

        melhor = max(
            placas,
            key=lambda p: (
                p["confidence"],
                p["yolo_confidence"]
            )
        )

    # ========================================================
    # MELHOR RESULTADO
    # ========================================================

    if melhor:

        resultado_melhor = (
            melhor["resultado"]
        )

        placa_principal = (
            melhor["placa"]
        )

        conf_principal = (
            melhor["confidence"]
        )

        yolo_principal = (
            melhor["yolo_confidence"]
        )

        imagem_principal = (
            melhor["task_id"]
        )

        legibilidade = (
            "LEGÍVEL"
        )

    else:

        resultado_melhor = {}

        placa_principal = ""

        conf_principal = 0

        yolo_principal = 0

        imagem_principal = ""

        legibilidade = (
            "SEM IMAGEM"
            if not passagem[
                "imagens"
            ]
            else "ILEGÍVEL"
        )

    # ========================================================
    # LISTAS DAS LEITURAS
    # ========================================================

    placas_detectadas = "; ".join(
        p["placa"]
        for p in placas
    )

    confiancas = "; ".join(
        f"{p['confidence']:.4f}"
        for p in placas
    )

    imagens_usadas = "; ".join(
        p["task_id"]
        for p in placas
    )

    # ========================================================
    # RESULTADO
    # ========================================================

    return {
        "Linha Excel":
            passagem[
                "linha"
            ],

        "ID_PASSAGEM":
            passagem[
                "id_passagem"
            ],

        "Placa no sistema":
            passagem[
                "placa_sistema"
            ],

        "Qtd Imagens":
            len(
                passagem[
                    "imagens"
                ]
            ),

        "Qtd YOLO":
            len(
                leituras
            ),

        "Qtd Placas Legíveis":
            len(
                placas
            ),

        "Placas Detectadas":
            placas_detectadas,

        "Confianças OCR":
            confiancas,

        "Imagens Usadas":
            imagens_usadas,

        "Placa Principal":
            placa_principal,

        "OCR Confidence":
            conf_principal,

        "Legibilidade":
            legibilidade,

        "Imagem Principal":
            imagem_principal,

        "YOLO Confidence":
            yolo_principal,

        "SR Usada":
            resultado_melhor.get(
                "sr_used",
                False
            ),

        "SR OCR":
            resultado_melhor.get(
                "sr_ocr",
                ""
            ),

        "SR Confidence":
            resultado_melhor.get(
                "sr_confidence",
                0
            ),

        "Resultado Graças SR":
            resultado_melhor.get(
                "resultado_gracas_sr",
                False
            ),

        "Best Variant":
            resultado_melhor.get(
                "best_variant",
                ""
            ),

        "OCR Cubic":
            resultado_melhor.get(
                "ocr_cubic",
                ""
            ),

        "Conf Cubic":
            resultado_melhor.get(
                "confidence_cubic",
                0
            ),

        "OCR Cubic Auxiliar":
            resultado_melhor.get(
                "ocr_cubic_auxiliar",
                ""
            ),

        "OCR CLAHE":
            resultado_melhor.get(
                "ocr_clahe",
                ""
            ),

        "Conf CLAHE":
            resultado_melhor.get(
                "confidence_clahe",
                0
            ),

        "OCR CLAHE Auxiliar":
            resultado_melhor.get(
                "ocr_clahe_auxiliar",
                ""
            ),

        "OCR Best Original":
            resultado_melhor.get(
                "ocr_best_original",
                ""
            ),

        "OCR Best Normal":
            resultado_melhor.get(
                "ocr_best_normal",
                ""
            ),

        "Conf Best Normal":
            resultado_melhor.get(
                "confidence_best_normal",
                0
            ),

        "Auxiliar Alterou":
            resultado_melhor.get(
                "auxiliar_alterou",
                False
            ),

        "OCR Best Final":
            resultado_melhor.get(
                "ocr_best",
                ""
            ),

        "Conf Best Final":
            resultado_melhor.get(
                "confidence_best",
                0
            )
    }


# ============================================================
# PROCESSA UM EXCEL
# ============================================================

def processar_excel(
    excel_path,
    processo_yolo,
    processo_ocr
):

    temp_dir = Path(
        tempfile.mkdtemp(
            prefix="fraude_"
        )
    )

    try:

        (
            passagens,
            colunas
        ) = carregar_planilha(
            excel_path,
            temp_dir
        )

        todas_imagens = []

        for passagem in passagens:

            todas_imagens.extend(
                passagem[
                    "imagens"
                ]
            )

        print(
            f"[PASSAGENS] "
            f"{len(passagens)}"
        )

        print(
            f"[IMAGENS] "
            f"{len(todas_imagens)}"
        )

        todos_yolo = {}

        todos_ocr = {}

        # ====================================================
        # BATCH
        # ====================================================

        total_lotes = (
            len(todas_imagens)
            + BATCH_SIZE
            - 1
        ) // BATCH_SIZE

        for lote_idx in range(
            total_lotes
        ):

            inicio = (
                lote_idx
                * BATCH_SIZE
            )

            fim = min(
                inicio
                + BATCH_SIZE,
                len(todas_imagens)
            )

            lote = (
                todas_imagens[
                    inicio:fim
                ]
            )

            print(
                "\n"
                + "-" * 70
            )

            print(
                f"LOTE "
                f"{lote_idx + 1}/"
                f"{total_lotes}"
            )

            print(
                f"Imagens: "
                f"{len(lote)}"
            )

            print(
                "-" * 70
            )

            # ================================================
            # YOLO - TODAS AS DETECÇÕES
            # ================================================

            yolo_results = (
                executar_yolo_fraude(
                    processo_yolo,
                    lote
                )
            )

            todos_yolo.update(
                yolo_results
            )

            # ================================================
            # PREPARA TODOS OS CROPS
            # ================================================

            crops = []

            for imagem in lote:

                task_id = imagem[
                    "task_id"
                ]

                yolo_result = (
                    yolo_results.get(
                        task_id,
                        {}
                    )
                )

                if not yolo_result.get(
                    "detected",
                    False
                ):
                    continue

                detections = (
                    yolo_result.get(
                        "detections",
                        []
                    )
                )

                if not detections:

                    crop_b64 = (
                        yolo_result.get(
                            "crop_b64"
                        )
                    )

                    if crop_b64:

                        detections = [
                            {
                                "detection_index":
                                    1,

                                "crop_b64":
                                    crop_b64
                            }
                        ]

                for detection in detections:

                    detection_index = (
                        detection.get(
                            "detection_index",
                            1
                        )
                    )

                    crop_b64 = (
                        detection.get(
                            "crop_b64"
                        )
                    )

                    if not crop_b64:
                        continue

                    crops.append(
                        {
                            "task_id":
                                (
                                    f"{task_id}"
                                    f"_d{detection_index}"
                                ),

                            "crop_b64":
                                crop_b64
                        }
                    )

            print(
                f"[YOLO] "
                f"Imagens analisadas: "
                f"{len(lote)}"
            )

            print(
                f"[YOLO] "
                f"Detecções enviadas ao OCR: "
                f"{len(crops)}"
            )

            # ================================================
            # OCR
            # ================================================

            ocr_results = executar_ocr(
                processo_ocr,
                crops
            )

            todos_ocr.update(
                ocr_results
            )

        # ====================================================
        # ANALISA PASSAGENS
        # ====================================================

        resultados_finais = []

        for passagem in passagens:

            resultado = (
                analisar_passagem(
                    passagem,
                    todos_yolo,
                    todos_ocr
                )
            )

            resultados_finais.append(
                resultado
            )

        return resultados_finais

    finally:

        shutil.rmtree(
            temp_dir,
            ignore_errors=True
        )


# ============================================================
# PROCESSA PASTA
# ============================================================

def processar_pasta(
    pasta
):

    pasta = Path(
        pasta
    ).resolve()

    arquivos = sorted(
        [
            arquivo

            for arquivo in pasta.rglob(
                "*.xlsx"
            )

            if not arquivo.name.startswith(
                "~$"
            )

            and
            not arquivo.stem.endswith(
                "_analisado"
            )
        ]
    )

    if not arquivos:

        print(
            "[ERRO] "
            "Nenhum XLSX encontrado."
        )

        return

    print(
        "\n"
        + "=" * 75
    )

    print(
        "ANÁLISE DE FRAUDES"
    )

    print(
        "=" * 75
    )

    print(
        f"Pasta: "
        f"{pasta}"
    )

    print(
        f"Arquivos: "
        f"{len(arquivos)}"
    )

    print(
        "=" * 75
    )

    for arquivo in [
        PYTHON_YOLO,
        PYTHON_OCR,
        DETECTOR_PY,
        OCR_PY
    ]:

        if not Path(
            arquivo
        ).exists():

            raise FileNotFoundError(
                f"Arquivo não encontrado:\n"
                f"{arquivo}"
            )

    processo_yolo = iniciar_yolo()

    processo_ocr = iniciar_ocr()

    try:

        import time

        print(
            "\n[INIT] "
            "Aguardando modelos..."
        )

        time.sleep(15)

        total_passagens = 0
        total_legiveis = 0
        total_parciais = 0
        total_ilegiveis = 0
        total_sem_imagem = 0

        for indice, excel_path in enumerate(
            arquivos,
            start=1
        ):

            print(
                "\n"
                + "#" * 75
            )

            print(
                f"ARQUIVO "
                f"{indice}/"
                f"{len(arquivos)}"
            )

            print(
                excel_path
            )

            print(
                "#" * 75
            )

            try:

                resultados = (
                    processar_excel(
                        excel_path,
                        processo_yolo,
                        processo_ocr
                    )
                )

                df = pd.DataFrame(
                    resultados
                )

                # =================================================
                # SAÍDA
                # =================================================

                saida = (
                    excel_path.parent
                    /
                    (
                        excel_path.stem
                        + "_analisado.xlsx"
                    )
                )

                df.to_excel(
                    saida,
                    index=False
                )

                # =================================================
                # ESTATÍSTICAS
                # =================================================

                total_passagens += len(
                    df
                )

                legiveis = int(
                    (
                        df[
                            "Legibilidade"
                        ]
                        == "LEGÍVEL"
                    ).sum()
                )

                parciais = int(
                    (
                        df[
                            "Legibilidade"
                        ]
                        == "PARCIALMENTE LEGÍVEL"
                    ).sum()
                )

                ilegiveis = int(
                    (
                        df[
                            "Legibilidade"
                        ]
                        == "ILEGÍVEL"
                    ).sum()
                )

                sem_imagem = int(
                    (
                        df[
                            "Legibilidade"
                        ]
                        == "SEM IMAGEM"
                    ).sum()
                )

                total_legiveis += (
                    legiveis
                )

                total_parciais += (
                    parciais
                )

                total_ilegiveis += (
                    ilegiveis
                )

                total_sem_imagem += (
                    sem_imagem
                )

                print(
                    "\n[OK] "
                    "Arquivo finalizado."
                )

                print(
                    f"Passagens: "
                    f"{len(df)}"
                )

                print(
                    f"Legíveis: "
                    f"{legiveis}"
                )

                print(
                    f"Parciais: "
                    f"{parciais}"
                )

                print(
                    f"Ilegíveis: "
                    f"{ilegiveis}"
                )

                print(
                    f"Sem imagem: "
                    f"{sem_imagem}"
                )

                print(
                    f"Saída: "
                    f"{saida}"
                )

            except KeyboardInterrupt:

                raise

            except Exception as e:

                print(
                    "\n[ERRO] "
                    f"{type(e).__name__}: "
                    f"{e}"
                )

    finally:

        encerrar_processo(
            processo_yolo,
            "YOLO"
        )

        encerrar_processo(
            processo_ocr,
            "OCR"
        )

    # ========================================================
    # RESUMO FINAL
    # ========================================================

    print(
        "\n"
        + "=" * 75
    )

    print(
        "RESUMO FINAL"
    )

    print(
        "=" * 75
    )

    print(
        f"Passagens: "
        f"{total_passagens}"
    )

    print(
        f"LEGÍVEIS: "
        f"{total_legiveis}"
    )

    print(
        f"PARCIALMENTE LEGÍVEIS: "
        f"{total_parciais}"
    )

    print(
        f"ILEGÍVEIS: "
        f"{total_ilegiveis}"
    )

    print(
        f"SEM IMAGEM: "
        f"{total_sem_imagem}"
    )

    print(
        "=" * 75
    )


# ============================================================
# MAIN
# ============================================================

def main():

    if len(sys.argv) < 2:

        print(
            'Uso: python analise_fraudes.py '
            '"C:\\caminho\\da\\pasta"'
        )

        sys.exit(1)

    pasta = Path(
        sys.argv[1]
    ).resolve()

    if not pasta.exists():

        print(
            f"[ERRO] "
            f"Pasta não encontrada:\n"
            f"{pasta}"
        )

        sys.exit(1)

    if not pasta.is_dir():

        print(
            f"[ERRO] "
            f"O caminho não é uma pasta:\n"
            f"{pasta}"
        )

        sys.exit(1)

    processar_pasta(
        pasta
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()