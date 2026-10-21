from __future__ import annotations


# ============================================================
# IMPORTS
# ============================================================

import base64
import json
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# IMPORT DO CONFIG
# ============================================================

BASE_DIR_LOCAL = (
    Path(__file__)
    .resolve()
    .parent
    .parent
)

if str(BASE_DIR_LOCAL) not in sys.path:
    sys.path.insert(
        0,
        str(BASE_DIR_LOCAL)
    )

from config import (
    BASE_DIR,
    PYTHON_OCR,
    OCR_SAU_PY,
)


# ============================================================
# CONFIGURAÇÕES
# ============================================================

RESULTADOS_DIR = (
    BASE_DIR
    / "resultados_ocr"
)

NOME_SAIDA = (
    "relatorio_administrativos.xlsx"
)


# ============================================================
# MÁSCARAS OCR
# ============================================================

MAPA_LETRAS = {
    "0": "O",
    "1": "I",
    "2": "Z",
    "4": "A",
    "5": "S",
    "6": "G",
    "7": "Z",
    "8": "B",
}


MAPA_DIGITOS = {
    "O": "0",
    "Q": "0",
    "D": "0",

    "I": "1",
    "L": "1",
    "T": "7",

    "Z": "2",

    "A": "4",

    "S": "5",

    "G": "6",

    "B": "8",
}


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def limpar_caminho_digitado(
    valor
):
    """
    Remove aspas de caminhos copiados/arrastados
    para o terminal do Windows.
    """

    return (
        str(valor)
        .strip()
        .strip('"')
        .strip("'")
    )


def normalizar_chave(
    valor
):
    """
    Remove:
        - acentos;
        - espaços;
        - hífens;
        - pontos;
        - demais caracteres especiais.

    Mantém somente A-Z e 0-9.
    """

    if valor is None:
        return ""

    try:

        if pd.isna(
            valor
        ):
            return ""

    except Exception:
        pass

    valor = (
        str(valor)
        .strip()
        .replace(
            "\ufeff",
            ""
        )
    )

    valor = (
        unicodedata.normalize(
            "NFKD",
            valor
        )
    )

    valor = "".join(
        c
        for c in valor
        if not unicodedata.combining(
            c
        )
    )

    return re.sub(
        r"[^A-Z0-9]",
        "",
        valor.upper()
    )


def normalizar_nome_coluna(
    valor
):
    return normalizar_chave(
        valor
    )


def formatar_placa(
    placa
):
    """
    ABC1234 -> ABC-1234
    ABC1D23 -> ABC-1D23
    """

    placa = normalizar_chave(
        placa
    )

    if len(
        placa
    ) == 7:

        return (
            f"{placa[:3]}-"
            f"{placa[3:]}"
        )

    return placa


# ============================================================
# CARREGA CADASTRO
# ============================================================

def carregar_cadastro_excel():

    print(
        "\n"
        + "=" * 70
    )

    print(
        "CADASTRO DE PREFIXOS E PLACAS"
    )

    print(
        "=" * 70
    )

    while True:

        caminho_digitado = input(
            "\nInforme o caminho do Excel contendo as colunas "
            "'Prefixo' e 'Placa':\n> "
        )

        caminho_excel = Path(
            limpar_caminho_digitado(
                caminho_digitado
            )
        )

        if not caminho_excel.is_file():

            print(
                "\nERRO: o arquivo informado não existe:\n"
                f"{caminho_excel}\n"
            )

            continue

        try:

            df = pd.read_excel(
                caminho_excel,
                dtype=str
            )

        except Exception as erro:

            print(
                "\nERRO ao abrir o Excel:\n"
                f"{erro}\n"
            )

            continue

        mapa_colunas = {
            normalizar_nome_coluna(
                coluna
            ): coluna

            for coluna in df.columns
        }

        coluna_prefixo = (
            mapa_colunas.get(
                "PREFIXO"
            )
        )

        coluna_placa = (
            mapa_colunas.get(
                "PLACA"
            )
        )

        if (
            not coluna_prefixo
            or
            not coluna_placa
        ):

            print(
                "\nERRO: a planilha precisa conter "
                "'Prefixo' e 'Placa'."
            )

            print(
                "Colunas encontradas:",
                ", ".join(
                    str(c)
                    for c in df.columns
                )
            )

            continue

        cadastro = {}

        ignorados = 0

        for _, linha in df.iterrows():

            prefixo = (
                normalizar_chave(
                    linha.get(
                        coluna_prefixo
                    )
                )
            )

            placa = (
                normalizar_chave(
                    linha.get(
                        coluna_placa
                    )
                )
            )

            if (
                not prefixo
                or
                not placa
            ):

                ignorados += 1
                continue

            if prefixo not in cadastro:

                cadastro[
                    prefixo
                ] = {
                    "prefixo":
                        prefixo,

                    "placas":
                        [],
                }

            if placa not in cadastro[
                prefixo
            ][
                "placas"
            ]:

                cadastro[
                    prefixo
                ][
                    "placas"
                ].append(
                    placa
                )

        if not cadastro:

            print(
                "\nERRO: nenhuma combinação válida "
                "de Prefixo + Placa foi encontrada."
            )

            continue

        print(
            "\nCadastro carregado com sucesso:"
        )

        print(
            f"  Prefixos: "
            f"{len(cadastro)}"
        )

        print(
            f"  Linhas ignoradas: "
            f"{ignorados}"
        )

        print(
            f"  Arquivo: "
            f"{caminho_excel}\n"
        )

        return cadastro


# ============================================================
# MÁSCARA DO PREFIXO
# ============================================================

def aplicar_mascara_prefixo(
    texto
):
    """
    Procura prefixos no formato:

        AA-0000
        AA-00000
        AAA-0000
        AAA-00000

    Hífen obrigatório.

    Antes do hífen:
        letras.

    Depois do hífen:
        números.

    Corrige confusões comuns do OCR.
    """

    if texto is None:
        return ""

    texto = (
        str(texto)
        .strip()
        .upper()
    )

    padrao = re.compile(
        r"\b"
        r"([A-Z0-9]{2,3})"
        r"-"
        r"([A-Z0-9]{4,5})"
        r"\b",
        re.IGNORECASE
    )

    match = padrao.search(
        texto
    )

    if not match:
        return ""

    parte_letras = (
        match.group(1)
        .upper()
    )

    parte_numeros = (
        match.group(2)
        .upper()
    )

    # --------------------------------------------------------
    # LETRAS
    # --------------------------------------------------------

    letras_corrigidas = []

    for char in parte_letras:

        if char.isalpha():

            letras_corrigidas.append(
                char
            )

        elif char in MAPA_LETRAS:

            letras_corrigidas.append(
                MAPA_LETRAS[
                    char
                ]
            )

        else:

            letras_corrigidas.append(
                char
            )

    # --------------------------------------------------------
    # NÚMEROS
    # --------------------------------------------------------

    numeros_corrigidos = []

    for char in parte_numeros:

        if char.isdigit():

            numeros_corrigidas = char

            numeros_corrigidos.append(
                char
            )

        elif char in MAPA_DIGITOS:

            numeros_corrigidos.append(
                MAPA_DIGITOS[
                    char
                ]
            )

        else:

            numeros_corrigidos.append(
                char
            )

    return (
        "".join(
            letras_corrigidas
        )
        +
        "-"
        +
        "".join(
            numeros_corrigidos
        )
    )


# ============================================================
# MÁSCARA DA PLACA
# ============================================================

def aplicar_mascara_placa(
    candidato,
    indices_originais=None,
):
    """
    Padrão:

        L L L D X D D

    A posição X (5ª posição) NUNCA é alterada.
    """

    candidato = normalizar_chave(
        candidato
    )

    if not candidato:
        return ""

    if indices_originais is None:

        indices_originais = list(
            range(
                len(candidato)
            )
        )

    resultado = []

    for indice_candidato, char in enumerate(
        candidato
    ):

        if indice_candidato >= len(
            indices_originais
        ):

            resultado.append(
                char
            )

            continue

        indice_original = (
            indices_originais[
                indice_candidato
            ]
        )

        # ----------------------------------------------------
        # 5ª POSIÇÃO
        # ----------------------------------------------------

        if indice_original == 4:

            resultado.append(
                char
            )

            continue

        # ----------------------------------------------------
        # LETRAS
        # ----------------------------------------------------

        if indice_original in {
            0,
            1,
            2,
        }:

            if char.isalpha():

                resultado.append(
                    char
                )

            elif char in MAPA_LETRAS:

                resultado.append(
                    MAPA_LETRAS[
                        char
                    ]
                )

            else:

                resultado.append(
                    char
                )

        # ----------------------------------------------------
        # NÚMEROS
        # ----------------------------------------------------

        else:

            if char.isdigit():

                resultado.append(
                    char
                )

            elif char in MAPA_DIGITOS:

                resultado.append(
                    MAPA_DIGITOS[
                        char
                    ]
                )

            else:

                resultado.append(
                    char
                )

    return "".join(
        resultado
    )


# ============================================================
# GERA CANDIDATOS
# ============================================================

def gerar_candidatos_placa(
    texto
):

    normalizado = normalizar_chave(
        texto
    )

    candidatos = []

    for tamanho in (
        7,
        6,
    ):

        if len(
            normalizado
        ) < tamanho:

            continue

        for inicio in range(
            0,
            len(normalizado)
            -
            tamanho
            +
            1
        ):

            candidato = normalizado[
                inicio:
                inicio
                +
                tamanho
            ]

            if candidato not in candidatos:

                candidatos.append(
                    candidato
                )

    return candidatos


# ============================================================
# PONTUAÇÃO DA PLACA
# ============================================================

def pontuar_candidato_placa(
    candidato,
    esperado,
):

    candidato = normalizar_chave(
        candidato
    )

    esperado = normalizar_chave(
        esperado
    )

    if len(
        esperado
    ) != 7:

        return None

    # ========================================================
    # 7 CARACTERES
    # ========================================================

    if len(
        candidato
    ) == 7:

        mascarado = (
            aplicar_mascara_placa(
                candidato
            )
        )

        iguais = sum(
            a == b
            for a, b in zip(
                mascarado,
                esperado
            )
        )

        if iguais >= 6:

            return {
                "pontuacao":
                    iguais,

                "candidato":
                    mascarado,
            }

        return None

    # ========================================================
    # 6 CARACTERES
    # ========================================================

    if len(
        candidato
    ) == 6:

        melhor = None

        for remover in range(
            7
        ):

            esperado_reduzido = (
                esperado[:remover]
                +
                esperado[remover + 1:]
            )

            indices_originais = [
                indice
                for indice
                in range(7)
                if indice != remover
            ]

            mascarado = (
                aplicar_mascara_placa(
                    candidato,
                    indices_originais
                )
            )

            iguais = sum(
                a == b
                for a, b in zip(
                    mascarado,
                    esperado_reduzido
                )
            )

            if iguais == 6:

                atual = {
                    "pontuacao":
                        6,

                    "candidato":
                        mascarado,
                }

                if (
                    melhor is None
                    or
                    atual[
                        "pontuacao"
                    ]
                    >
                    melhor[
                        "pontuacao"
                    ]
                ):

                    melhor = atual

        return melhor

    return None


# ============================================================
# VALIDA PLACA
# ============================================================

def validar_placa_ocr(
    textos,
    placas_esperadas,
):

    if not placas_esperadas:
        return None

    melhor = None

    for item in textos:

        texto = str(
            item.get(
                "texto",
                ""
            )
        ).strip()

        if not texto:
            continue

        candidatos = (
            gerar_candidatos_placa(
                texto
            )
        )

        try:

            confianca = float(
                item.get(
                    "confianca",
                    0.0
                )
                or
                0.0
            )

        except Exception:

            confianca = 0.0

        for candidato in candidatos:

            for placa_esperada in (
                placas_esperadas
            ):

                resultado = (
                    pontuar_candidato_placa(
                        candidato,
                        placa_esperada
                    )
                )

                if resultado is None:
                    continue

                atual = {
                    "placa_esperada":
                        placa_esperada,

                    "placa_ocr":
                        resultado[
                            "candidato"
                        ],

                    "pontuacao":
                        resultado[
                            "pontuacao"
                        ],

                    "texto_ocr":
                        texto,

                    "confianca":
                        confianca,
                }

                if (
                    melhor is None
                    or
                    atual[
                        "pontuacao"
                    ]
                    >
                    melhor[
                        "pontuacao"
                    ]
                    or
                    (
                        atual[
                            "pontuacao"
                        ]
                        ==
                        melhor[
                            "pontuacao"
                        ]
                        and
                        atual[
                            "confianca"
                        ]
                        >
                        melhor[
                            "confianca"
                        ]
                    )
                ):

                    melhor = atual

    return melhor


# ============================================================
# PROCESSO OCR
# ============================================================

def enviar_json(
    processo,
    dados
):

    if processo.poll() is not None:

        raise RuntimeError(
            "Processo OCR já encerrou. "
            f"Código: {processo.returncode}"
        )

    processo.stdin.write(
        json.dumps(
            dados,
            ensure_ascii=False
        )
        +
        "\n"
    )

    processo.stdin.flush()


def ler_json(
    processo
):

    linha = (
        processo.stdout.readline()
    )

    if not linha:

        raise RuntimeError(
            "O processo OCR encerrou "
            "antes de responder."
        )

    try:

        return json.loads(
            linha
        )

    except json.JSONDecodeError as erro:

        raise RuntimeError(
            "Resposta inválida do OCR: "
            f"{linha[:500]}"
        ) from erro


def iniciar_ocr():

    if not PYTHON_OCR.is_file():

        raise FileNotFoundError(
            "Python do ambiente OCR não encontrado:\n"
            f"{PYTHON_OCR}"
        )

    if not OCR_SAU_PY.is_file():

        raise FileNotFoundError(
            "OCR SAU não encontrado:\n"
            f"{OCR_SAU_PY}"
        )

    print(
        "\n"
        + "=" * 70
    )

    print(
        "INICIANDO OCR - VEÍCULOS ADMINISTRATIVOS"
    )

    print(
        "=" * 70
    )

    print(
        f"Python OCR: {PYTHON_OCR}"
    )

    print(
        f"Script OCR: {OCR_SAU_PY}\n"
    )

    return subprocess.Popen(
        [
            str(
                PYTHON_OCR
            ),
            str(
                OCR_SAU_PY
            ),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=None,
        text=True,
        bufsize=1,
        cwd=str(
            BASE_DIR
        ),
    )


def encerrar_ocr(
    processo
):

    if processo is None:
        return

    if processo.poll() is not None:
        return

    try:

        enviar_json(
            processo,
            {
                "command":
                    "exit"
            }
        )

        processo.wait(
            timeout=10
        )

    except Exception:

        try:

            processo.terminate()

            processo.wait(
                timeout=5
            )

        except Exception:

            try:
                processo.kill()

            except Exception:
                pass


# ============================================================
# EXTRAI OCR ITEMS
# ============================================================

def extrair_textos_resultado_ocr(
    resultado
):

    if not resultado:
        return []

    itens = (
        resultado.get(
            "ocr_items"
        )
    )

    if not isinstance(
        itens,
        list
    ):

        return []

    textos = []

    for item in itens:

        if not isinstance(
            item,
            dict
        ):

            continue

        texto = str(
            item.get(
                "texto",
                ""
            )
            or
            ""
        ).strip()

        if not texto:
            continue

        try:

            confianca = float(
                item.get(
                    "confianca",
                    0.0
                )
                or
                0.0
            )

        except Exception:

            confianca = 0.0

        textos.append(
            {
                "texto":
                    texto,

                "confianca":
                    confianca,

                "box":
                    item.get(
                        "box"
                    ),
            }
        )

    return textos


# ============================================================
# OCR DA IMAGEM
# ============================================================

def ler_imagem(
    processo_ocr,
    imagem,
):

    try:

        dados_imagem = (
            imagem.read_bytes()
        )

        imagem_b64 = (
            base64.b64encode(
                dados_imagem
            ).decode(
                "ascii"
            )
        )

    except Exception as erro:

        raise RuntimeError(
            f"Erro lendo a imagem "
            f"{imagem.name}: {erro}"
        ) from erro

    print(
        "\n"
        + "-" * 80
    )

    print(
        f"OCR DA IMAGEM: "
        f"{imagem.name}"
    )

    print(
        "-" * 80
    )

    enviar_json(
        processo_ocr,
        {
            "command":
                "ocr_batch",

            "images": [
                {
                    "task_id":
                        imagem.stem,

                    "crop_b64":
                        imagem_b64,
                }
            ],
        }
    )

    resposta = (
        ler_json(
            processo_ocr
        )
    )

    if resposta.get(
        "error"
    ):

        raise RuntimeError(
            f"OCR retornou erro: "
            f"{resposta['error']}"
        )

    resultados = (
        resposta.get(
            "results",
            []
        )
    )

    if not resultados:

        print(
            "[OCR] Nenhum resultado."
        )

        return []

    resultado = resultados[0]

    if resultado.get(
        "error"
    ):

        raise RuntimeError(
            f"OCR retornou erro para "
            f"{imagem.name}: "
            f"{resultado['error']}"
        )

    textos = (
        extrair_textos_resultado_ocr(
            resultado
        )
    )

    for item in textos:

        print(
            f"[TEXTO] "
            f"{item['texto']!r} | "
            f"CONF: "
            f"{item['confianca']:.4f}"
        )

    print(
        "-" * 80
    )

    return textos


# ============================================================
# ANALISA UMA IMAGEM
# ============================================================

def analisar_imagem(
    textos,
    prefixo_esperado,
    placas_esperadas,
):

    prefixo_ok = False

    prefixo_lido = None

    # ========================================================
    # PREFIXO
    # ========================================================

    for item in textos:

        texto = item.get(
            "texto",
            ""
        )

        prefixo_ocr = (
            aplicar_mascara_prefixo(
                texto
            )
        )

        if not prefixo_ocr:
            continue

        if (
            normalizar_chave(
                prefixo_ocr
            )
            ==
            normalizar_chave(
                prefixo_esperado
            )
        ):

            prefixo_ok = True

            prefixo_lido = (
                prefixo_esperado
            )

            break

    # ========================================================
    # PLACA
    # ========================================================

    placa_resultado = None

    if placas_esperadas:

        placa_resultado = (
            validar_placa_ocr(
                textos,
                placas_esperadas
            )
        )

    return {
        "prefixo_ok":
            prefixo_ok,

        "prefixo_lido":
            prefixo_lido,

        "placa_resultado":
            placa_resultado,
    }


# ============================================================
# MAIN
# ============================================================

def main():

    if len(
        sys.argv
    ) < 2:

        print(
            "\nUso:\n"
            '    py ler_imagem_administrativos.py '
            '"C:\\caminho\\da\\pasta"\n'
        )

        return

    pasta_imagens = Path(
        limpar_caminho_digitado(
            sys.argv[1]
        )
    )

    if not pasta_imagens.is_dir():

        print(
            "\nERRO: pasta inválida:\n"
            f"{pasta_imagens}\n"
        )

        return

    # ========================================================
    # CADASTRO
    # ========================================================

    cadastro = (
        carregar_cadastro_excel()
    )

    # ========================================================
    # OCR
    # ========================================================

    processo_ocr = (
        iniciar_ocr()
    )

    # ========================================================
    # IMAGENS
    # ========================================================

    arquivos = sorted(
        [
            arquivo
            for arquivo
            in pasta_imagens.rglob("*")
            if (
                arquivo.is_file()
                and
                arquivo.suffix.lower()
                in {
                    ".jpg",
                    ".jpeg",
                    ".png",
                }
            )
        ],
        key=lambda p:
            str(p.relative_to(pasta_imagens)).lower()
    )

    print(
        f"[{len(arquivos)}] imagens encontradas. "
        "Iniciando processamento..."
    )

    # ========================================================
    # PADRÃO DO ARQUIVO
    # ========================================================

    padrao_arquivo = re.compile(
        r"^("
        r"[A-Z]{2,3}"
        r"-"
        r"\d{4,5}"
        r")"
        r"[_-]"
        r"(LD|LE|FR|TR)"
        r"(?:[_-]|\.[^.]*$|$)",
        re.IGNORECASE
    )

    # ========================================================
    # VEÍCULOS
    # ========================================================

    veiculos = {}

    try:

        for idx, imagem in enumerate(
            arquivos,
            start=1
        ):

            nome_arquivo = (
                imagem.name
            )

            match = (
                padrao_arquivo.search(
                    nome_arquivo
                )
            )

            if not match:

                print(
                    f"[{idx}/{len(arquivos)}] "
                    f"Ignorado: "
                    f"{nome_arquivo}"
                )

                continue

            prefixo = (
                match.group(1)
                .upper()
            )

            lado = (
                match.group(2)
                .upper()
            )

            registro = cadastro.get(
                normalizar_chave(
                    prefixo
                )
            )

            if registro:

                placas_esperadas = (
                    registro[
                        "placas"
                    ]
                )

            else:

                placas_esperadas = []

            print(
                f"\n[{idx}/{len(arquivos)}] "
                f"Analisando {prefixo} "
                f"(Lado: {lado})"
            )

            textos = ler_imagem(
                processo_ocr,
                imagem
            )

            analise = (
                analisar_imagem(
                    textos,
                    prefixo,
                    placas_esperadas
                )
            )

            # =================================================
            # LOG PREFIXO
            # =================================================

            if analise[
                "prefixo_ok"
            ]:

                print(
                    f"PREFIXO OK "
                    f"({prefixo})"
                )

            else:

                print(
                    "PREFIXO NÃO ENCONTRADO"
                )

            # =================================================
            # LOG PLACA
            # =================================================

            placa_resultado = (
                analise[
                    "placa_resultado"
                ]
            )

            if placa_resultado:

                print(
                    f"PLACA OK "
                    f"({formatar_placa(placa_resultado['placa_ocr'])} | "
                    f"{placa_resultado['pontuacao']}/7)"
                )

            elif placas_esperadas:

                print(
                    "PLACA NÃO VALIDADA"
                )

            else:

                print(
                    "SEM CADASTRO DE PLACA"
                )

            # =================================================
            # CRIA VEÍCULO
            # =================================================

            if prefixo not in veiculos:

                veiculos[
                    prefixo
                ] = {

                    "placas":
                        placas_esperadas,

                    "lados": {

                        "FR":
                            {
                                "prefixo_ok":
                                    False
                            },

                        "LD":
                            {
                                "prefixo_ok":
                                    False
                            },

                        "LE":
                            {
                                "prefixo_ok":
                                    False
                            },

                        "TR":
                            {
                                "prefixo_ok":
                                    False
                            },
                    },

                    "melhor_placa":
                        None,
                }

            veiculo = (
                veiculos[
                    prefixo
                ]
            )

            # =================================================
            # SALVA RESULTADO DO LADO
            # =================================================

            veiculo[
                "lados"
            ][
                lado
            ] = {

                "prefixo_ok":
                    analise[
                        "prefixo_ok"
                    ],

                "prefixo_lido":
                    analise[
                        "prefixo_lido"
                    ],
            }

            # =================================================
            # MELHOR PLACA DO VEÍCULO
            # =================================================

            if placa_resultado:

                melhor_atual = (
                    veiculo[
                        "melhor_placa"
                    ]
                )

                if (
                    melhor_atual is None
                    or
                    placa_resultado[
                        "pontuacao"
                    ]
                    >
                    melhor_atual[
                        "pontuacao"
                    ]
                    or
                    (
                        placa_resultado[
                            "pontuacao"
                        ]
                        ==
                        melhor_atual[
                            "pontuacao"
                        ]
                        and
                        placa_resultado[
                            "confianca"
                        ]
                        >
                        melhor_atual[
                            "confianca"
                        ]
                    )
                ):

                    veiculo[
                        "melhor_placa"
                    ] = (
                        placa_resultado
                    )

    finally:

        encerrar_ocr(
            processo_ocr
        )

    # ========================================================
    # CONSOLIDAÇÃO
    # ========================================================

    print(
        "\n"
        + "=" * 70
    )

    print(
        "CONSOLIDANDO DADOS E GERANDO PLANILHA"
    )

    print(
        "=" * 70
    )

    colunas = [

        "Prefixo",

        "Frente (FR)",

        "Direita (LD)",

        "Esquerda (LE)",

        "Trás (TR)",

        "Placa informada",

        "Placa OCR",

        "Status placa",
    ]

    linhas = []

    for prefixo, veiculo in (
        veiculos.items()
    ):

        placas = veiculo[
            "placas"
        ]

        if placas:

            placa_informada = (
                " / ".join(
                    formatar_placa(
                        placa
                    )
                    for placa in placas
                )
            )

        else:

            placa_informada = (
                "Não cadastrada"
            )

        melhor_placa = veiculo[
            "melhor_placa"
        ]

        if melhor_placa:

            placa_ocr = formatar_placa(
                melhor_placa[
                    "placa_ocr"
                ]
            )

            status_placa = (
                f"OK "
                f"({melhor_placa['pontuacao']}/7)"
            )

        else:

            placa_ocr = (
                "Não validada"
            )

            if placas:

                status_placa = (
                    "Não encontrada"
                )

            else:

                status_placa = (
                    "Sem cadastro"
                )

        # ====================================================
        # PREFIXOS DOS 4 LADOS
        # ====================================================

        valores_lados = {}

        for lado in (
            "FR",
            "LD",
            "LE",
            "TR",
        ):

            resultado_lado = (
                veiculo[
                    "lados"
                ].get(
                    lado,
                    {}
                )
            )

            if resultado_lado.get(
                "prefixo_ok",
                False
            ):

                # Quando está correto, escreve o prefixo
                # esperado/cadastrado.

                valores_lados[
                    lado
                ] = prefixo

            else:

                valores_lados[
                    lado
                ] = "Não encontrado"

        linhas.append(
            {
                "Prefixo":
                    prefixo,

                "Frente (FR)":
                    valores_lados["FR"],

                "Direita (LD)":
                    valores_lados["LD"],

                "Esquerda (LE)":
                    valores_lados["LE"],

                "Trás (TR)":
                    valores_lados["TR"],

                "Placa informada":
                    placa_informada,

                "Placa OCR":
                    placa_ocr,

                "Status placa":
                    status_placa,
            }
        )

    # ========================================================
    # DATAFRAME
    # ========================================================

    df = pd.DataFrame(
        linhas,
        columns=colunas
    )

    # ========================================================
    # DIRETÓRIO
    # ========================================================

    RESULTADOS_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    caminho_saida = (
        RESULTADOS_DIR
        /
        NOME_SAIDA
    )

    # ========================================================
    # EXCEL
    # ========================================================

    workbook = Workbook()

    ws = (
        workbook.active
    )

    ws.title = (
        "Administrativos"
    )

    # ========================================================
    # CABEÇALHO
    # ========================================================

    for coluna, nome in enumerate(
        colunas,
        start=1
    ):

        celula = ws.cell(
            row=1,
            column=coluna,
            value=nome
        )

        celula.font = Font(
            bold=True
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ========================================================
    # DADOS
    # ========================================================

    for numero_linha, (
        indice,
        registro
    ) in enumerate(
        df.iterrows(),
        start=2
    ):

        for numero_coluna, nome in enumerate(
            colunas,
            start=1
        ):

            valor = registro[
                nome
            ]

            celula = ws.cell(
                row=numero_linha,
                column=numero_coluna,
                value=valor
            )

            celula.alignment = Alignment(
                vertical="center",
                horizontal="center"
            )

    # ========================================================
    # CORES DOS 4 LADOS
    # ========================================================

    colunas_lados = {
        "Frente (FR)": 2,
        "Direita (LD)": 3,
        "Esquerda (LE)": 4,
        "Trás (TR)": 5,
    }

    preenchimento_verde = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE"
    )

    preenchimento_vermelho = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE"
    )

    fonte_verde = Font(
        bold=True,
        color="006100"
    )

    fonte_vermelha = Font(
        bold=True,
        color="9C0006"
    )

    for linha in range(
        2,
        ws.max_row + 1
    ):

        for nome_coluna, numero_coluna in (
            colunas_lados.items()
        ):

            celula = ws.cell(
                row=linha,
                column=numero_coluna
            )

            if (
                celula.value
                and
                celula.value
                !=
                "Não encontrado"
            ):

                celula.fill = (
                    preenchimento_verde
                )

                celula.font = (
                    fonte_verde
                )

            else:

                celula.value = (
                    "Não encontrado"
                )

                celula.fill = (
                    preenchimento_vermelho
                )

                celula.font = (
                    fonte_vermelha
                )

    # ========================================================
    # COR DO STATUS DA PLACA
    # ========================================================

    coluna_status = 8

    for linha in range(
        2,
        ws.max_row + 1
    ):

        celula = ws.cell(
            row=linha,
            column=coluna_status
        )

        if str(
            celula.value
        ).startswith(
            "OK"
        ):

            celula.fill = (
                preenchimento_verde
            )

            celula.font = (
                fonte_verde
            )

        else:

            celula.fill = (
                preenchimento_vermelho
            )

            celula.font = (
                fonte_vermelha
            )

    # ========================================================
    # LARGURAS
    # ========================================================

    larguras = {

        "Prefixo":
            18,

        "Frente (FR)":
            18,

        "Direita (LD)":
            18,

        "Esquerda (LE)":
            18,

        "Trás (TR)":
            18,

        "Placa informada":
            22,

        "Placa OCR":
            18,

        "Status placa":
            20,
    }

    for nome, largura in (
        larguras.items()
    ):

        numero_coluna = (
            colunas.index(
                nome
            )
            + 1
        )

        ws.column_dimensions[
            get_column_letter(
                numero_coluna
            )
        ].width = largura

    # ========================================================
    # CONGELA CABEÇALHO
    # ========================================================

    ws.freeze_panes = (
        "A2"
    )

    # ========================================================
    # FILTRO
    # ========================================================

    ws.auto_filter.ref = (
        ws.dimensions
    )

    # ========================================================
    # SALVA
    # ========================================================

    workbook.save(
        caminho_saida
    )

    # ========================================================
    # RESUMO
    # ========================================================

    total_veiculos = len(
        veiculos
    )

    total_prefixos_fr = sum(
        1
        for v in veiculos.values()
        if v[
            "lados"
        ][
            "FR"
        ].get(
            "prefixo_ok",
            False
        )
    )

    total_prefixos_ld = sum(
        1
        for v in veiculos.values()
        if v[
            "lados"
        ][
            "LD"
        ].get(
            "prefixo_ok",
            False
        )
    )

    total_prefixos_le = sum(
        1
        for v in veiculos.values()
        if v[
            "lados"
        ][
            "LE"
        ].get(
            "prefixo_ok",
            False
        )
    )

    total_prefixos_tr = sum(
        1
        for v in veiculos.values()
        if v[
            "lados"
        ][
            "TR"
        ].get(
            "prefixo_ok",
            False
        )
    )

    total_placas = sum(
        1
        for v in veiculos.values()
        if v[
            "melhor_placa"
        ]
    )

    print(
        "\n"
        + "=" * 70
    )

    print(
        "RESULTADO FINAL"
    )

    print(
        "=" * 70
    )

    print(
        f"Veículos processados: "
        f"{total_veiculos}"
    )

    print(
        f"Prefixos OK - Frente: "
        f"{total_prefixos_fr}"
    )

    print(
        f"Prefixos OK - Direita: "
        f"{total_prefixos_ld}"
    )

    print(
        f"Prefixos OK - Esquerda: "
        f"{total_prefixos_le}"
    )

    print(
        f"Prefixos OK - Trás: "
        f"{total_prefixos_tr}"
    )

    print(
        f"Placas validadas: "
        f"{total_placas}"
    )

    print(
        f"Planilha gerada em:"
        f"\n{caminho_saida}"
    )

    print(
        "=" * 70
    )


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    main()


