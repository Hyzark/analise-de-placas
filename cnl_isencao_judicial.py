from __future__ import annotations
from placas import (
    aplicar_mascara as _aplicar_mascara_comum,
    descobrir_mascara as _descobrir_mascara_comum,
    distancia_placas as _distancia_placas_comum,
    distancia_ponderada as _distancia_ponderada,
    normalizar_ocr as _normalizar_ocr_comum,
    normalizar_placa as _normalizar_placa,
    CONFUNDIVEIS as _CONFUNDIVEIS,
    MAPA_LETRAS as _MAPA_LETRAS,
    MAPA_DIGITOS as _MAPA_DIGITOS,
)

from common import extrair_bytes_imagem, enviar_json

import functools
import contextlib
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

from config import (
    BASE_DIR,
    caminho_resultado,
    PYTHON_YOLO,
    PYTHON_OCR,
    DETECTOR_PY,
    OCR_PY,
)


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BATCH_SIZE = 64

# 0, 1 ou 2 diferenças = APROVADO
MAX_DIFERENCAS_APROVACAO = 2

MAX_CARACTERES_OCR = 100

MAX_IMAGENS_POR_PASSAGEM = 50

TAMANHO_PLACA = 7


# ============================================================
# COLUNAS DA SAÍDA
# ============================================================

CAMPOS_DADOS_PASSAGEM = [
    "ID_PASSAGEM",
    "Data/hora da passagem",
    "Praça de pedágio",
    "Rodovia",
    "km",
    "Sentido",
    "Faixa",
    "Placa no sistema",
    "Categoria",
    "Marca",
    "Modelo",
    "Cor",
    "Valor R$",
]

CAMPOS_OCR = [
    "OCR Bruto",
    "OCR Limpo",
    "OCR Confidence",
]

CAMPOS_COMPARACAO = [
    "Diferença OCR",
    "Diferenças de caracteres",
    "Resultado",
]

CAMPOS_SAIDA = (
    CAMPOS_DADOS_PASSAGEM
    + CAMPOS_OCR
    + CAMPOS_COMPARACAO
)


# ============================================================
# CONFUNDÍVEIS
# ============================================================



# ============================================================
# MAPAS DE CORREÇÃO POR POSIÇÃO
# ============================================================




# ============================================================
# LEVENSHTEIN
# ============================================================

def _levenshtein(a: str, b: str) -> float:
    return _distancia_ponderada(a, b, _CONFUNDIVEIS)


# ============================================================
# LIMPA TEXTO
# ============================================================

def limpar_texto(valor) -> str:
    return _normalizar_placa(valor)


# ============================================================
# DISTÂNCIA
# ============================================================

def distancia_placas(placa_a: str, placa_b: str) -> float:
    return _distancia_placas_comum(placa_a, placa_b, _CONFUNDIVEIS)


# ============================================================
# DESCOBRE MÁSCARA DA PLACA
# ============================================================

def descobrir_mascara(placa: str) -> str:
    return _descobrir_mascara_comum(placa, tamanho=TAMANHO_PLACA, mascara_invalida="LLLDLDD")


# ============================================================
# APLICA MÁSCARA DA PLACA
# ============================================================

def aplicar_mascara_placa(candidato: str, mascara: str) -> str:
    return _aplicar_mascara_comum(candidato, mascara, _MAPA_LETRAS, _MAPA_DIGITOS, tamanho=TAMANHO_PLACA, comprimento_invalido="manter")


# ============================================================
# ENCONTRA MELHOR TRECHO DE 7 CARACTERES
# ============================================================

def encontrar_melhor_trecho(
    placa_esperada,
    placa_ocr,
):

    esperada = limpar_texto(
        placa_esperada
    )

    ocr = limpar_texto(
        placa_ocr
    )

    if (
        not esperada
        or not ocr
    ):

        return (
            "",
            999.0,
        )

    # --------------------------------------------------------
    # SE A PLACA DO SISTEMA NÃO TIVER 7 CARACTERES
    # --------------------------------------------------------

    if len(esperada) != TAMANHO_PLACA:

        return (
            ocr,
            distancia_placas(
                esperada,
                ocr,
            ),
        )

    mascara = descobrir_mascara(
        esperada
    )

    # --------------------------------------------------------
    # OCR COM 7 CARACTERES
    # --------------------------------------------------------

    if len(ocr) == TAMANHO_PLACA:

        corrigido = aplicar_mascara_placa(
            ocr,
            mascara,
        )

        distancia_original = (
            distancia_placas(
                esperada,
                ocr,
            )
        )

        distancia_corrigida = (
            distancia_placas(
                esperada,
                corrigido,
            )
        )

        if distancia_corrigida <= distancia_original:

            return (
                corrigido,
                distancia_corrigida,
            )

        return (
            ocr,
            distancia_original,
        )

    # --------------------------------------------------------
    # OCR MENOR QUE 7
    # --------------------------------------------------------

    if len(ocr) < TAMANHO_PLACA:

        return (
            ocr,
            distancia_placas(
                esperada,
                ocr,
            ),
        )

    # --------------------------------------------------------
    # OCR MAIOR QUE 7
    #
    # Procura a melhor janela de 7.
    # --------------------------------------------------------

    candidatos = []

    for inicio in range(
        len(ocr) - TAMANHO_PLACA + 1
    ):

        trecho = ocr[
            inicio:
            inicio + TAMANHO_PLACA
        ]

        corrigido = aplicar_mascara_placa(
            trecho,
            mascara,
        )

        distancia_original = (
            distancia_placas(
                esperada,
                trecho,
            )
        )

        distancia_corrigida = (
            distancia_placas(
                esperada,
                corrigido,
            )
        )

        # ----------------------------------------------------
        # PRIORIZAÇÃO
        #
        # 0 = trecho bruto exatamente igual
        # 1 = trecho corrigido exatamente igual
        # 2 = menor distância
        # ----------------------------------------------------

        if trecho == esperada:

            prioridade = 0

        elif corrigido == esperada:

            prioridade = 1

        else:

            prioridade = 2

        distancia = min(
            distancia_original,
            distancia_corrigida,
        )

        candidato_final = (
            corrigido
            if distancia_corrigida
            <= distancia_original
            else trecho
        )

        candidatos.append(
            (
                prioridade,
                distancia,
                inicio,
                candidato_final,
            )
        )

    candidatos.sort(
        key=lambda item: (
            item[0],
            item[1],
            item[2],
        )
    )

    (
        _,
        melhor_distancia,
        _,
        melhor_trecho,
    ) = candidatos[0]

    return (
        melhor_trecho,
        melhor_distancia,
    )


# ============================================================
# DESCREVE DIFERENÇAS DE CARACTERES
# ============================================================

def descrever_diferencas(
    placa_esperada,
    placa_ocr,
) -> str:

    esperada = limpar_texto(
        placa_esperada
    )

    ocr = limpar_texto(
        placa_ocr
    )

    if not esperada or not ocr:
        return "Sem comparação"

    # --------------------------------------------------------
    # MATRIZ DE DISTÂNCIA + BACKTRACE
    # --------------------------------------------------------

    linhas = len(esperada) + 1
    colunas = len(ocr) + 1

    matriz = [
        [0.0 for _ in range(colunas)]
        for _ in range(linhas)
    ]

    operacoes = [
        [None for _ in range(colunas)]
        for _ in range(linhas)
    ]

    for i in range(linhas):
        matriz[i][0] = float(i)

    for j in range(colunas):
        matriz[0][j] = float(j)

    for i in range(1, linhas):

        for j in range(1, colunas):

            ca = esperada[i - 1]
            cb = ocr[j - 1]

            if ca == cb:

                custo_substituicao = 0.0

            elif (
                ca,
                cb,
            ) in _CONFUNDIVEIS:

                custo_substituicao = 0.35

            else:

                custo_substituicao = 1.0

            substituicao = (
                matriz[i - 1][j - 1]
                +
                custo_substituicao
            )

            insercao = (
                matriz[i][j - 1]
                +
                1.0
            )

            remocao = (
                matriz[i - 1][j]
                +
                1.0
            )

            menor = min(
                substituicao,
                insercao,
                remocao,
            )

            matriz[i][j] = menor

            if menor == substituicao:

                if custo_substituicao == 0:

                    operacoes[i][j] = (
                        "igual",
                        ca,
                        cb,
                    )

                else:

                    operacoes[i][j] = (
                        "substituicao",
                        ca,
                        cb,
                    )

            elif menor == insercao:

                operacoes[i][j] = (
                    "insercao",
                    "",
                    cb,
                )

            else:

                operacoes[i][j] = (
                    "remocao",
                    ca,
                    "",
                )

    # --------------------------------------------------------
    # BACKTRACE
    # --------------------------------------------------------

    i = len(esperada)
    j = len(ocr)

    diferencas = []

    while i > 0 or j > 0:

        if i > 0 and j > 0:

            operacao = operacoes[i][j]
            tipo = operacao[0]

        elif i > 0:

            tipo = "remocao"

        else:

            tipo = "insercao"

        if tipo == "igual":

            i -= 1
            j -= 1

        elif tipo == "substituicao":

            posicao = i

            esperado = esperada[i - 1]
            recebido = ocr[j - 1]

            diferencas.append(
                (
                    posicao,
                    recebido,
                    esperado,
                )
            )

            i -= 1
            j -= 1

        elif tipo == "insercao":

            posicao = j
            recebido = ocr[j - 1]

            diferencas.append(
                (
                    posicao,
                    recebido,
                    "—",
                )
            )

            j -= 1

        elif tipo == "remocao":

            posicao = i
            esperado = esperada[i - 1]

            diferencas.append(
                (
                    posicao,
                    "—",
                    esperado,
                )
            )

            i -= 1

    diferencas.reverse()

    if not diferencas:
        return "Nenhuma"

    resultado = []

    for posicao, recebido, esperado in (
        diferencas
    ):

        if esperado == "—":

            resultado.append(
                f"Posição {posicao}: "
                f"OCR={recebido} "
                f"(caractere extra)"
            )

        elif recebido == "—":

            resultado.append(
                f"Posição {posicao}: "
                f"OCR ausente "
                f"(esperado={esperado})"
            )

        else:

            resultado.append(
                f"Posição {posicao}: "
                f"OCR={recebido} "
                f"≠ Sistema={esperado}"
            )

    return "; ".join(
        resultado
    )


# ============================================================
# NORMALIZA CABEÇALHO
# ============================================================

def normalizar_coluna(
    nome,
) -> str:

    if nome is None:
        return ""

    texto = str(
        nome
    ).strip().lower()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto


# ============================================================
# ALIASES
# ============================================================

ALIASES_COLUNAS = {

    "id_passagem": [
        "id_passagem",
        "id passagem",
        "idpassagem",
        "id",
        "passagem",
        "id transacao",
        "id transação",
    ],

    "data_hora": [
        "data/hora da passagem",
        "data hora da passagem",
        "data/hora",
        "data hora",
        "data_hora",
        "data e hora",
        "data e hora da passagem",
    ],

    "praca": [
        "praça de pedágio",
        "praca de pedagio",
        "praça pedágio",
        "praca pedagio",
        "praça",
        "praca",
    ],

    "rodovia": [
        "rodovia",
    ],

    "km": [
        "km",
        "km.",
        "quilometro",
        "quilômetro",
    ],

    "sentido": [
        "sentido",
    ],

    "faixa": [
        "faixa",
    ],

    "placa_sistema": [
        "placa no sistema",
        "placa_no_sistema",
        "placa sistema",
        "placa",
    ],

    "categoria": [
        "categoria",
        "categoria cca",
        "cat",
        "classe veículo",
        "classe veiculo",
        "classeveiculo",
    ],

    "marca": [
        "marca",
    ],

    "modelo": [
        "modelo",
    ],

    "cor": [
        "cor",
    ],

    "valor": [
        "valor r$",
        "valor",
        "valor rs",
        "valor tarifa",
        "tarifa",
    ],
}


# ============================================================
# LOCALIZA COLUNAS
# ============================================================

def localizar_colunas(
    sheet,
):

    colunas = {}

    for coluna in range(
        1,
        sheet.max_column + 1,
    ):

        cabecalho = sheet.cell(
            row=1,
            column=coluna,
        ).value

        normalizado = normalizar_coluna(
            cabecalho
        )

        if normalizado:

            colunas[
                normalizado
            ] = coluna

    resultado = {}

    for campo, aliases in (
        ALIASES_COLUNAS.items()
    ):

        for alias in aliases:

            alias_normalizado = (
                normalizar_coluna(
                    alias
                )
            )

            if (
                alias_normalizado
                in colunas
            ):

                resultado[
                    campo
                ] = colunas[
                    alias_normalizado
                ]

                break

    return resultado


# ============================================================
# EXTRAI IMAGEM
# ============================================================



# ============================================================
# EXTENSÃO
# ============================================================

def descobrir_extensao(
    imagem,
    dados,
):

    formato = getattr(
        imagem,
        "format",
        None,
    )

    if formato:

        formato = str(
            formato
        ).lower()

        if formato in (
            "jpg",
            "jpeg",
        ):

            return ".jpg"

        if formato == "png":

            return ".png"

        if formato == "webp":

            return ".webp"

        if formato == "bmp":

            return ".bmp"

        if formato in (
            "tif",
            "tiff",
        ):

            return ".tiff"

    if dados[:3] == (
        b"\xff\xd8\xff"
    ):

        return ".jpg"

    if dados[:8] == (
        b"\x89PNG\r\n\x1a\n"
    ):

        return ".png"

    if dados[:4] == b"RIFF":

        return ".webp"

    return ".png"


# ============================================================
# JSON
# ============================================================



def ler_json(
    processo,
    nome,
):

    linha = processo.stdout.readline()

    if not linha:

        raise RuntimeError(
            f"Processo {nome} "
            "encerrou antes de responder."
        )

    try:

        return json.loads(
            linha
        )

    except json.JSONDecodeError as erro:

        raise RuntimeError(
            f"Resposta inválida do processo "
            f"{nome}: {linha[:500]}"
        ) from erro


# ============================================================
# ENCERRA
# ============================================================

def encerrar_processo(
    processo,
    nome,
):

    if processo is None:
        return

    if processo.poll() is not None:
        return

    try:

        enviar_json(
            processo,
            {
                "command": "exit"
            },
        )

        processo.wait(
            timeout=10
        )

    except Exception:

        try:

            processo.terminate()

            processo.wait(
                timeout=5
            )

        except Exception:

            try:

                processo.kill()

            except Exception:
                pass


# ============================================================
# INICIA YOLO
# ============================================================

def iniciar_yolo():

    print(
        "[INIT] Iniciando YOLO..."
    )

    return subprocess.Popen(
        [
            str(PYTHON_YOLO),
            str(DETECTOR_PY),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )


# ============================================================
# INICIA OCR
# ============================================================

def iniciar_ocr():

    print(
        "[INIT] Iniciando PaddleOCR..."
    )

    return subprocess.Popen(
        [
            str(PYTHON_OCR),
            str(OCR_PY),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )


# ============================================================
# CARREGA EXCEL
# ============================================================

def carregar_planilha(
    caminho_excel,
    temp_dir,
):

    print(
        f"\n[EXCEL] {caminho_excel}"
    )

    workbook = openpyxl.load_workbook(
        caminho_excel,
        data_only=True,
    )

    if "Conferencia" in workbook.sheetnames:

        sheet = workbook[
            "Conferencia"
        ]

    else:

        sheet = workbook.active

    colunas = localizar_colunas(
        sheet
    )

    print(
        "[EXCEL] Colunas localizadas:"
    )

    print(
        colunas
    )

    if "id_passagem" not in colunas:

        raise ValueError(
            "A coluna ID_PASSAGEM não foi encontrada."
        )

    if "placa_sistema" not in colunas:

        raise ValueError(
            "A coluna 'Placa no sistema' "
            "não foi encontrada."
        )

    # ========================================================
    # IMAGENS POR LINHA
    # ========================================================

    imagens_por_linha = {}

    for indice_imagem, imagem in enumerate(
        sheet._images
    ):

        try:

            linha = (
                imagem.anchor._from.row
                + 1
            )

        except Exception:

            continue

        imagens_por_linha.setdefault(
            linha,
            [],
        ).append(
            (
                indice_imagem,
                imagem,
            )
        )

    print(
        f"[EXCEL] Imagens encontradas: "
        f"{len(sheet._images)}"
    )

    # ========================================================
    # AGRUPAMENTO
    # ========================================================

    grupos = {}
    ordem_ids = []

    for linha in range(
        2,
        sheet.max_row + 1,
    ):

        id_valor = sheet.cell(
            row=linha,
            column=colunas[
                "id_passagem"
            ],
        ).value

        id_limpo = limpar_texto(
            id_valor
        )

        if not id_limpo:
            continue

        if id_limpo not in grupos:

            dados_saida = {}

            for campo in (
                "id_passagem",
                "data_hora",
                "praca",
                "rodovia",
                "km",
                "sentido",
                "faixa",
                "placa_sistema",
                "categoria",
                "marca",
                "modelo",
                "cor",
                "valor",
            ):

                coluna = colunas.get(
                    campo
                )

                if coluna is None:

                    dados_saida[
                        campo
                    ] = ""

                else:

                    dados_saida[
                        campo
                    ] = sheet.cell(
                        row=linha,
                        column=coluna,
                    ).value

            grupos[
                id_limpo
            ] = {

                "id_passagem":
                    dados_saida[
                        "id_passagem"
                    ],

                "dados_saida":
                    dados_saida,

                "placas_sistema":
                    [],

                "imagens":
                    [],
            }

            ordem_ids.append(
                id_limpo
            )

        grupo = grupos[
            id_limpo
        ]

        # ----------------------------------------------------
        # PLACAS DO SISTEMA
        # ----------------------------------------------------

        placa = sheet.cell(
            row=linha,
            column=colunas[
                "placa_sistema"
            ],
        ).value

        placa_limpa = limpar_texto(
            placa
        )

        if (
            placa_limpa
            and
            placa_limpa
            not in grupo[
                "placas_sistema"
            ]
        ):

            grupo[
                "placas_sistema"
            ].append(
                placa_limpa
            )

        # ----------------------------------------------------
        # IMAGENS
        # ----------------------------------------------------

        for indice_imagem, imagem in (
            imagens_por_linha.get(
                linha,
                [],
            )
        ):

            dados = extrair_bytes_imagem(
                imagem
            )

            if not dados:
                continue

            extensao = descobrir_extensao(
                imagem,
                dados,
            )

            caminho = (
                temp_dir
                /
                (
                    f"id_{id_limpo}_"
                    f"img_{indice_imagem}"
                    f"{extensao}"
                )
            )

            caminho.write_bytes(
                dados
            )

            grupo[
                "imagens"
            ].append(
                {
                    "task_id":
                        (
                            f"id_{id_limpo}_"
                            f"img_{indice_imagem}"
                        ),

                    "image_path":
                        str(caminho),

                    "indice":
                        indice_imagem,
                }
            )

    # ========================================================
    # LISTA FINAL
    # ========================================================

    passagens = []

    for id_limpo in ordem_ids:

        grupo = grupos[
            id_limpo
        ]

        imagens = grupo[
            "imagens"
        ]

        if (
            len(imagens)
            > MAX_IMAGENS_POR_PASSAGEM
        ):

            imagens = imagens[
                :MAX_IMAGENS_POR_PASSAGEM
            ]

        passagens.append(
            {
                "id":
                    id_limpo,

                "id_passagem":
                    grupo[
                        "id_passagem"
                    ],

                "dados_saida":
                    grupo[
                        "dados_saida"
                    ],

                "placas_sistema":
                    grupo[
                        "placas_sistema"
                    ],

                "imagens":
                    imagens,
            }
        )

    print(
        f"[EXCEL] IDs únicos: "
        f"{len(passagens)}"
    )

    return passagens


# ============================================================
# YOLO
# ============================================================

def executar_yolo(
    processo_yolo,
    imagens,
):

    if not imagens:
        return {}

    items = []

    for imagem in imagens:

        caminho = Path(
            imagem[
                "image_path"
            ]
        ).resolve()

        if (
            not caminho.exists()
            or
            caminho.stat().st_size == 0
        ):

            continue

        items.append(
            {
                "task_id":
                    imagem[
                        "task_id"
                    ],

                "image_path":
                    str(caminho),
            }
        )

    if not items:
        return {}

    enviar_json(
        processo_yolo,
        {
            "command":
                "detect_batch_all",

            "images":
                items,
        },
    )

    resposta = ler_json(
        processo_yolo,
        "YOLO",
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        [],
    ):

        task_id = resultado.get(
            "task_id"
        )

        if task_id:

            resultados[
                task_id
            ] = resultado

    return resultados


# ============================================================
# OCR
# ============================================================

def executar_ocr(
    processo_ocr,
    yolo_results,
):

    if not yolo_results:
        return {}

    crops = []

    for task_id, resultado in (
        yolo_results.items()
    ):

        if not resultado.get(
            "detected",
            False,
        ):

            continue

        detections = resultado.get(
            "detections",
            [],
        )

        # Compatibilidade
        if not detections:

            crop_b64 = resultado.get(
                "crop_b64"
            )

            if crop_b64:

                detections = [
                    {
                        "detection_index":
                            1,

                        "crop_b64":
                            crop_b64,
                    }
                ]

        for detection in detections:

            detection_index = (
                detection.get(
                    "detection_index",
                    1,
                )
            )

            crop_b64 = detection.get(
                "crop_b64"
            )

            if not crop_b64:
                continue

            crops.append(
                {
                    "task_id":
                        (
                            f"{task_id}_"
                            f"d{detection_index}"
                        ),

                    "crop_b64":
                        crop_b64,
                }
            )

    if not crops:
        return {}

    enviar_json(
        processo_ocr,
        {
            "command":
                "ocr_batch",

            "images":
                crops,
        },
    )

    resposta = ler_json(
        processo_ocr,
        "OCR",
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        [],
    ):

        task_id = resultado.get(
            "task_id"
        )

        if task_id:

            resultados[
                task_id
            ] = resultado

    return resultados


# ============================================================
# AVALIA OCR
# ============================================================

def avaliar_ocr(
    placa_planilha,
    resultado_ocr,
):

    placa_planilha = limpar_texto(
        placa_planilha
    )

    ocr_bruto = (
        resultado_ocr.get(
            "ocr_best_original",
            "",
        )
    )

    ocr_best = (
        resultado_ocr.get(
            "ocr_best",
            "",
        )
    )

    # --------------------------------------------------------
    # LIMITA OCR EXCESSIVAMENTE GRANDE
    # --------------------------------------------------------

    ocr_best = str(
        ocr_best or ""
    )[:MAX_CARACTERES_OCR]

    # --------------------------------------------------------
    # CONFIDENCE
    # --------------------------------------------------------

    try:

        confianca = float(
            resultado_ocr.get(
                "confidence_best",
                0.0,
            )
            or 0.0
        )

    except Exception:

        confianca = 0.0

    # --------------------------------------------------------
    # SEM OCR
    # --------------------------------------------------------

    if not limpar_texto(
        ocr_best
    ):

        return {
            "aprovado":
                False,

            "ocr_bruto":
                ocr_bruto,

            "ocr_limpo":
                "",

            "confianca":
                confianca,

            "diferenca":
                999.0,

            "diferencas_caracteres":
                "Sem OCR",
        }

    # --------------------------------------------------------
    # NORMALIZA OCR
    # --------------------------------------------------------

    ocr_normalizado = limpar_texto(
        ocr_best
    )

    # --------------------------------------------------------
    # ENCONTRA MELHOR TRECHO
    # --------------------------------------------------------

    ocr_limpo, diferenca = (
        encontrar_melhor_trecho(
            placa_planilha,
            ocr_normalizado,
        )
    )

    # --------------------------------------------------------
    # DIFERENÇAS
    # --------------------------------------------------------

    diferencas = (
        descrever_diferencas(
            placa_planilha,
            ocr_limpo,
        )
    )

    # --------------------------------------------------------
    # APROVAÇÃO
    # --------------------------------------------------------

    aprovado = (
        diferenca
        <= MAX_DIFERENCAS_APROVACAO
    )

    return {
        "aprovado":
            aprovado,

        "ocr_bruto":
            ocr_bruto,

        "ocr_limpo":
            ocr_limpo,

        "confianca":
            confianca,

        "diferenca":
            round(
                diferenca,
                2,
            ),

        "diferencas_caracteres":
            diferencas,
    }


# ============================================================
# AVALIA UMA IMAGEM
# ============================================================

def avaliar_imagem(
    passagem,
    imagem,
    yolo_result,
    ocr_results,
):

    if not yolo_result:

        return {
            "aprovado":
                False,

            "diferenca":
                999.0,

            "ocr_bruto":
                "",

            "ocr_limpo":
                "",

            "confianca":
                0.0,

            "diferencas_caracteres":
                "Sem detecção",
        }

    if not yolo_result.get(
        "detected",
        False,
    ):

        return {
            "aprovado":
                False,

            "diferenca":
                999.0,

            "ocr_bruto":
                "",

            "ocr_limpo":
                "",

            "confianca":
                0.0,

            "diferencas_caracteres":
                "Sem detecção",
        }

    detections = yolo_result.get(
        "detections",
        [],
    )

    if not detections:

        crop_b64 = yolo_result.get(
            "crop_b64"
        )

        if crop_b64:

            detections = [
                {
                    "detection_index":
                        1,
                }
            ]

    melhor = None

    for detection in detections:

        detection_index = (
            detection.get(
                "detection_index",
                1,
            )
        )

        detection_task_id = (
            f"{imagem['task_id']}_"
            f"d{detection_index}"
        )

        ocr_result = ocr_results.get(
            detection_task_id
        )

        if not ocr_result:
            continue

        # ----------------------------------------------------
        # TESTA TODAS AS PLACAS DO SISTEMA
        # ----------------------------------------------------

        for placa_sistema in (
            passagem[
                "placas_sistema"
            ]
        ):

            avaliacao = avaliar_ocr(
                placa_sistema,
                ocr_result,
            )

            avaliacao[
                "placa_sistema_comparada"
            ] = placa_sistema

            if (
                melhor is None
                or
                avaliacao[
                    "diferenca"
                ]
                <
                melhor[
                    "diferenca"
                ]
            ):

                melhor = avaliacao

            if avaliacao[
                "aprovado"
            ]:

                return avaliacao

    if melhor is not None:
        return melhor

    return {
        "aprovado":
            False,

        "diferenca":
            999.0,

        "ocr_bruto":
            "",

        "ocr_limpo":
            "",

        "confianca":
            0.0,

        "diferencas_caracteres":
            "Sem OCR",
    }


# ============================================================
# CRIA LINHA DE SAÍDA
# ============================================================

def criar_linha_saida(
    passagem,
    resultado_ocr,
    resultado_final,
):

    dados = passagem[
        "dados_saida"
    ]

    linha = {}

    # ========================================================
    # DADOS DA PASSAGEM
    # ========================================================

    linha[
        "ID_PASSAGEM"
    ] = dados[
        "id_passagem"
    ]

    linha[
        "Data/hora da passagem"
    ] = dados[
        "data_hora"
    ]

    linha[
        "Praça de pedágio"
    ] = dados[
        "praca"
    ]

    linha[
        "Rodovia"
    ] = dados[
        "rodovia"
    ]

    linha[
        "km"
    ] = dados[
        "km"
    ]

    linha[
        "Sentido"
    ] = dados[
        "sentido"
    ]

    linha[
        "Faixa"
    ] = dados[
        "faixa"
    ]

    linha[
        "Placa no sistema"
    ] = dados[
        "placa_sistema"
    ]

    linha[
        "Categoria"
    ] = dados[
        "categoria"
    ]

    linha[
        "Marca"
    ] = dados[
        "marca"
    ]

    linha[
        "Modelo"
    ] = dados[
        "modelo"
    ]

    linha[
        "Cor"
    ] = dados[
        "cor"
    ]

    linha[
        "Valor R$"
    ] = dados[
        "valor"
    ]

    # ========================================================
    # OCR
    # ========================================================

    linha[
        "OCR Bruto"
    ] = resultado_ocr.get(
        "ocr_bruto",
        "",
    )

    linha[
        "OCR Limpo"
    ] = resultado_ocr.get(
        "ocr_limpo",
        "",
    )

    linha[
        "OCR Confidence"
    ] = resultado_ocr.get(
        "confianca",
        0.0,
    )

    # ========================================================
    # COMPARAÇÃO
    # ========================================================

    linha[
        "Diferença OCR"
    ] = resultado_ocr.get(
        "diferenca",
        999.0,
    )

    linha[
        "Diferenças de caracteres"
    ] = resultado_ocr.get(
        "diferencas_caracteres",
        "",
    )

    linha[
        "Resultado"
    ] = resultado_final

    return linha


# ============================================================
# PROCESSA EXCEL
# ============================================================

def processar_excel(
    caminho_excel,
    processo_yolo,
    processo_ocr,
):

    inicio = time.perf_counter()

    temp_dir = Path(
        tempfile.mkdtemp(
            prefix="validacao_passagens_"
        )
    )

    try:

        passagens = carregar_planilha(
            caminho_excel,
            temp_dir,
        )

        total_ids = len(
            passagens
        )

        total_imagens = sum(
            len(
                passagem[
                    "imagens"
                ]
            )
            for passagem in passagens
        )

        print(
            f"\n[PASSAGENS] {total_ids}"
        )

        print(
            f"[IMAGENS] {total_imagens}"
        )

        if not passagens:

            raise ValueError(
                "Nenhuma passagem encontrada."
            )

        # ====================================================
        # STATUS
        # ====================================================

        ids_aprovados = set()
        ids_revisar = set()

        proxima_imagem = {
            passagem["id"]: 0
            for passagem in passagens
        }

        # ====================================================
        # MELHOR RESULTADO POR ID
        # ====================================================

        melhor_resultado_por_id = {}

        # ====================================================
        # BARRA
        # ====================================================

        barra = tqdm(
            total=total_imagens,
            desc="YOLO + OCR",
            unit="img",
            dynamic_ncols=True,
            leave=True,
        )

        total_enviado_yolo = 0
        total_detectadas = 0
        total_ocr = 0

        try:

            while True:

                lote_rodada = []

                # =================================================
                # UMA IMAGEM POR ID PENDENTE
                # =================================================

                for passagem in passagens:

                    id_valor = passagem[
                        "id"
                    ]

                    if (
                        id_valor
                        in ids_aprovados
                    ):

                        continue

                    if (
                        id_valor
                        in ids_revisar
                    ):

                        continue

                    indice = proxima_imagem[
                        id_valor
                    ]

                    imagens = passagem[
                        "imagens"
                    ]

                    if (
                        indice
                        >= len(imagens)
                    ):

                        ids_revisar.add(
                            id_valor
                        )

                        continue

                    imagem = imagens[
                        indice
                    ]

                    proxima_imagem[
                        id_valor
                    ] = indice + 1

                    lote_rodada.append(
                        (
                            passagem,
                            imagem,
                        )
                    )

                if not lote_rodada:
                    break

                # =================================================
                # PROCESSA EM SUBLOTES
                # =================================================

                for inicio_lote in range(
                    0,
                    len(lote_rodada),
                    BATCH_SIZE,
                ):

                    lote = lote_rodada[
                        inicio_lote:
                        inicio_lote
                        + BATCH_SIZE
                    ]

                    imagens = [
                        item[1]
                        for item in lote
                    ]

                    # --------------------------------------------
                    # YOLO
                    # --------------------------------------------

                    yolo_results = (
                        executar_yolo(
                            processo_yolo,
                            imagens,
                        )
                    )

                    total_enviado_yolo += len(
                        imagens
                    )

                    total_detectadas += sum(
                        1
                        for resultado
                        in yolo_results.values()
                        if resultado.get(
                            "detected",
                            False,
                        )
                    )

                    # --------------------------------------------
                    # OCR
                    # --------------------------------------------

                    ocr_results = (
                        executar_ocr(
                            processo_ocr,
                            yolo_results,
                        )
                    )

                    total_ocr += len(
                        ocr_results
                    )

                    # --------------------------------------------
                    # AVALIA
                    # --------------------------------------------

                    for passagem, imagem in lote:

                        id_valor = passagem[
                            "id"
                        ]

                        if (
                            id_valor
                            in ids_aprovados
                        ):

                            barra.update(
                                1
                            )

                            continue

                        yolo_result = (
                            yolo_results.get(
                                imagem[
                                    "task_id"
                                ],
                                {},
                            )
                        )

                        avaliacao = (
                            avaliar_imagem(
                                passagem,
                                imagem,
                                yolo_result,
                                ocr_results,
                            )
                        )

                        # ----------------------------------------
                        # GUARDA MELHOR OCR
                        # ----------------------------------------

                        distancia = avaliacao[
                            "diferenca"
                        ]

                        anterior = (
                            melhor_resultado_por_id.get(
                                id_valor
                            )
                        )

                        if (
                            anterior is None
                            or
                            distancia
                            <
                            anterior[
                                "diferenca"
                            ]
                        ):

                            melhor_resultado_por_id[
                                id_valor
                            ] = avaliacao.copy()

                        # ----------------------------------------
                        # APROVAÇÃO
                        # ----------------------------------------

                        if avaliacao[
                            "aprovado"
                        ]:

                            ids_aprovados.add(
                                id_valor
                            )

                            # Guarda resultado que aprovou.
                            melhor_resultado_por_id[
                                id_valor
                            ] = avaliacao.copy()

                            # ------------------------------------
                            # NÃO PROCESSA IMAGENS FUTURAS
                            # ------------------------------------

                            restante = (
                                len(
                                    passagem[
                                        "imagens"
                                    ]
                                )
                                -
                                proxima_imagem[
                                    id_valor
                                ]
                            )

                            if restante > 0:

                                barra.update(
                                    restante
                                )

                                proxima_imagem[
                                    id_valor
                                ] = len(
                                    passagem[
                                        "imagens"
                                    ]
                                )

                        barra.update(
                            1
                        )

                # =================================================
                # IDS ESGOTADOS
                # =================================================

                for passagem in passagens:

                    id_valor = passagem[
                        "id"
                    ]

                    if (
                        id_valor
                        in ids_aprovados
                    ):

                        continue

                    if (
                        id_valor
                        in ids_revisar
                    ):

                        continue

                    if (
                        proxima_imagem[
                            id_valor
                        ]
                        >=
                        len(
                            passagem[
                                "imagens"
                            ]
                        )
                    ):

                        ids_revisar.add(
                            id_valor
                        )

                # =================================================
                # POSTFIX
                # =================================================

                resolvidos = (
                    len(
                        ids_aprovados
                    )
                    +
                    len(
                        ids_revisar
                    )
                )

                pendentes = max(
                    0,
                    total_ids - resolvidos,
                )

                barra.set_postfix(
                    aprovados=len(
                        ids_aprovados
                    ),
                    revisar=len(
                        ids_revisar
                    ),
                    pendentes=pendentes,
                )

        finally:

            barra.close()

        # ====================================================
        # GARANTE REVISAR NOS RESTANTES
        # ====================================================

        for passagem in passagens:

            id_valor = passagem[
                "id"
            ]

            if (
                id_valor
                not in ids_aprovados
            ):

                ids_revisar.add(
                    id_valor
                )

        # ====================================================
        # GERA RESULTADOS
        # ====================================================

        registros = []

        for passagem in passagens:

            id_valor = passagem[
                "id"
            ]

            melhor = (
                melhor_resultado_por_id.get(
                    id_valor
                )
            )

            if melhor is None:

                melhor = {
                    "ocr_bruto":
                        "",

                    "ocr_limpo":
                        "",

                    "confianca":
                        0.0,

                    "diferenca":
                        999.0,

                    "diferencas_caracteres":
                        "Nenhuma leitura OCR",
                }

            if (
                id_valor
                in ids_aprovados
            ):

                resultado_final = (
                    "Aprovado"
                )

            else:

                resultado_final = (
                    "Revisar"
                )

            registros.append(
                criar_linha_saida(
                    passagem,
                    melhor,
                    resultado_final,
                )
            )

        # ====================================================
        # DATAFRAME
        # ====================================================

        df_resultado = pd.DataFrame(
            registros,
            columns=CAMPOS_SAIDA,
        )

        # ====================================================
        # NÃO SALVA AQUI
        # ====================================================
        # O DataFrame é devolvido ao chamador.
        # O Excel final será criado somente depois que
        # todos os arquivos forem processados.

        # ====================================================
        # ESTATÍSTICAS
        # ====================================================

        aprovados = len(
            ids_aprovados
        )

        revisar = len(
            ids_revisar
        )

        tempo_total = (
            time.perf_counter()
            - inicio
        )

        print(
            "\n"
            + "=" * 75
        )

        print(
            "RESULTADO FINAL"
        )

        print(
            "=" * 75
        )

        print(
            f"IDs únicos: "
            f"{total_ids}"
        )

        print(
            f"Aprovados: "
            f"{aprovados}"
        )

        print(
            f"Revisar: "
            f"{revisar}"
        )

        print(
            f"Imagens existentes: "
            f"{total_imagens}"
        )

        print(
            f"Imagens enviadas ao YOLO: "
            f"{total_enviado_yolo}"
        )

        print(
            f"Detecções YOLO: "
            f"{total_detectadas}"
        )

        print(
            f"OCR executados: "
            f"{total_ocr}"
        )

        print(
            f"Tempo total: "
            f"{tempo_total:.1f}s"
        )

        print(
            "=" * 75
        )

        return df_resultado

    finally:

        shutil.rmtree(
            temp_dir,
            ignore_errors=True,
        )


# ============================================================
# SALVA RESULTADO FINAL
# ============================================================

def salvar_resultado_final(
    df_resultado,
    caminho_saida,
):
    """
    Cria um único Excel final contendo os resultados
    acumulados de todos os Excels processados.
    """

    saida = Path(
        caminho_saida
    ).resolve()

    saida.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook_saida = openpyxl.Workbook()

    sheet_resultado = workbook_saida.active
    sheet_resultado.title = "Resultado"

    # ========================================================
    # CABEÇALHO
    # ========================================================

    for coluna, nome in enumerate(
        CAMPOS_SAIDA,
        start=1,
    ):

        celula = sheet_resultado.cell(
            row=1,
            column=coluna,
            value=nome,
        )

        celula.font = Font(
            bold=True
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # ========================================================
    # DADOS
    # ========================================================

    for linha_excel, (_, registro) in enumerate(
        df_resultado.iterrows(),
        start=2,
    ):

        for coluna, nome in enumerate(
            CAMPOS_SAIDA,
            start=1,
        ):

            valor = registro[nome]

            if pd.isna(valor):
                valor = ""

            celula = sheet_resultado.cell(
                row=linha_excel,
                column=coluna,
                value=valor,
            )

            celula.alignment = Alignment(
                vertical="top",
                wrap_text=(
                    nome
                    ==
                    "Diferenças de caracteres"
                ),
            )

    # ========================================================
    # CONGELA CABEÇALHO
    # ========================================================

    sheet_resultado.freeze_panes = "A2"

    # ========================================================
    # FILTRO
    # ========================================================

    sheet_resultado.auto_filter.ref = (
        sheet_resultado.dimensions
    )

    # ========================================================
    # ALTURA CABEÇALHO
    # ========================================================

    sheet_resultado.row_dimensions[1].height = 35

    # ========================================================
    # LARGURA DAS COLUNAS
    # ========================================================

    for coluna in range(
        1,
        sheet_resultado.max_column + 1,
    ):

        nome_coluna = sheet_resultado.cell(
            row=1,
            column=coluna,
        ).value

        maior = len(
            str(nome_coluna or "")
        )

        for linha in range(
            2,
            sheet_resultado.max_row + 1,
        ):

            valor = sheet_resultado.cell(
                row=linha,
                column=coluna,
            ).value

            if valor is not None:

                maior = max(
                    maior,
                    len(str(valor)),
                )

        if nome_coluna == "Diferenças de caracteres":

            largura = 55

        elif nome_coluna in (
            "OCR Bruto",
            "OCR Limpo",
            "Placa no sistema",
        ):

            largura = 20

        elif nome_coluna == "Data/hora da passagem":

            largura = 22

        else:

            largura = min(
                max(maior + 2, 12),
                30,
            )

        sheet_resultado.column_dimensions[
            get_column_letter(coluna)
        ].width = largura

    # ========================================================
    # FORMATA NÚMEROS
    # ========================================================

    mapa_colunas = {
        sheet_resultado.cell(
            row=1,
            column=coluna,
        ).value: coluna
        for coluna in range(
            1,
            sheet_resultado.max_column + 1,
        )
    }

    coluna_conf = mapa_colunas.get(
        "OCR Confidence"
    )

    if coluna_conf:

        for linha in range(
            2,
            sheet_resultado.max_row + 1,
        ):

            sheet_resultado.cell(
                row=linha,
                column=coluna_conf,
            ).number_format = "0.0000"

    coluna_diferenca = mapa_colunas.get(
        "Diferença OCR"
    )

    if coluna_diferenca:

        for linha in range(
            2,
            sheet_resultado.max_row + 1,
        ):

            sheet_resultado.cell(
                row=linha,
                column=coluna_diferenca,
            ).number_format = "0.00"

    coluna_resultado = mapa_colunas.get(
        "Resultado"
    )

    # ========================================================
    # DESTAQUE DO RESULTADO
    # ========================================================

    if coluna_resultado:

        for linha in range(
            2,
            sheet_resultado.max_row + 1,
        ):

            celula = sheet_resultado.cell(
                row=linha,
                column=coluna_resultado,
            )

            if celula.value == "Aprovado":

                celula.fill = PatternFill(
                    fill_type="solid",
                    fgColor="C6EFCE",
                )

                celula.font = Font(
                    bold=True,
                    color="006100",
                )

            else:

                celula.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFC7CE",
                )

                celula.font = Font(
                    bold=True,
                    color="9C0006",
                )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # ========================================================
    # SEPARAÇÃO VISUAL DOS BLOCOS
    # ========================================================

    from openpyxl.styles import Border, Side

    ultima_coluna_dados = len(
        CAMPOS_DADOS_PASSAGEM
    )

    ultima_coluna_ocr = (
        ultima_coluna_dados
        + len(CAMPOS_OCR)
    )

    linha_divisoria = Side(
        style="medium",
        color="808080",
    )

    for linha in range(
        1,
        sheet_resultado.max_row + 1,
    ):

        for coluna in (
            ultima_coluna_dados,
            ultima_coluna_ocr,
        ):

            celula = sheet_resultado.cell(
                row=linha,
                column=coluna,
            )

            celula.border = Border(
                right=linha_divisoria,
                top=celula.border.top,
                bottom=celula.border.bottom,
                left=celula.border.left,
            )

    workbook_saida.save(
        saida
    )

    print(
        f"\n[SAÍDA] Excel final criado:\n{saida}"
    )

    return saida


# ============================================================
# MAIN
# ============================================================

def main():

    print(
        "\n"
        + "=" * 75
    )

    print(
        "VALIDADOR DE PASSAGENS — ISENÇÃO JUDICIAL"
    )

    print(
        "=" * 75
    )

    print(
        f"[CONFIG] Aprovação: até "
        f"{MAX_DIFERENCAS_APROVACAO} diferenças"
    )

    # ========================================================
    # ENTRADA
    # ========================================================

    caminho_entrada = input(
        "\nCaminho do Excel ou pasta com os Excels: "
    ).strip().strip('"')

    if not caminho_entrada:

        print(
            "[ERRO] Nenhum caminho informado."
        )

        sys.exit(1)

    caminho_entrada = Path(
        caminho_entrada
    ).resolve()

    # ========================================================
    # DESCOBRE OS ARQUIVOS
    # ========================================================

    if caminho_entrada.is_file():

        if caminho_entrada.suffix.lower() != ".xlsx":

            print(
                "[ERRO] O arquivo informado não é .xlsx."
            )

            sys.exit(1)

        arquivos_excel = [
            caminho_entrada
        ]

        pasta_base = caminho_entrada.parent

    elif caminho_entrada.is_dir():

        arquivos_excel = []

        for arquivo in sorted(
            caminho_entrada.rglob("*.xlsx")
        ):

            if arquivo.name.startswith("~$"):
                continue

            if arquivo.name == "resultado_isencao_judicial.xlsx":
                continue

            if arquivo.stem.endswith(
                (
                    "_analisado",
                    "_validado",
                    "_resultado",
                )
            ):

                continue

            arquivos_excel.append(
                arquivo
            )

        pasta_base = caminho_entrada

    else:

        print(
            f"[ERRO] Caminho não encontrado:\n"
            f"{caminho_entrada}"
        )

        sys.exit(1)

    if not arquivos_excel:

        print(
            f"[ERRO] Nenhum arquivo .xlsx encontrado em:\n"
            f"{caminho_entrada}"
        )

        sys.exit(1)

    print(
        f"\n[ARQUIVOS] "
        f"{len(arquivos_excel)} Excel(s) serão processados."
    )

    for indice, arquivo in enumerate(
        arquivos_excel,
        start=1,
    ):

        print(
            f"  [{indice}/{len(arquivos_excel)}] "
            f"{arquivo}"
        )

    # ========================================================
    # DEPENDÊNCIAS
    # ========================================================

    arquivos_obrigatorios = [
        PYTHON_YOLO,
        PYTHON_OCR,
        DETECTOR_PY,
        OCR_PY,
    ]

    print(
        "\n[CHECK] Verificando arquivos..."
    )

    for caminho in arquivos_obrigatorios:

        caminho = Path(caminho)

        if not caminho.exists():

            print(
                f"[ERRO] Arquivo não encontrado:\n"
                f"{caminho}"
            )

            sys.exit(1)

        print(
            f"[OK] {caminho}"
        )

    # ========================================================
    # SAÍDA ÚNICA
    # ========================================================

    caminho_saida = caminho_resultado(
        "cnl_isencao",
        "resultado_isencao_judicial.xlsx",
    )

    # ========================================================
    # PROCESSOS
    # ========================================================

    processo_yolo = iniciar_yolo()
    processo_ocr = iniciar_ocr()

    todos_resultados = []
    erros_arquivos = []

    inicio_geral = time.perf_counter()

    try:

        print(
            "\n[INIT] Aguardando modelos..."
        )

        time.sleep(15)

        total = len(
            arquivos_excel
        )

        for indice, caminho_excel in enumerate(
            arquivos_excel,
            start=1,
        ):

            print(
                "\n"
                + "#" * 75
            )

            print(
                f"  ARQUIVO {indice}/{total}: "
                f"{caminho_excel.name}"
            )

            print(
                "#" * 75
            )

            try:

                with contextlib.redirect_stdout(io.StringIO()):
                    df_resultado = processar_excel(
                        caminho_excel,
                        processo_yolo,
                        processo_ocr,
                    )

                if (
                    df_resultado is not None
                    and not df_resultado.empty
                ):
                    origem = (
                        str(caminho_excel.relative_to(caminho_entrada))
                        if caminho_entrada.is_dir()
                        else caminho_excel.name
                    )
                    df_resultado.insert(0, "Arquivo de Origem", origem)

                    todos_resultados.append(
                        df_resultado
                    )

            except Exception as erro:

                erros_arquivos.append(
                    (str(caminho_excel), f"{type(erro).__name__}: {erro}")
                )

        # ====================================================
        # CONSOLIDA TUDO
        # ====================================================

        if not todos_resultados:

            print(
                "\n[ERRO] Nenhum resultado foi gerado."
            )

            for caminho, mensagem in erros_arquivos:
                print(f"[ERRO] {caminho}: {mensagem}")

            sys.exit(1)

        df_final = pd.concat(
            todos_resultados,
            ignore_index=True,
        )

        # ====================================================
        # SALVA APENAS AGORA
        # ====================================================

        salvar_resultado_final(
            df_final,
            caminho_saida,
        )

        aprovados = int(
            (
                df_final["Resultado"]
                == "Aprovado"
            ).sum()
        )

        revisar = int(
            (
                df_final["Resultado"]
                == "Revisar"
            ).sum()
        )

        tempo_geral = (
            time.perf_counter()
            - inicio_geral
        )

        print(
            "\n"
            + "=" * 75
        )

        print(
            "RESULTADO GERAL"
        )

        print(
            "=" * 75
        )

        print(
            f"Excels processados: "
            f"{len(arquivos_excel)}"
        )

        print(
            f"Registros finais: "
            f"{len(df_final)}"
        )

        print(
            f"Aprovados: "
            f"{aprovados}"
        )

        print(
            f"Revisar: "
            f"{revisar}"
        )

        print(
            f"Tempo total: "
            f"{tempo_geral:.1f}s"
        )

        print(
            f"Saída única: "
            f"{caminho_saida.resolve()}"
        )

        if erros_arquivos:
            print("Arquivos com erro:")
            for caminho, mensagem in erros_arquivos:
                print(f"- {caminho}: {mensagem}")

        print(
            "=" * 75
        )

    except KeyboardInterrupt:

        print(
            "\n[INFO] Processamento interrompido."
        )

    except Exception as erro:

        print(
            "\n[ERRO]"
        )

        print(
            f"{type(erro).__name__}: {erro}"
        )

        raise

    finally:

        encerrar_processo(
            processo_yolo,
            "YOLO",
        )

        encerrar_processo(
            processo_ocr,
            "OCR",
        )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()
