from __future__ import annotations

import functools
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from pathlib import Path
from typing import Any

try:
    import fitz  # PyMuPDF
except ImportError as exc:
    raise SystemExit(
        "Dependência ausente: PyMuPDF. Instale no ambiente principal com: python -m pip install pymupdf"
    ) from exc

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from tqdm import tqdm


# ============================================================================
# CONFIGURAÇÕES LOCAIS DO VALIDADOR
# Este arquivo não importa config.py nem módulos do validador Excel.
# ============================================================================

BASE_DIR = Path(__file__).resolve().parent
YOLO_VENV = BASE_DIR / ".venv_yolo"
OCR_VENV = BASE_DIR / ".venv_ocr"
YOLO_MODEL = BASE_DIR / "models" / "bestnew1024.onnx"

BATCH_SIZE = 64
MAX_DIFERENCAS_APROVACAO = 2
MAX_CARACTERES_OCR = 100
MAX_IMAGENS_POR_PASSAGEM = 50
TAMANHO_PLACA = 7
YOLO_CONF_THRESHOLD = 0.35
YOLO_MARGEM_X = 0.15
YOLO_MARGEM_Y = 0.20


# ============================================================================
# LAYOUT DO PDF (modelo recebido: página A4 paisagem, 842 x 595 pt)
# ============================================================================

PDF_PAGE_REFERENCE_WIDTH = 842.0
PDF_PAGE_REFERENCE_HEIGHT = 595.0
PDF_DATA_Y_MIN = 108.5
PDF_DATA_Y_MAX = 125.5
PDF_MIN_IMAGE_WIDTH = 250
PDF_MIN_IMAGE_HEIGHT = 150

PDF_FIELD_BANDS = {
    "id_arrecadacao": (34.5, 126.8, True),
    "data": (126.8, 157.5, False),
    "hora": (157.5, 180.8, False),
    "sp": (180.8, 196.5, True),
    "km_m": (196.5, 219.0, True),
    "praca": (219.0, 279.8, False),
    "numero_pista": (279.8, 318.8, False),
    # Faixas históricas usadas como fallback. A extração normal identifica
    # as divisórias reais da linha de dados em cada PDF, pois a largura das
    # colunas pode variar entre documentos do mesmo modelo.
    "sentido": (318.8, 339.8, False),
    "municipio": (339.8, 368.2, False),
    "codigo": (368.2, 388.5, True),
    "placa_sistema": (388.5, 412.5, True),
    "marca_modelo": (412.5, 453.0, False),
    "quantidade_eixos": (453.0, 498.0, False),
    "valor_tarifa": (498.0, 533.2, False),
    "responsavel_anotacao": (533.2, 574.5, False),
    "numero_auto": (574.5, 651.0, False),
    "envio_email_der": (651.0, 699.8, False),
    "confirmacao_email_der": (699.8, 807.8, False),
}

OUTPUT_COLUMNS = [
    "Arquivo PDF", "Página", "ID do Sist. de Arrecadação", "Data", "Hora",
    "SP", "KM + M", "Nome da Praça de Pedágio", "Número da Pista", "Sentido",
    "Município", "Código", "Placa no Sistema", "Marca/Modelo", "Quantidade de Eixos",
    "Valor da Tarifa", "Nome (Resp. p/ Anotação)", "Número Auto de Infração/Evento",
    "Envio de E-mail ao DER", "Confirmação de Recebimento do E-mail pelo DER",
    "Placa Comparada", "OCR Bruto", "OCR Limpo", "OCR Confidence", "Diferença OCR",
    "Diferenças de Caracteres", "Resultado",
]

OUTPUT_FIELD_MAP = {
    "ID do Sist. de Arrecadação": "id_arrecadacao",
    "Data": "data",
    "Hora": "hora",
    "SP": "sp",
    "KM + M": "km_m",
    "Nome da Praça de Pedágio": "praca",
    "Número da Pista": "numero_pista",
    "Sentido": "sentido",
    "Município": "municipio",
    "Código": "codigo",
    "Placa no Sistema": "placa_sistema",
    "Marca/Modelo": "marca_modelo",
    "Quantidade de Eixos": "quantidade_eixos",
    "Valor da Tarifa": "valor_tarifa",
    "Nome (Resp. p/ Anotação)": "responsavel_anotacao",
    "Número Auto de Infração/Evento": "numero_auto",
    "Envio de E-mail ao DER": "envio_email_der",
    "Confirmação de Recebimento do E-mail pelo DER": "confirmacao_email_der",
}


# ============================================================================
# NORMALIZAÇÃO E COMPARAÇÃO DE PLACAS (implementação própria neste arquivo)
# ============================================================================

_CONFUNDIVEIS: set[tuple[str, str]] = {
    ("0", "O"), ("O", "0"), ("1", "I"), ("I", "1"), ("1", "L"), ("L", "1"),
    ("T", "I"), ("I", "T"), ("2", "Z"), ("Z", "2"), ("Z", "7"), ("7", "Z"),
    ("4", "A"), ("A", "4"), ("5", "S"), ("S", "5"), ("6", "G"), ("G", "6"),
    ("8", "B"), ("B", "8"), ("B", "E"), ("E", "B"), ("D", "0"), ("0", "D"),
    ("D", "O"), ("O", "D"),
}

_MAPA_LETRAS = {"0": "O", "1": "I", "2": "Z", "4": "A", "5": "S", "6": "G", "7": "Z", "8": "B"}
_MAPA_DIGITOS = {"O": "0", "Q": "0", "D": "0", "I": "1", "L": "1", "T": "7", "J": "1", "Z": "2", "A": "4", "S": "5", "G": "6", "B": "8", "E": "8", "U": "0"}


@functools.lru_cache(maxsize=16384)
def _levenshtein(a: str, b: str) -> float:
    anterior = list(range(len(b) + 1))
    for i, ca in enumerate(a, start=1):
        atual = [i]
        for j, cb in enumerate(b, start=1):
            if ca == cb:
                custo = 0.0
            elif (ca, cb) in _CONFUNDIVEIS:
                custo = 0.35
            else:
                custo = 1.0
            atual.append(min(anterior[j] + 1, atual[j - 1] + 1, anterior[j - 1] + custo))
        anterior = atual
    return anterior[-1]


def limpar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(valor).upper())


def distancia_placas(placa_a: str, placa_b: str) -> float:
    a = limpar_texto(placa_a)
    b = limpar_texto(placa_b)
    if not a or not b:
        return 999.0
    return _levenshtein(a, b)


def descobrir_mascara(placa: str) -> str:
    placa = limpar_texto(placa)
    if len(placa) != TAMANHO_PLACA:
        return "LLLDLDD"
    if placa[4].isalpha():
        return "LLLDLDD"
    return "LLLDDDD"


def aplicar_mascara_placa(candidato: str, mascara: str) -> str:
    candidato = limpar_texto(candidato)
    if len(candidato) != TAMANHO_PLACA:
        return candidato
    resultado = []
    for caractere, tipo in zip(candidato, mascara):
        if tipo == "L":
            resultado.append(_MAPA_LETRAS.get(caractere, caractere))
        else:
            resultado.append(_MAPA_DIGITOS.get(caractere, caractere))
    return "".join(resultado)


def encontrar_melhor_trecho(placa_esperada: str, placa_ocr: str) -> tuple[str, float]:
    esperada = limpar_texto(placa_esperada)
    ocr = limpar_texto(placa_ocr)
    if not esperada or not ocr:
        return "", 999.0
    if len(esperada) != TAMANHO_PLACA:
        return ocr, distancia_placas(esperada, ocr)

    mascara = descobrir_mascara(esperada)
    if len(ocr) == TAMANHO_PLACA:
        corrigido = aplicar_mascara_placa(ocr, mascara)
        distancia_original = distancia_placas(esperada, ocr)
        distancia_corrigida = distancia_placas(esperada, corrigido)
        if distancia_corrigida <= distancia_original:
            return corrigido, distancia_corrigida
        return ocr, distancia_original

    if len(ocr) < TAMANHO_PLACA:
        return ocr, distancia_placas(esperada, ocr)

    candidatos = []
    for inicio in range(len(ocr) - TAMANHO_PLACA + 1):
        trecho = ocr[inicio:inicio + TAMANHO_PLACA]
        corrigido = aplicar_mascara_placa(trecho, mascara)
        distancia_original = distancia_placas(esperada, trecho)
        distancia_corrigida = distancia_placas(esperada, corrigido)
        if trecho == esperada:
            prioridade = 0
        elif corrigido == esperada:
            prioridade = 1
        else:
            prioridade = 2
        distancia = min(distancia_original, distancia_corrigida)
        candidato_final = corrigido if distancia_corrigida <= distancia_original else trecho
        candidatos.append((prioridade, distancia, inicio, candidato_final))
    candidatos.sort(key=lambda item: (item[0], item[1], item[2]))
    _, melhor_distancia, _, melhor_trecho = candidatos[0]
    return melhor_trecho, melhor_distancia


def descrever_diferencas(placa_esperada: str, placa_ocr: str) -> str:
    esperada = limpar_texto(placa_esperada)
    ocr = limpar_texto(placa_ocr)
    if not esperada or not ocr:
        return "Sem comparação"

    linhas, colunas = len(esperada) + 1, len(ocr) + 1
    matriz = [[0.0 for _ in range(colunas)] for _ in range(linhas)]
    operacoes = [[None for _ in range(colunas)] for _ in range(linhas)]
    for i in range(linhas):
        matriz[i][0] = float(i)
    for j in range(colunas):
        matriz[0][j] = float(j)

    for i in range(1, linhas):
        for j in range(1, colunas):
            ca, cb = esperada[i - 1], ocr[j - 1]
            if ca == cb:
                custo_substituicao = 0.0
            elif (ca, cb) in _CONFUNDIVEIS:
                custo_substituicao = 0.35
            else:
                custo_substituicao = 1.0
            substituicao = matriz[i - 1][j - 1] + custo_substituicao
            insercao = matriz[i][j - 1] + 1.0
            remocao = matriz[i - 1][j] + 1.0
            menor = min(substituicao, insercao, remocao)
            matriz[i][j] = menor
            if menor == substituicao:
                operacoes[i][j] = ("igual", ca, cb) if custo_substituicao == 0 else ("substituicao", ca, cb)
            elif menor == insercao:
                operacoes[i][j] = ("insercao", "", cb)
            else:
                operacoes[i][j] = ("remocao", ca, "")

    i, j = len(esperada), len(ocr)
    diferencas = []
    while i > 0 or j > 0:
        if i > 0 and j > 0:
            operacao = operacoes[i][j]
            tipo = operacao[0]
        elif i > 0:
            tipo = "remocao"
        else:
            tipo = "insercao"

        if tipo == "igual":
            i -= 1
            j -= 1
        elif tipo == "substituicao":
            diferencas.append((i, ocr[j - 1], esperada[i - 1]))
            i -= 1
            j -= 1
        elif tipo == "insercao":
            diferencas.append((j, ocr[j - 1], "—"))
            j -= 1
        elif tipo == "remocao":
            diferencas.append((i, "—", esperada[i - 1]))
            i -= 1

    diferencas.reverse()
    if not diferencas:
        return "Nenhuma"
    resultado = []
    for posicao, recebido, esperado in diferencas:
        if esperado == "—":
            resultado.append(f"Posição {posicao}: OCR={recebido} (caractere extra)")
        elif recebido == "—":
            resultado.append(f"Posição {posicao}: OCR ausente (esperado={esperado})")
        else:
            resultado.append(f"Posição {posicao}: OCR={recebido} ≠ Sistema={esperado}")
    return "; ".join(resultado)


def avaliar_ocr(placa_planilha: str, resultado_ocr: dict[str, Any]) -> dict[str, Any]:
    placa_planilha = limpar_texto(placa_planilha)
    ocr_bruto = resultado_ocr.get("ocr_best_original", "")
    ocr_best = str(resultado_ocr.get("ocr_best", "") or "")[:MAX_CARACTERES_OCR]
    try:
        confianca = float(resultado_ocr.get("confidence_best", 0.0) or 0.0)
    except (TypeError, ValueError):
        confianca = 0.0

    if not limpar_texto(ocr_best):
        return {
            "aprovado": False, "ocr_bruto": ocr_bruto, "ocr_limpo": "",
            "confianca": confianca, "diferenca": 999.0,
            "diferencas_caracteres": "Sem OCR",
        }

    ocr_normalizado = limpar_texto(ocr_best)
    ocr_limpo, diferenca = encontrar_melhor_trecho(placa_planilha, ocr_normalizado)
    diferencas = descrever_diferencas(placa_planilha, ocr_limpo)
    aprovado = diferenca <= MAX_DIFERENCAS_APROVACAO
    return {
        "aprovado": aprovado, "ocr_bruto": ocr_bruto, "ocr_limpo": ocr_limpo,
        "confianca": confianca, "diferenca": round(diferenca, 2),
        "diferencas_caracteres": diferencas,
    }


def avaliar_imagem(
    registro: dict[str, Any],
    imagem: dict[str, Any],
    yolo_result: dict[str, Any],
    ocr_results: dict[str, Any],
) -> dict[str, Any]:
    if not yolo_result or not yolo_result.get("detected", False):
        return {
            "aprovado": False, "diferenca": 999.0, "ocr_bruto": "", "ocr_limpo": "",
            "confianca": 0.0, "diferencas_caracteres": "Sem detecção",
        }

    detections = yolo_result.get("detections", [])
    if not detections and yolo_result.get("crop_b64"):
        detections = [{"detection_index": 1}]

    melhor = None
    for detection in detections:
        detection_index = detection.get("detection_index", 1)
        detection_task_id = f"{imagem['task_id']}_d{detection_index}"
        ocr_result = ocr_results.get(detection_task_id)
        if not ocr_result:
            continue
        for placa_sistema in registro["placas_sistema"]:
            avaliacao = avaliar_ocr(placa_sistema, ocr_result)
            avaliacao["placa_sistema_comparada"] = placa_sistema
            if melhor is None or avaliacao["diferenca"] < melhor["diferenca"]:
                melhor = avaliacao
            if avaliacao["aprovado"]:
                return avaliacao

    if melhor is not None:
        return melhor
    return {
        "aprovado": False, "diferenca": 999.0, "ocr_bruto": "", "ocr_limpo": "",
        "confianca": 0.0, "diferencas_caracteres": "Sem OCR",
    }


# ============================================================================
# EXTRAÇÃO DO PDF: texto por coordenadas e imagens incorporadas
# ============================================================================

def limpar_id(valor: Any) -> str:
    if valor is None:
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(valor).upper())


def corrigir_texto_pdf(texto: str, compactar: bool = False) -> str:
    texto = re.sub(r"\s+", " ", str(texto or "")).strip()
    texto = re.sub(r"\bS\s+o\b", "São", texto, flags=re.IGNORECASE)
    substituicoes = {
        r"\bPraca\b": "Praça", r"\bautuacao\b": "autuação",
        r"\bInfracao\b": "Infração", r"\barrecadacao\b": "arrecadação",
    }
    for padrao, substituto in substituicoes.items():
        texto = re.sub(padrao, substituto, texto, flags=re.IGNORECASE)
    if compactar:
        texto = re.sub(r"\s+", "", texto)
    return texto


def extrair_campo_da_pagina(
    page: fitz.Page, x_inicial: float, x_final: float,
    y_min: float = PDF_DATA_Y_MIN, y_max: float = PDF_DATA_Y_MAX,
    compactar: bool = False,
) -> str:
    sx = page.rect.width / PDF_PAGE_REFERENCE_WIDTH
    sy = page.rect.height / PDF_PAGE_REFERENCE_HEIGHT
    x0, x1 = x_inicial * sx, x_final * sx
    top, bottom = y_min * sy, y_max * sy
    palavras = []
    for palavra in page.get_text("words"):
        wx0, wy0, wx1, wy1, texto = palavra[:5]
        centro_x = (wx0 + wx1) / 2
        if x0 <= centro_x < x1 and wy1 >= top and wy0 <= bottom:
            palavras.append((wy0, wx0, str(texto)))
    palavras.sort(key=lambda item: (round(item[0], 1), item[1]))
    return corrigir_texto_pdf(" ".join(item[2] for item in palavras), compactar=compactar)


def detectar_faixas_colunas(page: fitz.Page) -> dict[str, tuple[float, float, bool]]:
    """Detecta as divisórias verticais da linha de dados do PDF.

    Os PDFs desse relatório têm 18 campos na linha principal, mas a largura
    de algumas colunas varia entre documentos. As divisórias vetoriais do
    próprio PDF são mais confiáveis do que usar coordenadas fixas para todos.
    Retorna as faixas em coordenadas da página de referência (842 x 595 pt).
    """
    escala_x = page.rect.width / PDF_PAGE_REFERENCE_WIDTH
    escala_y = page.rect.height / PDF_PAGE_REFERENCE_HEIGHT
    topo_ref = PDF_DATA_Y_MIN
    base_ref = PDF_DATA_Y_MAX
    candidatos = []

    for desenho in page.get_drawings():
        retangulo = desenho.get("rect")
        if retangulo is None:
            continue

        largura = retangulo.width
        altura = retangulo.height
        topo = retangulo.y0 / escala_y
        base = retangulo.y1 / escala_y

        # Divisórias das células da linha de dados (não a borda das fotos).
        if (
            largura <= 1.5 * escala_x
            and altura >= 13.0 * escala_y
            and abs(topo - topo_ref) <= 4.0
            and abs(base - base_ref) <= 4.0
        ):
            centro_x_ref = ((retangulo.x0 + retangulo.x1) / 2.0) / escala_x
            candidatos.append(centro_x_ref)

    # Vários desenhos representam a mesma linha; agrupa as duplicatas.
    candidatos.sort()
    grupos: list[list[float]] = []
    for x in candidatos:
        if not grupos or x - grupos[-1][-1] > 1.2:
            grupos.append([x])
        else:
            grupos[-1].append(x)

    limites = [sum(grupo) / len(grupo) for grupo in grupos]
    nomes_campos = list(PDF_FIELD_BANDS.keys())

    # Uma linha com 18 campos precisa de 19 limites, incluindo as bordas.
    if len(limites) != len(nomes_campos) + 1:
        return PDF_FIELD_BANDS

    faixas = {}
    for indice, nome in enumerate(nomes_campos):
        _, _, compactar = PDF_FIELD_BANDS[nome]
        faixas[nome] = (limites[indice], limites[indice + 1], compactar)
    return faixas


def extrair_campos_da_pagina(page: fitz.Page) -> dict[str, str]:
    campos = {}
    faixas = detectar_faixas_colunas(page)
    for nome, (x0, x1, compactar) in faixas.items():
        campos[nome] = extrair_campo_da_pagina(page, x0, x1, compactar=compactar)
    campos["id_arrecadacao"] = re.sub(r"\s+", "", campos["id_arrecadacao"])
    campos["placa_sistema"] = re.sub(r"[^A-Z0-9]", "", campos["placa_sistema"].upper())
    campos["sp"] = campos["sp"].upper()
    return campos


def extrair_imagens_da_pagina(
    documento: fitz.Document, page: fitz.Page, pasta_temporaria: Path, prefixo: str,
) -> list[dict[str, str]]:
    colocacoes: list[tuple[float, float, int]] = []
    vistos: set[tuple[int, int, int, int, int]] = set()
    for imagem in page.get_images(full=True):
        xref, largura, altura = int(imagem[0]), int(imagem[2]), int(imagem[3])
        if largura < PDF_MIN_IMAGE_WIDTH or altura < PDF_MIN_IMAGE_HEIGHT:
            continue
        for retangulo in page.get_image_rects(xref):
            chave = (xref, round(retangulo.x0), round(retangulo.y0), round(retangulo.x1), round(retangulo.y1))
            if chave in vistos:
                continue
            vistos.add(chave)
            colocacoes.append((retangulo.y0, retangulo.x0, xref))

    colocacoes.sort(key=lambda item: (round(item[0], 2), item[1]))
    imagens_saida = []
    for indice, (_, _, xref) in enumerate(colocacoes[:MAX_IMAGENS_POR_PASSAGEM], start=1):
        try:
            extraida = documento.extract_image(xref)
            bytes_imagem = extraida.get("image")
            extensao = str(extraida.get("ext") or "jpg").lower()
            if not bytes_imagem:
                continue
            if extensao not in {"jpg", "jpeg", "png", "webp", "bmp", "tif", "tiff"}:
                extensao = "jpg"
            caminho = pasta_temporaria / f"{prefixo}_img_{indice:02d}.{extensao}"
            caminho.write_bytes(bytes_imagem)
            imagens_saida.append({
                "task_id": f"{prefixo}_img_{indice:02d}",
                "image_path": str(caminho.resolve()),
                "indice": str(indice),
            })
        except Exception as erro:
            print(f"[AVISO] Não foi possível extrair imagem da página: {erro}")
    return imagens_saida


def extrair_registros_pdf(caminho_pdf: Path, pasta_temporaria: Path, indice_pdf: int) -> list[dict[str, Any]]:
    registros = []
    try:
        documento = fitz.open(caminho_pdf)
    except Exception as erro:
        raise RuntimeError(f"Não foi possível abrir o PDF: {erro}") from erro

    with documento:
        if documento.is_encrypted:
            raise RuntimeError("O PDF está protegido por senha.")
        for numero_pagina, page in enumerate(documento, start=1):
            campos = extrair_campos_da_pagina(page)
            id_original = campos.get("id_arrecadacao", "").strip()
            id_normalizado = limpar_id(id_original)
            chave_base = id_normalizado or f"PDF{indice_pdf}PAG{numero_pagina}"
            prefixo = f"pdf{indice_pdf:04d}_p{numero_pagina:04d}_{chave_base[:45]}"
            imagens = extrair_imagens_da_pagina(documento, page, pasta_temporaria, prefixo)
            placa = campos.get("placa_sistema", "")
            if not id_original:
                print(f"[AVISO] ID não encontrado em {caminho_pdf.name}, página {numero_pagina}; registro irá para revisão.")
            if not placa:
                print(f"[AVISO] Placa não encontrada em {caminho_pdf.name}, página {numero_pagina}; registro irá para revisão.")
            if not imagens:
                print(f"[AVISO] Nenhuma foto encontrada em {caminho_pdf.name}, página {numero_pagina}; registro irá para revisão.")
            registros.append({
                "id": prefixo, "id_arrecadacao": id_original, "id_normalizado": id_normalizado,
                "arquivo_pdf": caminho_pdf.name, "caminho_pdf": str(caminho_pdf.resolve()),
                "pagina": numero_pagina, "campos": campos,
                "placas_sistema": [placa] if placa else [], "imagens": imagens,
            })
    return registros


def localizar_pdfs(caminho_entrada: Path) -> list[Path]:
    if caminho_entrada.is_file():
        if caminho_entrada.suffix.lower() != ".pdf":
            raise ValueError("O arquivo informado não é um PDF.")
        return [caminho_entrada]
    if not caminho_entrada.is_dir():
        raise FileNotFoundError(f"Caminho não encontrado: {caminho_entrada}")
    arquivos = []
    for arquivo in caminho_entrada.rglob("*"):
        if not arquivo.is_file() or arquivo.suffix.lower() != ".pdf":
            continue
        if arquivo.name.startswith("~$") or arquivo.name.startswith("."):
            continue
        if arquivo.stem.lower().startswith("resultado_"):
            continue
        arquivos.append(arquivo)
    return sorted(arquivos)


# ============================================================================
# PROCESSOS: JSON para os workers incluídos neste próprio arquivo
# ============================================================================

def localizar_python_venv(pasta_venv: Path) -> Path:
    candidatos = [pasta_venv / "Scripts" / "python.exe", pasta_venv / "bin" / "python"]
    for candidato in candidatos:
        if candidato.exists():
            return candidato
    raise FileNotFoundError(
        f"Não encontrei o Python do ambiente virtual em {pasta_venv}. "
        "Confira se .venv_yolo e .venv_ocr estão na pasta do projeto."
    )


def iniciar_worker(python_exe: Path, script_worker: Path, argumentos: list[str], nome: str, pasta_logs: Path):
    log_path = pasta_logs / f"{nome.lower()}_worker.log"
    log_handle = log_path.open("w", encoding="utf-8", errors="replace")
    processo = subprocess.Popen(
        [str(python_exe), str(script_worker), *argumentos],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=log_handle,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )
    WORKER_LOG_HANDLES[processo.pid] = log_handle
    WORKER_LOG_PATHS[processo.pid] = log_path
    return processo


WORKER_LOG_HANDLES: dict[int, Any] = {}
WORKER_LOG_PATHS: dict[int, Path] = {}


def enviar_json(processo: subprocess.Popen, dados: dict[str, Any]) -> None:
    if processo.poll() is not None:
        raise RuntimeError(f"Processo já encerrou. Código: {processo.returncode}")
    if processo.stdin is None:
        raise RuntimeError("A entrada do processo não está disponível.")
    processo.stdin.write(json.dumps(dados, ensure_ascii=False) + "\n")
    processo.stdin.flush()


def ler_json(processo: subprocess.Popen, nome: str) -> dict[str, Any]:
    if processo.stdout is None:
        raise RuntimeError(f"Saída do processo {nome} não disponível.")
    linha = processo.stdout.readline()
    if not linha:
        log_path = WORKER_LOG_PATHS.get(processo.pid)
        detalhe = ""
        if log_path and log_path.exists():
            try:
                detalhe = "\nLog do worker:\n" + "".join(log_path.read_text(encoding="utf-8", errors="replace").splitlines(True)[-40:])
            except OSError:
                pass
        raise RuntimeError(f"Processo {nome} encerrou antes de responder. Código: {processo.poll()}.{detalhe}")
    try:
        retorno = json.loads(linha)
    except json.JSONDecodeError as erro:
        raise RuntimeError(f"Resposta inválida do processo {nome}: {linha[:500]}") from erro
    if retorno.get("error"):
        raise RuntimeError(f"Erro do processo {nome}: {retorno['error']}")
    return retorno


def encerrar_processo(processo: subprocess.Popen | None, nome: str) -> None:
    if processo is None:
        return
    if processo.poll() is None:
        try:
            enviar_json(processo, {"command": "exit"})
            processo.wait(timeout=10)
        except Exception:
            try:
                processo.terminate()
                processo.wait(timeout=5)
            except Exception:
                try:
                    processo.kill()
                except Exception:
                    pass
    log_handle = WORKER_LOG_HANDLES.pop(processo.pid, None)
    if log_handle is not None:
        try:
            log_handle.close()
        except Exception:
            pass
    WORKER_LOG_PATHS.pop(processo.pid, None)


def executar_yolo(processo_yolo: subprocess.Popen, imagens: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    if not imagens:
        return {}
    items = []
    for imagem in imagens:
        caminho = Path(imagem["image_path"]).resolve()
        if not caminho.exists() or caminho.stat().st_size == 0:
            continue
        items.append({"task_id": imagem["task_id"], "image_path": str(caminho)})
    if not items:
        return {}
    enviar_json(processo_yolo, {"command": "detect_batch_all", "images": items})
    resposta = ler_json(processo_yolo, "YOLO")
    return {
        resultado["task_id"]: resultado
        for resultado in resposta.get("results", [])
        if resultado.get("task_id")
    }


def executar_ocr(processo_ocr: subprocess.Popen, yolo_results: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    if not yolo_results:
        return {}
    crops = []
    for task_id, resultado in yolo_results.items():
        if not resultado.get("detected", False):
            continue
        detections = resultado.get("detections", [])
        if not detections and resultado.get("crop_b64"):
            detections = [{"detection_index": 1, "crop_b64": resultado["crop_b64"]}]
        for detection in detections:
            detection_index = detection.get("detection_index", 1)
            crop_b64 = detection.get("crop_b64")
            if crop_b64:
                crops.append({"task_id": f"{task_id}_d{detection_index}", "crop_b64": crop_b64})
    if not crops:
        return {}
    enviar_json(processo_ocr, {"command": "ocr_batch", "images": crops})
    resposta = ler_json(processo_ocr, "OCR")
    return {
        resultado["task_id"]: resultado
        for resultado in resposta.get("results", [])
        if resultado.get("task_id")
    }


# ============================================================================
# PROCESSAMENTO E PLANILHA DE SAÍDA
# ============================================================================

def processar_registros_com_yolo_ocr(
    registros: list[dict[str, Any]], processo_yolo: subprocess.Popen, processo_ocr: subprocess.Popen,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    aprovados: set[str] = set()
    revisar: set[str] = set()
    proxima_imagem = {registro["id"]: 0 for registro in registros}
    melhor_por_registro: dict[str, dict[str, Any]] = {}
    total_imagens = sum(len(r["imagens"]) for r in registros)
    total_enviadas_yolo = total_detectadas = total_resultados_ocr = 0
    barra = tqdm(total=total_imagens, desc="YOLO + OCR (PDF)", unit="img", dynamic_ncols=True)

    try:
        while True:
            rodada = []
            for registro in registros:
                chave = registro["id"]
                if chave in aprovados or chave in revisar:
                    continue
                indice = proxima_imagem[chave]
                imagens = registro["imagens"]
                if indice >= len(imagens):
                    revisar.add(chave)
                    continue
                imagem = imagens[indice]
                proxima_imagem[chave] = indice + 1
                rodada.append((registro, imagem))

            if not rodada:
                break

            for inicio_lote in range(0, len(rodada), BATCH_SIZE):
                lote = rodada[inicio_lote:inicio_lote + BATCH_SIZE]
                imagens_lote = [item[1] for item in lote]
                yolo_resultados = executar_yolo(processo_yolo, imagens_lote)
                total_enviadas_yolo += len(imagens_lote)
                total_detectadas += sum(1 for r in yolo_resultados.values() if r.get("detected", False))
                ocr_resultados = executar_ocr(processo_ocr, yolo_resultados)
                total_resultados_ocr += len(ocr_resultados)

                for registro, imagem in lote:
                    chave = registro["id"]
                    if chave in aprovados:
                        barra.update(1)
                        continue
                    avaliacao = avaliar_imagem(
                        registro, imagem, yolo_resultados.get(imagem["task_id"], {}), ocr_resultados,
                    )
                    distancia = avaliacao.get("diferenca", 999.0)
                    anterior = melhor_por_registro.get(chave)
                    if anterior is None or distancia < anterior.get("diferenca", 999.0):
                        melhor_por_registro[chave] = avaliacao.copy()
                    if avaliacao.get("aprovado", False):
                        aprovados.add(chave)
                        melhor_por_registro[chave] = avaliacao.copy()
                        restante = len(registro["imagens"]) - proxima_imagem[chave]
                        if restante > 0:
                            barra.update(restante)
                            proxima_imagem[chave] = len(registro["imagens"])
                    barra.update(1)

            resolvidos = len(aprovados) + len(revisar)
            barra.set_postfix(aprovados=len(aprovados), revisar=len(revisar), pendentes=max(0, len(registros) - resolvidos))
    finally:
        barra.close()

    for registro in registros:
        if registro["id"] not in aprovados:
            revisar.add(registro["id"])

    resultado_final = []
    for registro in registros:
        chave = registro["id"]
        avaliacao = melhor_por_registro.get(chave, {})
        dados = registro["campos"]
        row: dict[str, Any] = {"Arquivo PDF": registro["arquivo_pdf"], "Página": registro["pagina"]}
        for coluna_saida, campo_pdf in OUTPUT_FIELD_MAP.items():
            row[coluna_saida] = dados.get(campo_pdf, "")
        row["Placa Comparada"] = avaliacao.get("placa_sistema_comparada", "")
        row["OCR Bruto"] = avaliacao.get("ocr_bruto", "")
        row["OCR Limpo"] = avaliacao.get("ocr_limpo", "")
        row["OCR Confidence"] = avaliacao.get("confianca", 0.0)
        row["Diferença OCR"] = avaliacao.get("diferenca", 999.0)
        row["Diferenças de Caracteres"] = avaliacao.get(
            "diferencas_caracteres", "Nenhuma leitura OCR" if not avaliacao else "Sem comparação",
        )
        row["Resultado"] = "Aprovado" if chave in aprovados else "Revisar"
        resultado_final.append(row)

    estatisticas = {
        "registros": len(registros), "aprovados": len(aprovados), "revisar": len(revisar),
        "imagens": total_imagens, "imagens_enviadas_yolo": total_enviadas_yolo,
        "deteccoes_yolo": total_detectadas, "resultados_ocr": total_resultados_ocr,
    }
    return resultado_final, estatisticas


def salvar_resultado_excel(registros: list[dict[str, Any]], caminho_saida: Path) -> Path:
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resultado"
    for indice_coluna, nome in enumerate(OUTPUT_COLUMNS, start=1):
        celula = sheet.cell(row=1, column=indice_coluna, value=nome)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for indice_linha, registro in enumerate(registros, start=2):
        for indice_coluna, nome in enumerate(OUTPUT_COLUMNS, start=1):
            valor = registro.get(nome, "")
            if isinstance(valor, float) and math.isnan(valor):
                valor = ""
            celula = sheet.cell(row=indice_linha, column=indice_coluna, value=valor)
            celula.alignment = Alignment(vertical="top", wrap_text=(nome == "Diferenças de Caracteres"))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 42
    larguras = {
        "Arquivo PDF": 34, "Página": 9, "ID do Sist. de Arrecadação": 38, "Data": 13, "Hora": 12,
        "SP": 12, "KM + M": 13, "Nome da Praça de Pedágio": 26, "Número da Pista": 15,
        "Sentido": 14, "Município": 22, "Código": 12, "Placa no Sistema": 18, "Marca/Modelo": 22,
        "Quantidade de Eixos": 18, "Valor da Tarifa": 15, "Nome (Resp. p/ Anotação)": 22,
        "Número Auto de Infração/Evento": 28, "Envio de E-mail ao DER": 22,
        "Confirmação de Recebimento do E-mail pelo DER": 36, "Placa Comparada": 18, "OCR Bruto": 20,
        "OCR Limpo": 18, "OCR Confidence": 16, "Diferença OCR": 15, "Diferenças de Caracteres": 55,
        "Resultado": 14,
    }
    coluna_por_nome = {nome: idx for idx, nome in enumerate(OUTPUT_COLUMNS, start=1)}
    for nome, largura in larguras.items():
        sheet.column_dimensions[get_column_letter(coluna_por_nome[nome])].width = largura

    for linha in range(2, sheet.max_row + 1):
        sheet.cell(row=linha, column=coluna_por_nome["OCR Confidence"]).number_format = "0.0000"
        sheet.cell(row=linha, column=coluna_por_nome["Diferença OCR"]).number_format = "0.00"
        resultado_celula = sheet.cell(row=linha, column=coluna_por_nome["Resultado"])
        if resultado_celula.value == "Aprovado":
            resultado_celula.fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
            resultado_celula.font = Font(bold=True, color="006100")
        else:
            resultado_celula.fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
            resultado_celula.font = Font(bold=True, color="9C0006")
        resultado_celula.alignment = Alignment(horizontal="center", vertical="center")

    lado = Side(style="medium", color="808080")
    for linha in range(1, sheet.max_row + 1):
        for numero_coluna in (20, 24, 26):
            celula = sheet.cell(row=linha, column=numero_coluna)
            celula.border = Border(left=celula.border.left, right=lado, top=celula.border.top, bottom=celula.border.bottom)
    workbook.save(caminho_saida)
    return caminho_saida


# ============================================================================
# MAIN
# ============================================================================

def main() -> None:
    print("\n" + "=" * 76)
    print("VALIDADOR DE EVASÃO DE PEDÁGIO — PDF + YOLO + OCR (INDEPENDENTE)")
    print("=" * 76)
    print("O código de extração, comparação, comunicação e saída está neste arquivo.")
    print("Não serão importados config.py, validador_excel.py, detector_process.py ou ocr_process.py.")
    print(f"[CONFIG] Modelo YOLO: {YOLO_MODEL}")
    print(f"[CONFIG] Limite de aprovação (distância ponderada): {MAX_DIFERENCAS_APROVACAO}")

    entrada = input("\nCaminho do PDF ou pasta com PDFs: ").strip().strip('"')
    if not entrada:
        print("[ERRO] Nenhum caminho informado.")
        sys.exit(1)
    caminho_entrada = Path(entrada).expanduser().resolve()
    try:
        arquivos_pdf = localizar_pdfs(caminho_entrada)
    except Exception as erro:
        print(f"[ERRO] {erro}")
        sys.exit(1)
    if not arquivos_pdf:
        print(f"[ERRO] Nenhum PDF encontrado em: {caminho_entrada}")
        sys.exit(1)

    python_yolo = localizar_python_venv(YOLO_VENV)
    python_ocr = localizar_python_venv(OCR_VENV)
    arquivos_obrigatorios = [(python_yolo, "Python do ambiente YOLO"), (python_ocr, "Python do ambiente OCR"), (YOLO_MODEL, "Modelo ONNX do YOLO")]
    for caminho, descricao in arquivos_obrigatorios:
        if not caminho.exists():
            raise FileNotFoundError(f"{descricao} não encontrado: {caminho}")

    pasta_saida = caminho_entrada.parent if caminho_entrada.is_file() else caminho_entrada
    caminho_saida = pasta_saida / "resultado_validacao_evasao_pedagio.xlsx"
    pasta_temporaria = Path(tempfile.mkdtemp(prefix="validacao_evasao_pdf_"))
    pasta_workers = pasta_temporaria / "workers"
    pasta_workers.mkdir(parents=True, exist_ok=True)

    # O conteúdo dos workers é gravado do texto embutido mais abaixo neste arquivo.
    script_yolo = pasta_workers / "yolo_worker_embutido.py"
    script_ocr = pasta_workers / "ocr_worker_embutido.py"
    script_yolo.write_text(YOLO_WORKER_SOURCE, encoding="utf-8")
    script_ocr.write_text(OCR_WORKER_SOURCE, encoding="utf-8")

    processo_yolo = processo_ocr = None
    inicio_geral = time.perf_counter()
    try:
        print(f"\n[ARQUIVOS] {len(arquivos_pdf)} PDF(s)")
        todos_registros: list[dict[str, Any]] = []
        erros_pdf: list[tuple[str, str]] = []
        for indice_pdf, caminho_pdf in enumerate(arquivos_pdf, start=1):
            print(f"[PDF {indice_pdf}/{len(arquivos_pdf)}] {caminho_pdf}")
            try:
                registros = extrair_registros_pdf(caminho_pdf, pasta_temporaria, indice_pdf)
                todos_registros.extend(registros)
                print(f"  [OK] {len(registros)} página(s)/registro(s); {sum(len(r['imagens']) for r in registros)} foto(s) extraída(s).")
            except Exception as erro:
                erros_pdf.append((str(caminho_pdf), f"{type(erro).__name__}: {erro}"))
                print(f"  [ERRO] {type(erro).__name__}: {erro}")

        if not todos_registros:
            raise RuntimeError("Nenhum registro foi extraído dos PDFs informados.")
        print(f"\n[EXTRAÇÃO] Registros: {len(todos_registros)}")
        print(f"[EXTRAÇÃO] Fotos: {sum(len(r['imagens']) for r in todos_registros)}")
        print("[INIT] Iniciando workers YOLO e PaddleOCR incluídos no validador...")
        processo_yolo = iniciar_worker(python_yolo, script_yolo, [str(YOLO_MODEL.resolve())], "YOLO", pasta_temporaria)
        processo_ocr = iniciar_worker(python_ocr, script_ocr, [], "OCR", pasta_temporaria)
        print("[INIT] Aguardando os modelos carregarem (15 segundos)...")
        time.sleep(15)
        for nome, processo in (("YOLO", processo_yolo), ("OCR", processo_ocr)):
            if processo.poll() is not None:
                log_path = WORKER_LOG_PATHS.get(processo.pid)
                detalhe = log_path.read_text(encoding="utf-8", errors="replace") if log_path and log_path.exists() else ""
                raise RuntimeError(f"O processo {nome} encerrou na inicialização (código {processo.returncode}).\n{detalhe}")

        linhas, estatisticas = processar_registros_com_yolo_ocr(todos_registros, processo_yolo, processo_ocr)
        salvar_resultado_excel(linhas, caminho_saida)
        tempo_total = time.perf_counter() - inicio_geral
        print("\n" + "=" * 76)
        print("RESULTADO GERAL")
        print("=" * 76)
        print(f"PDFs encontrados: {len(arquivos_pdf)}")
        print(f"Registros/páginas: {estatisticas['registros']}")
        print(f"Aprovados: {estatisticas['aprovados']}")
        print(f"Revisar: {estatisticas['revisar']}")
        print(f"Fotos extraídas: {estatisticas['imagens']}")
        print(f"Imagens enviadas ao YOLO: {estatisticas['imagens_enviadas_yolo']}")
        print(f"Detecções YOLO: {estatisticas['deteccoes_yolo']}")
        print(f"Resultados OCR: {estatisticas['resultados_ocr']}")
        print(f"Tempo total: {tempo_total:.1f}s")
        print(f"[SAÍDA] {caminho_saida.resolve()}")
        if erros_pdf:
            print(f"[AVISO] {len(erros_pdf)} PDF(s) não puderam ser processados:")
            for arquivo, erro in erros_pdf:
                print(f"  - {arquivo}: {erro}")
        print("=" * 76)
    except KeyboardInterrupt:
        print("\n[INFO] Processamento interrompido pelo usuário.")
    finally:
        encerrar_processo(processo_yolo, "YOLO")
        encerrar_processo(processo_ocr, "OCR")
        shutil.rmtree(pasta_temporaria, ignore_errors=True)


# As duas implementações abaixo foram incorporadas ao arquivo a partir dos
# scripts que o usuário forneceu, com a leitura da configuração externa removida.
OCR_WORKER_SOURCE = r'''import sys
import json
import base64

import numpy as np
import cv2
from paddleocr import PaddleOCR


# ============================================================
# CONFIGURAÇÕES
# ============================================================

LARGURA_MINIMA = 600
MAX_SIDE_LIMIT = 4000

OCR_CONF_MINIMA = 0.70


# ============================================================
# MAPAS OCR AUXILIAR
# ============================================================

NUMERO_PARA_LETRA = {
    "0": "O",
    "1": "I",
    "2": "Z",
    "3": "B",
    "4": "A",
    "5": "S",
    "6": "G",
    "7": "T",
    "8": "B",
    "9": "G",
}


LETRA_PARA_NUMERO = {
    "O": "0",
    "Q": "0",
    "D": "0",
    "I": "1",
    "L": "1",
    "Z": "2",
    "B": "8",
    "A": "4",
    "S": "5",
    "G": "6",
    "T": "7",
}


# ============================================================
# LIMPA TEXTO
# ============================================================

def limpar_texto_ocr(texto):

    if texto is None:
        return ""

    texto = str(
        texto
    ).upper()

    return "".join(
        char
        for char in texto
        if (
            "A" <= char <= "Z"
            or "0" <= char <= "9"
        )
    )


# ============================================================
# PREPARA TEXTO PARA AUXILIAR
# ============================================================

def preparar_para_auxiliar(texto):

    texto = limpar_texto_ocr(
        texto
    )

    if len(texto) == 8:
        return texto[1:]

    return texto


# ============================================================
# OCR AUXILIAR
# ============================================================

def OCR_auxiliar(texto):

    if not texto:
        return ""

    texto = limpar_texto_ocr(
        texto
    )

    if len(texto) != 7:
        return texto

    chars = list(texto)

    # --------------------------------------------------------
    # POSIÇÕES 1-3
    # DEVEM SER LETRAS
    # --------------------------------------------------------

    for i in range(0, 3):

        char = chars[i]

        if char.isdigit():

            chars[i] = (
                NUMERO_PARA_LETRA.get(
                    char,
                    char
                )
            )

    # --------------------------------------------------------
    # POSIÇÃO 4
    # DEVE SER NÚMERO
    # --------------------------------------------------------

    char = chars[3]

    if char.isalpha():

        chars[3] = (
            LETRA_PARA_NUMERO.get(
                char,
                char
            )
        )

    # --------------------------------------------------------
    # POSIÇÃO 5
    # NÃO ALTERAR
    # --------------------------------------------------------

    # intencionalmente vazio

    # --------------------------------------------------------
    # POSIÇÃO 6
    # DEVE SER NÚMERO
    # --------------------------------------------------------

    char = chars[5]

    if char.isalpha():

        chars[5] = (
            LETRA_PARA_NUMERO.get(
                char,
                char
            )
        )

    # --------------------------------------------------------
    # POSIÇÃO 7
    # DEVE SER NÚMERO
    # --------------------------------------------------------

    char = chars[6]

    if char.isalpha():

        chars[6] = (
            LETRA_PARA_NUMERO.get(
                char,
                char
            )
        )

    return "".join(chars)


# ============================================================
# VALIDA PLACA
# ============================================================

def placa_valida(texto):

    if not texto:
        return False

    if len(texto) != 7:
        return False

    if not (
        texto[0].isalpha()
        and texto[1].isalpha()
        and texto[2].isalpha()
    ):
        return False

    if not texto[3].isdigit():
        return False

    if not texto[4].isalnum():
        return False

    if not texto[5].isdigit():
        return False

    if not texto[6].isdigit():
        return False

    return True


# ============================================================
# RESULTADO SATISFATÓRIO
# ============================================================

def resultado_satisfatorio(
    texto,
    confianca
):

    if not texto:
        return False

    if not placa_valida(texto):
        return False

    if confianca < OCR_CONF_MINIMA:
        return False

    return True


# ============================================================
# LIMITA TAMANHO
# ============================================================

def limitar_tamanho_imagem(
    image,
    max_side=MAX_SIDE_LIMIT
):

    if (
        image is None
        or image.size == 0
    ):
        return image

    h, w = image.shape[:2]

    maior_lado = max(
        h,
        w
    )

    if maior_lado <= max_side:
        return image

    escala = (
        max_side
        / float(maior_lado)
    )

    nova_w = max(
        1,
        int(
            round(
                w * escala
            )
        )
    )

    nova_h = max(
        1,
        int(
            round(
                h * escala
            )
        )
    )

    return cv2.resize(
        image,
        (
            nova_w,
            nova_h
        ),
        interpolation=cv2.INTER_AREA
    )


# ============================================================
# PREPARA VARIANTES
# ============================================================

def preparar_variantes(crop):

    if (
        crop is None
        or crop.size == 0
    ):

        return None, None

    altura, largura = (
        crop.shape[:2]
    )

    escala = max(
        3.0,
        LARGURA_MINIMA
        / max(
            largura,
            1
        )
    )

    maior_dimensao = max(
        altura * escala,
        largura * escala
    )

    if (
        maior_dimensao
        > MAX_SIDE_LIMIT
    ):

        escala = (
            MAX_SIDE_LIMIT
            / float(
                max(
                    altura,
                    largura
                )
            )
        )

    crop_cubic = cv2.resize(
        crop,
        None,
        fx=escala,
        fy=escala,
        interpolation=cv2.INTER_CUBIC
    )

    crop_cubic = limitar_tamanho_imagem(
        crop_cubic
    )

    # ========================================================
    # CLAHE
    # ========================================================

    gray = cv2.cvtColor(
        crop_cubic,
        cv2.COLOR_BGR2GRAY
    )

    clahe = cv2.createCLAHE(
        clipLimit=1.2,
        tileGridSize=(8, 8)
    )

    gray_clahe = clahe.apply(
        gray
    )

    crop_clahe = cv2.cvtColor(
        gray_clahe,
        cv2.COLOR_GRAY2BGR
    )

    crop_clahe = limitar_tamanho_imagem(
        crop_clahe
    )

    return (
        crop_cubic,
        crop_clahe
    )


# ============================================================
# PARSE PADDLEOCR
# ============================================================

def parse_ocr_result(
    resultado
):

    textos = []
    confiancas = []

    if resultado is None:
        return "", 0.0

    try:

        dados = resultado.json.get(
            "res",
            {}
        )

        for texto in dados.get(
            "rec_texts",
            []
        ):

            if texto:
                textos.append(
                    str(texto)
                )

        for score in dados.get(
            "rec_scores",
            []
        ):

            try:

                confiancas.append(
                    float(score)
                )

            except Exception:
                pass

    except Exception:

        return "", 0.0

    texto_bruto = "".join(
        textos
    )

    texto_limpo = (
        limpar_texto_ocr(
            texto_bruto
        )
    )

    confianca = (
        float(
            np.mean(
                confiancas
            )
        )
        if confiancas
        else 0.0
    )

    return (
        texto_limpo,
        confianca
    )


# ============================================================
# EXECUTA OCR
#
# Sem tqdm.
# Sem prints por imagem.
# ============================================================

def executar_ocr_variante(
    ocr,
    imagens,
    indices
):

    resultados = {}

    if not indices:
        return resultados

    for start in range(
        0,
        len(indices),
        8
    ):

        sub_indices = indices[
            start:start + 8
        ]

        sub_imagens = [
            limitar_tamanho_imagem(
                imagens[i]
            )
            for i in sub_indices
        ]

        try:

            outputs = ocr.predict(
                sub_imagens
            )

            outputs = list(
                outputs
            )

            for indice, output in zip(
                sub_indices,
                outputs
            ):

                resultados[
                    indice
                ] = parse_ocr_result(
                    output
                )

        except Exception:

            # ------------------------------------------------
            # FALLBACK INDIVIDUAL
            # ------------------------------------------------

            for indice, imagem in zip(
                sub_indices,
                sub_imagens
            ):

                try:

                    outputs = list(
                        ocr.predict(
                            limitar_tamanho_imagem(
                                imagem
                            )
                        )
                    )

                    output = (
                        outputs[0]
                        if outputs
                        else None
                    )

                    resultados[
                        indice
                    ] = parse_ocr_result(
                        output
                    )

                except Exception:

                    resultados[
                        indice
                    ] = (
                        "",
                        0.0
                    )

    return resultados


# ============================================================
# PREPARA BATCH
# ============================================================

def preparar_batch(
    items
):

    imagens = []
    metadados = []
    erros = []

    for item in items:

        task_id = item.get(
            "task_id",
            "sem_id"
        )

        try:

            img_bytes = (
                base64.b64decode(
                    item[
                        "crop_b64"
                    ]
                )
            )

            array = np.frombuffer(
                img_bytes,
                np.uint8
            )

            crop = cv2.imdecode(
                array,
                cv2.IMREAD_COLOR
            )

            if (
                crop is None
                or crop.size == 0
            ):

                erros.append(
                    {
                        "task_id":
                            task_id,

                        "error":
                            "Crop inválido"
                    }
                )

                continue

            crop = (
                limitar_tamanho_imagem(
                    crop
                )
            )

            (
                crop_cubic,
                crop_clahe
            ) = preparar_variantes(
                crop
            )

            if (
                crop_cubic is None
                or crop_clahe is None
            ):

                erros.append(
                    {
                        "task_id":
                            task_id,

                        "error":
                            "Falha preparando variantes"
                    }
                )

                continue

            imagens.append(
                crop_cubic
            )

            metadados.append(
                {
                    "task_id":
                        task_id,

                    "variant":
                        "CUBIC"
                }
            )

            imagens.append(
                crop_clahe
            )

            metadados.append(
                {
                    "task_id":
                        task_id,

                    "variant":
                        "CUBIC + CLAHE"
                }
            )

        except Exception as e:

            erros.append(
                {
                    "task_id":
                        task_id,

                    "error":
                        str(e)
                }
            )

    return (
        imagens,
        metadados,
        erros
    )


# ============================================================
# PROCESSA BATCH
# ============================================================

def processar_batch(
    ocr,
    items
):

    (
        imagens,
        metadados,
        erros
    ) = preparar_batch(
        items
    )

    resultados = list(
        erros
    )

    if not imagens:
        return resultados

    # ========================================================
    # ÍNDICES
    # ========================================================

    indices_cubic = [
        i
        for i, meta
        in enumerate(
            metadados
        )
        if meta["variant"]
        == "CUBIC"
    ]

    indices_clahe = [
        i
        for i, meta
        in enumerate(
            metadados
        )
        if meta["variant"]
        == "CUBIC + CLAHE"
    ]

    # ========================================================
    # OCR CUBIC
    # ========================================================

    resultados_cubic = (
        executar_ocr_variante(
            ocr,
            imagens,
            indices_cubic
        )
    )

    # ========================================================
    # OCR CLAHE
    # ========================================================

    resultados_clahe = (
        executar_ocr_variante(
            ocr,
            imagens,
            indices_clahe
        )
    )

    # ========================================================
    # AGRUPA
    # ========================================================

    agrupado = {}

    for indice, resultado in (
        resultados_cubic.items()
    ):

        task_id = (
            metadados[indice][
                "task_id"
            ]
        )

        if task_id not in agrupado:
            agrupado[
                task_id
            ] = {}

        agrupado[
            task_id
        ]["CUBIC"] = resultado

    for indice, resultado in (
        resultados_clahe.items()
    ):

        task_id = (
            metadados[indice][
                "task_id"
            ]
        )

        if task_id not in agrupado:
            agrupado[
                task_id
            ] = {}

        agrupado[
            task_id
        ]["CLAHE"] = resultado

    # ========================================================
    # CONSOLIDA
    # ========================================================

    for task_id, variantes in (
        agrupado.items()
    ):

        cubic = variantes.get(
            "CUBIC",
            (
                "",
                0.0
            )
        )

        clahe = variantes.get(
            "CLAHE",
            (
                "",
                0.0
            )
        )

        (
            cubic_original,
            cubic_conf
        ) = cubic

        (
            clahe_original,
            clahe_conf
        ) = clahe

        # ====================================================
        # AUXILIAR
        # ====================================================

        cubic_sem_primeiro = (
            preparar_para_auxiliar(
                cubic_original
            )
        )

        clahe_sem_primeiro = (
            preparar_para_auxiliar(
                clahe_original
            )
        )

        cubic_auxiliar = (
            OCR_auxiliar(
                cubic_sem_primeiro
            )
        )

        clahe_auxiliar = (
            OCR_auxiliar(
                clahe_sem_primeiro
            )
        )

        # ====================================================
        # ESCOLHE MELHOR NORMAL
        # ====================================================

        if clahe_conf > cubic_conf:

            normal_original = (
                clahe_original
            )

            normal_sem_primeiro = (
                clahe_sem_primeiro
            )

            normal_aux = (
                clahe_auxiliar
            )

            normal_conf = (
                clahe_conf
            )

            normal_variant = (
                "CUBIC + CLAHE"
            )

        else:

            normal_original = (
                cubic_original
            )

            normal_sem_primeiro = (
                cubic_sem_primeiro
            )

            normal_aux = (
                cubic_auxiliar
            )

            normal_conf = (
                cubic_conf
            )

            normal_variant = (
                "CUBIC"
            )

        # ====================================================
        # RESULTADO FINAL
        #
        # SR DESATIVADO
        # ====================================================

        ocr_best_original = (
            normal_original
        )

        ocr_best_sem_primeiro = (
            normal_sem_primeiro
        )

        ocr_best_normal = (
            normal_aux
        )

        confidence_best_normal = (
            normal_conf
        )

        ocr_best = (
            ocr_best_normal
        )

        confidence_best = (
            confidence_best_normal
        )

        best_variant = (
            normal_variant
        )

        auxiliar_alterou = (
            ocr_best_original
            != ocr_best_normal
        )

        # ====================================================
        # RESULTADO
        # ====================================================

        resultados.append(
            {
                "task_id":
                    task_id,

                # --------------------------------------------
                # CUBIC
                # --------------------------------------------

                "ocr_cubic":
                    cubic_original,

                "confidence_cubic":
                    cubic_conf,

                "ocr_cubic_sem_primeiro":
                    cubic_sem_primeiro,

                "ocr_cubic_auxiliar":
                    cubic_auxiliar,

                # --------------------------------------------
                # CLAHE
                # --------------------------------------------

                "ocr_clahe":
                    clahe_original,

                "confidence_clahe":
                    clahe_conf,

                "ocr_clahe_sem_primeiro":
                    clahe_sem_primeiro,

                "ocr_clahe_auxiliar":
                    clahe_auxiliar,

                # --------------------------------------------
                # NORMAL
                # --------------------------------------------

                "ocr_best_original":
                    ocr_best_original,

                "ocr_best_sem_primeiro":
                    ocr_best_sem_primeiro,

                "ocr_best_normal":
                    ocr_best_normal,

                "confidence_best_normal":
                    confidence_best_normal,

                # --------------------------------------------
                # SR
                #
                # Mantido somente por compatibilidade.
                # Não executado.
                # --------------------------------------------

                "sr_used":
                    False,

                "sr_ocr":
                    "",

                "sr_ocr_sem_primeiro":
                    "",

                "sr_ocr_auxiliar":
                    "",

                "sr_confidence":
                    0.0,

                "sr_variant":
                    "",

                "resultado_gracas_sr":
                    False,

                # --------------------------------------------
                # FINAL
                # --------------------------------------------

                "ocr_best":
                    ocr_best,

                "confidence_best":
                    confidence_best,

                "best_variant":
                    best_variant,

                "auxiliar_alterou":
                    auxiliar_alterou,
            }
        )

    return resultados


# ============================================================
# MAIN
# ============================================================

def main():

    # ========================================================
    # INICIALIZA OCR
    # ========================================================

    try:

        ocr = PaddleOCR(
            lang="en",
            device="gpu:0",
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
        )

    except Exception as e:

        print(
            f"[OCR] Erro inicializando: {e}",
            file=sys.stderr,
            flush=True
        )

        sys.exit(1)

    # ========================================================
    # MENSAGENS INICIAIS
    #
    # Apenas uma vez.
    # ========================================================

    print(
        "[OCR] PaddleOCR pronto.",
        file=sys.stderr,
        flush=True
    )

    print(
        "[OCR] CUBIC + CLAHE | SR desativado.",
        file=sys.stderr,
        flush=True
    )

    # ========================================================
    # STDIN
    # ========================================================

    for line in sys.stdin:

        line = line.strip()

        if not line:
            continue

        try:

            req = json.loads(
                line
            )

            # =================================================
            # EXIT
            # =================================================

            if req.get(
                "command"
            ) == "exit":

                break

            # =================================================
            # BATCH
            # =================================================

            if req.get(
                "command"
            ) == "ocr_batch":

                items = req.get(
                    "images",
                    []
                )

                response = {
                    "results":
                        processar_batch(
                            ocr,
                            items
                        )
                }

            # =================================================
            # INDIVIDUAL
            # =================================================

            else:

                response = {
                    "results":
                        processar_batch(
                            ocr,
                            [req]
                        )
                }

            # =================================================
            # RETORNO JSON
            # =================================================

            sys.stdout.write(
                json.dumps(
                    response,
                    ensure_ascii=False
                )
                + "\n"
            )

            sys.stdout.flush()

        except Exception as e:

            print(
                f"[OCR] Erro: {e}",
                file=sys.stderr,
                flush=True
            )

            sys.stdout.write(
                json.dumps(
                    {
                        "results": [],
                        "error": str(e)
                    },
                    ensure_ascii=False
                )
                + "\n"
            )

            sys.stdout.flush()


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()

import sys
import json
import base64
from pathlib import Path

import numpy as np
import cv2'''
YOLO_WORKER_SOURCE = r'''import sys
import json
import base64
from pathlib import Path

import numpy as np
import cv2
import onnxruntime as ort


# ============================================================
# CONFIGURAÇÃO INCORPORADA AO WORKER
# ============================================================

YOLO_MODEL = Path(sys.argv[1]).expanduser().resolve() if len(sys.argv) > 1 else Path("models/bestnew1024.onnx").resolve()
YOLO_CONF_THRESHOLD = 0.35
YOLO_MARGEM_X = 0.15
YOLO_MARGEM_Y = 0.20


BATCH_SIZE = 64


# ============================================================
# LETTERBOX
# ============================================================

def letterbox(
    img,
    target_shape=(1024, 1024),
    color=(114, 114, 114)
):

    h, w = img.shape[:2]

    tw, th = target_shape

    r = min(
        tw / w,
        th / h
    )

    new_unpad = (
        int(round(w * r)),
        int(round(h * r))
    )

    dw = (
        tw - new_unpad[0]
    ) / 2.0

    dh = (
        th - new_unpad[1]
    ) / 2.0

    if (w, h) != new_unpad:

        resized = cv2.resize(
            img,
            new_unpad,
            interpolation=cv2.INTER_CUBIC
        )

    else:

        resized = img.copy()

    top = int(round(dh - 0.1))
    bottom = int(round(dh + 0.1))
    left = int(round(dw - 0.1))
    right = int(round(dw + 0.1))

    padded = cv2.copyMakeBorder(
        resized,
        top,
        bottom,
        left,
        right,
        cv2.BORDER_CONSTANT,
        value=color
    )

    return padded, r, dw, dh


# ============================================================
# TENSOR
# ============================================================

def imagem_para_tensor(image):

    input_img, r, dw, dh = letterbox(
        image,
        (1024, 1024)
    )

    blob = cv2.cvtColor(
        input_img,
        cv2.COLOR_BGR2RGB
    )

    blob = (
        blob.astype(np.float32)
        / 255.0
    )

    blob = np.transpose(
        blob,
        (2, 0, 1)
    )

    meta = {
        "r": r,
        "dw": dw,
        "dh": dh
    }

    return blob, meta


# ============================================================
# EXPANDE BOX
# ============================================================

def expandir_box(
    x1,
    y1,
    x2,
    y2,
    largura,
    altura
):

    bw = x2 - x1
    bh = y2 - y1

    extra_x = bw * YOLO_MARGEM_X
    extra_y = bh * YOLO_MARGEM_Y

    x1 = int(round(x1 - extra_x))
    y1 = int(round(y1 - extra_y))
    x2 = int(round(x2 + extra_x))
    y2 = int(round(y2 + extra_y))

    x1 = max(
        0,
        min(
            x1,
            largura - 1
        )
    )

    y1 = max(
        0,
        min(
            y1,
            altura - 1
        )
    )

    x2 = max(
        1,
        min(
            x2,
            largura
        )
    )

    y2 = max(
        1,
        min(
            y2,
            altura
        )
    )

    return x1, y1, x2, y2


# ============================================================
# PROCESSA SAÍDA YOLO
# ============================================================

def processar_saida(
    output,
    meta,
    original,
    image_path,
    task_id
):

    if output is None:

        return {
            "detected": False,
            "detections": []
        }

    try:

        if output.ndim == 3:
            preds = output[0]

        elif output.ndim == 2:
            preds = output

        else:

            return {
                "detected": False,
                "detections": []
            }

        if (
            preds.ndim != 2
            or preds.shape[1] < 6
        ):

            return {
                "detected": False,
                "detections": []
            }

        scores = preds[:, 4]

        mask = (
            scores
            >= YOLO_CONF_THRESHOLD
        )

        preds = preds[mask]
        scores = scores[mask]

        if len(preds) == 0:

            return {
                "detected": False,
                "detections": []
            }

        boxes = preds[:, :4].copy()

        h_orig, w_orig = (
            original.shape[:2]
        )

        candidates = []

        # ====================================================
        # CONVERTE DETECÇÕES
        # ====================================================

        for idx in range(len(boxes)):

            box = boxes[idx]

            confidence = float(
                scores[idx]
            )

            x1 = (
                box[0]
                - meta["dw"]
            ) / meta["r"]

            y1 = (
                box[1]
                - meta["dh"]
            ) / meta["r"]

            x2 = (
                box[2]
                - meta["dw"]
            ) / meta["r"]

            y2 = (
                box[3]
                - meta["dh"]
            ) / meta["r"]

            x1 = max(
                0,
                min(
                    int(round(x1)),
                    w_orig - 1
                )
            )

            y1 = max(
                0,
                min(
                    int(round(y1)),
                    h_orig - 1
                )
            )

            x2 = max(
                1,
                min(
                    int(round(x2)),
                    w_orig
                )
            )

            y2 = max(
                1,
                min(
                    int(round(y2)),
                    h_orig
                )
            )

            bw = x2 - x1
            bh = y2 - y1

            if (
                bw <= 0
                or bh <= 0
            ):
                continue

            aspect_ratio = (
                bw / bh
            )

            if (
                bw < 20
                or bh < 10
                or aspect_ratio < 0.8
                or aspect_ratio > 6.0
            ):
                continue

            candidates.append(
                {
                    "x1": x1,
                    "y1": y1,
                    "x2": x2,
                    "y2": y2,
                    "confidence": confidence
                }
            )

        if not candidates:

            return {
                "detected": False,
                "detections": []
            }

        # ====================================================
        # NMS
        # ====================================================

        nms_boxes = []
        nms_scores = []

        for candidate in candidates:

            x1 = candidate["x1"]
            y1 = candidate["y1"]
            x2 = candidate["x2"]
            y2 = candidate["y2"]

            nms_boxes.append(
                [
                    x1,
                    y1,
                    x2 - x1,
                    y2 - y1
                ]
            )

            nms_scores.append(
                candidate["confidence"]
            )

        indices = cv2.dnn.NMSBoxes(
            nms_boxes,
            nms_scores,
            YOLO_CONF_THRESHOLD,
            0.45
        )

        if (
            indices is None
            or len(indices) == 0
        ):

            return {
                "detected": False,
                "detections": []
            }

        indices = (
            np.array(indices)
            .reshape(-1)
        )

        candidates_nms = [
            candidates[int(i)]
            for i in indices
        ]

        # ====================================================
        # MAIOR CONFIANÇA PRIMEIRO
        # ====================================================

        candidates_nms.sort(
            key=lambda c: (
                c["confidence"]
            ),
            reverse=True
        )

        detections = []

        # ====================================================
        # GERA CROPS
        # ====================================================

        for numero, candidate in enumerate(
            candidates_nms,
            start=1
        ):

            (
                cx1,
                cy1,
                cx2,
                cy2
            ) = expandir_box(
                candidate["x1"],
                candidate["y1"],
                candidate["x2"],
                candidate["y2"],
                w_orig,
                h_orig
            )

            crop = original[
                cy1:cy2,
                cx1:cx2
            ].copy()

            if (
                crop is None
                or crop.size == 0
            ):
                continue

            success, encoded = (
                cv2.imencode(
                    ".png",
                    crop
                )
            )

            if not success:
                continue

            crop_b64 = (
                base64.b64encode(
                    encoded
                ).decode("utf-8")
            )

            detections.append(
                {
                    "detection_index":
                        numero,

                    "box":
                        [
                            cx1,
                            cy1,
                            cx2,
                            cy2
                        ],

                    "raw_box":
                        [
                            candidate["x1"],
                            candidate["y1"],
                            candidate["x2"],
                            candidate["y2"]
                        ],

                    "confidence":
                        candidate[
                            "confidence"
                        ],

                    "crop_b64":
                        crop_b64
                }
            )

        if not detections:

            return {
                "detected": False,
                "detections": []
            }

        # ====================================================
        # PRIMEIRO CROP
        # ====================================================

        primeira = detections[0]

        return {
            "detected": True,

            "detections":
                detections,

            "box":
                primeira["box"],

            "raw_box":
                primeira["raw_box"],

            "confidence":
                primeira["confidence"],

            "crop_b64":
                primeira["crop_b64"],

            "strategy":
                "YOLO 1024"
        }

    except Exception as e:

        return {
            "detected": False,
            "detections": [],
            "error": str(e)
        }


# ============================================================
# PREPARA BATCH
# ============================================================

def preparar_batch(items):

    tensors = []
    metas = []
    originais = []
    valid_indices = []

    for index, item in enumerate(
        items
    ):

        image_path = item.get(
            "image_path"
        )

        if not image_path:
            continue

        image_path = str(
            Path(
                image_path
            ).resolve()
        )

        if not Path(
            image_path
        ).exists():

            continue

        image = cv2.imread(
            image_path
        )

        if image is None:
            continue

        original = image.copy()

        blob, meta = (
            imagem_para_tensor(
                original
            )
        )

        tensors.append(blob)
        metas.append(meta)
        originais.append(original)
        valid_indices.append(index)

    if not tensors:

        return (
            None,
            [],
            [],
            []
        )

    batch = np.stack(
        tensors,
        axis=0
    )

    return (
        batch,
        metas,
        originais,
        valid_indices
    )


# ============================================================
# COMPLETA BATCH
# ============================================================

def completar_batch(batch):

    quantidade = batch.shape[0]

    if quantidade == BATCH_SIZE:
        return batch

    if quantidade == 0:
        return None

    padding = np.repeat(
        batch[-1:],
        BATCH_SIZE - quantidade,
        axis=0
    )

    return np.concatenate(
        [
            batch,
            padding
        ],
        axis=0
    )


# ============================================================
# INFERÊNCIA
# ============================================================

def inferencia(
    session,
    batch
):

    input_name = (
        session
        .get_inputs()[0]
        .name
    )

    batch_completo = (
        completar_batch(batch)
    )

    outputs = session.run(
        None,
        {
            input_name:
                batch_completo
        }
    )[0]

    return outputs


# ============================================================
# EXECUTA BATCH
#
# UMA ÚNICA TENTATIVA.
# SEM ZOOM.
# ============================================================

def executar_batch(
    session,
    items
):

    results = [
        {
            "detected": False,
            "detections": []
        }
        for _ in items
    ]

    total = len(items)

    if total == 0:
        return results

    for start in range(
        0,
        total,
        BATCH_SIZE
    ):

        sub_items = items[
            start:start + BATCH_SIZE
        ]

        (
            batch,
            metas,
            originais,
            valid_indices
        ) = preparar_batch(
            sub_items
        )

        if batch is None:
            continue

        try:

            outputs = inferencia(
                session,
                batch
            )

        except Exception as e:

            # erro continua sendo enviado
            # somente no stderr, nunca no
            # stdout usado pelo protocolo JSON

            print(
                f"[YOLO] Erro batch: {e}",
                file=sys.stderr,
                flush=True
            )

            continue

        for local_idx, original_idx in enumerate(
            valid_indices
        ):

            item = sub_items[
                original_idx
            ]

            result = processar_saida(
                outputs[
                    local_idx:
                    local_idx + 1
                ],
                metas[local_idx],
                originais[local_idx],
                item["image_path"],
                item["task_id"]
            )

            results[
                start + original_idx
            ] = result

    return results


# ============================================================
# MAIN
# ============================================================

def main():

    # ========================================================
    # ORT
    # ========================================================

    try:

        ort.preload_dlls(
            directory=""
        )

    except Exception:
        pass

    providers = [
        "CUDAExecutionProvider",
        "CPUExecutionProvider"
    ]

    try:

        session = (
            ort.InferenceSession(
                str(YOLO_MODEL),
                providers=providers
            )
        )

    except Exception as e:

        print(
            f"[YOLO] Erro carregando modelo: {e}",
            file=sys.stderr,
            flush=True
        )

        sys.exit(1)

    # ========================================================
    # INFORMAÇÕES INICIAIS
    #
    # Apenas uma vez.
    # ========================================================

    print(
        "[YOLO] Modelo carregado.",
        file=sys.stderr,
        flush=True
    )

    print(
        "[YOLO] Providers: "
        f"{session.get_providers()}",
        file=sys.stderr,
        flush=True
    )

    print(
        "[YOLO] Input: "
        f"{session.get_inputs()[0].shape}",
        file=sys.stderr,
        flush=True
    )

    print(
        f"[YOLO] Confiança mínima: "
        f"{YOLO_CONF_THRESHOLD}",
        file=sys.stderr,
        flush=True
    )

    print(
        "[YOLO] Estratégia: 1024 único, sem zoom.",
        file=sys.stderr,
        flush=True
    )

    # ========================================================
    # STDIN
    # ========================================================

    for line in sys.stdin:

        line = line.strip()

        if not line:
            continue

        try:

            req = json.loads(line)

            # =================================================
            # EXIT
            # =================================================

            if req.get(
                "command"
            ) == "exit":

                break

            # =================================================
            # BATCH
            # =================================================

            items = req.get(
                "images",
                []
            )

            resultados = executar_batch(
                session,
                items
            )

            # =================================================
            # RESPOSTA
            # =================================================

            final_results = []

            for item, result in zip(
                items,
                resultados
            ):

                result[
                    "task_id"
                ] = item.get(
                    "task_id"
                )

                result[
                    "image_path"
                ] = item.get(
                    "image_path"
                )

                final_results.append(
                    result
                )

            response = {
                "results":
                    final_results
            }

            sys.stdout.write(
                json.dumps(
                    response,
                    ensure_ascii=False
                )
                + "\n"
            )

            sys.stdout.flush()

        except Exception as e:

            print(
                f"[YOLO] Erro: {e}",
                file=sys.stderr,
                flush=True
            )

            sys.stdout.write(
                json.dumps(
                    {
                        "results": [],
                        "error": str(e)
                    },
                    ensure_ascii=False
                )
                + "\n"
            )

            sys.stdout.flush()


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()'''


if __name__ == "__main__":
    main()
