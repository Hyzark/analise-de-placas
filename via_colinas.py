from __future__ import annotations
from placas import (
    aplicar_mascara as _aplicar_mascara_comum,
    descobrir_mascara as _descobrir_mascara_comum,
    distancia_placas as _distancia_placas_comum,
    distancia_ponderada as _distancia_ponderada,
    normalizar_ocr as _normalizar_ocr_comum,
    normalizar_placa as _normalizar_placa,
    CONFUNDIVEIS as _CONFUNDIVEIS,
    MAPA_LETRAS as MAPA_LETRAS,
    MAPA_DIGITOS as MAPA_DIGITOS,
)

from common import enviar_json

import functools
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
import warnings
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from tqdm import tqdm

from config import (
    BASE_DIR,
    caminho_resultado,
    PYTHON_YOLO,
    PYTHON_OCR,
    DETECTOR_PY,
    OCR_PY,
)


# ============================================================
# OCULTA SOMENTE O WARNING DE WMF DO OPENPYXL
# ============================================================

warnings.filterwarnings(
    "ignore",
    message=r"wmf image format is not supported so the image is being dropped",
    category=UserWarning,
    module=r"openpyxl\.reader\.drawings",
)


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BATCH_SIZE = 64

MAX_DIFERENCAS_APROVACAO = 3

MAX_IMAGENS_POR_PASSAGEM = 50

NOME_ABA = "Fotos"

LINHA_CABECALHO = 5

NOME_SAIDA = "Resultado_Validacao.xlsx"

TAMANHO_PLACA = 7


# ============================================================
# MÁSCARA DA PLACA MERCOSUL
# ============================================================

MASCARA_PLACA = "LLLDLDD"


# ============================================================
# COLUNAS DA PASSAGEM
# ============================================================

CAMPOS_PASSAGEM = [
    "ID DO SIST. DE ARRECADAÇÃO",
    "DATA",
    "HORA",
    "SP",
    "KM+M",
    "PLACA",
    "QUANTIDADE DE EIXOS",
]


# ============================================================
# COLUNAS DA SAÍDA
# ============================================================

CAMPOS_OCR = [
    "OCR Bruto",
    "OCR Limpo",
    "OCR Confidence",
]

CAMPOS_COMPARACAO = [
    "Diferença OCR",
    "Diferenças de caracteres",
    "Resultado",
]

CAMPOS_CONTROLE = [
    "Imagem Usada",
    "Qtd Imagens",
    "Qtd YOLO",
    "Qtd OCR",
]

CAMPOS_SAIDA = (
    ["Arquivo Excel"]
    + CAMPOS_PASSAGEM
    + CAMPOS_OCR
    + CAMPOS_COMPARACAO
    + CAMPOS_CONTROLE
)


# ============================================================
# XML
# ============================================================

_XML_NS = {
    "xdr":
        "http://schemas.openxmlformats.org/"
        "drawingml/2006/spreadsheetDrawing",

    "a":
        "http://schemas.openxmlformats.org/"
        "drawingml/2006/main",

    "rel":
        "http://schemas.openxmlformats.org/"
        "package/2006/relationships",
}

_R_EMBED = (
    "{http://schemas.openxmlformats.org/"
    "officeDocument/2006/relationships}"
    "embed"
)


# ============================================================
# NORMALIZA COLUNA
# ============================================================

def normalizar_coluna(nome) -> str:

    if nome is None:
        return ""

    texto = str(
        nome
    ).strip().upper()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        c
        for c in texto
        if not unicodedata.combining(c)
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto,
    )

    return texto


# ============================================================
# VALOR VAZIO
# ============================================================

def value_vazio(valor) -> bool:

    if valor is None:
        return True

    if isinstance(valor, str):
        return not valor.strip()

    return False


# ============================================================
# LOCALIZA COLUNAS
# ============================================================

def localizar_colunas(sheet):

    colunas = {}

    for coluna in range(
        1,
        sheet.max_column + 1,
    ):

        valor = sheet.cell(
            row=LINHA_CABECALHO,
            column=coluna,
        ).value

        nome = normalizar_coluna(
            valor
        )

        if nome:
            colunas[nome] = coluna

    resultado = {}

    for campo in CAMPOS_PASSAGEM:

        nome_normalizado = normalizar_coluna(
            campo
        )

        if nome_normalizado in colunas:

            resultado[campo] = colunas[
                nome_normalizado
            ]

    return resultado


# ============================================================
# LIMPA PLACA
# ============================================================

def limpar_placa(texto) -> str:
    return _normalizar_placa(texto)


# ============================================================
# CONFUNDÍVEIS
# ============================================================



# ============================================================
# MAPAS DE MÁSCARA
# ============================================================




# ============================================================
# NORMALIZA CANDIDATO PELA MÁSCARA
# ============================================================

def aplicar_mascara_placa(candidato: str) -> str:
    return _aplicar_mascara_comum(candidato, MASCARA_PLACA, MAPA_LETRAS, MAPA_DIGITOS, tamanho=TAMANHO_PLACA, comprimento_invalido="vazio", preservar_indice=4)


# ============================================================
# LEVENSHTEIN
# ============================================================

def _levenshtein(a: str, b: str) -> float:
    return float(_distancia_ponderada(a, b, _CONFUNDIVEIS))


# ============================================================
# DISTÂNCIA
# ============================================================

def distancia_placas(placa_a, placa_b) -> float:
    return _distancia_placas_comum(placa_a, placa_b, _CONFUNDIVEIS)


# ============================================================
# DESCREVE DIFERENÇAS
# ============================================================

def descrever_diferencas(
    placa_esperada,
    placa_ocr,
) -> str:

    esperado = limpar_placa(
        placa_esperada
    )

    ocr = limpar_placa(
        placa_ocr
    )

    if not esperado or not ocr:
        return "Sem comparação"

    linhas = len(esperado) + 1
    colunas = len(ocr) + 1

    matriz = [
        [
            0.0
            for _ in range(
                colunas
            )
        ]
        for _ in range(
            linhas
        )
    ]

    operacoes = [
        [
            None
            for _ in range(
                colunas
            )
        ]
        for _ in range(
            linhas
        )
    ]

    for i in range(linhas):
        matriz[i][0] = float(i)

    for j in range(colunas):
        matriz[0][j] = float(j)

    for i in range(
        1,
        linhas,
    ):

        for j in range(
            1,
            colunas,
        ):

            ca = esperado[
                i - 1
            ]

            cb = ocr[
                j - 1
            ]

            if ca == cb:

                custo = 0.0

            elif (
                ca,
                cb,
            ) in _CONFUNDIVEIS:

                custo = 0.35

            else:

                custo = 1.0

            substituicao = (
                matriz[
                    i - 1
                ][
                    j - 1
                ]
                + custo
            )

            insercao = (
                matriz[i][j - 1]
                + 1.0
            )

            remocao = (
                matriz[i - 1][j]
                + 1.0
            )

            menor = min(
                substituicao,
                insercao,
                remocao,
            )

            matriz[i][j] = menor

            if menor == substituicao:

                if custo == 0:

                    operacoes[i][j] = (
                        "igual",
                        ca,
                        cb,
                    )

                else:

                    operacoes[i][j] = (
                        "substituicao",
                        ca,
                        cb,
                    )

            elif menor == insercao:

                operacoes[i][j] = (
                    "insercao",
                    "",
                    cb,
                )

            else:

                operacoes[i][j] = (
                    "remocao",
                    ca,
                    "",
                )

    i = len(esperado)
    j = len(ocr)

    diferencas = []

    while (
        i > 0
        or
        j > 0
    ):

        if (
            i > 0
            and
            j > 0
        ):

            tipo = operacoes[
                i
            ][
                j
            ][0]

        elif i > 0:

            tipo = "remocao"

        else:

            tipo = "insercao"

        if tipo == "igual":

            i -= 1
            j -= 1

        elif tipo == "substituicao":

            diferencas.append(
                (
                    i,
                    ocr[j - 1],
                    esperado[i - 1],
                )
            )

            i -= 1
            j -= 1

        elif tipo == "insercao":

            diferencas.append(
                (
                    j,
                    ocr[j - 1],
                    "—",
                )
            )

            j -= 1

        else:

            diferencas.append(
                (
                    i,
                    "—",
                    esperado[i - 1],
                )
            )

            i -= 1

    diferencas.reverse()

    if not diferencas:
        return "Nenhuma"

    resultado = []

    for (
        posicao,
        ocr_char,
        esperado_char,
    ) in diferencas:

        if esperado_char == "—":

            resultado.append(
                f"Posição {posicao}: "
                f"OCR={ocr_char} "
                f"(caractere extra)"
            )

        elif ocr_char == "—":

            resultado.append(
                f"Posição {posicao}: "
                f"OCR ausente "
                f"(esperado={esperado_char})"
            )

        else:

            resultado.append(
                f"Posição {posicao}: "
                f"OCR={ocr_char} "
                f"≠ Sistema={esperado_char}"
            )

    return "; ".join(
        resultado
    )


# ============================================================
# BUSCA MELHOR JANELA DE 7
# ============================================================

def encontrar_melhor_placa_no_texto(
    texto_ocr: str,
    placa_sistema: str,
):

    texto = limpar_placa(
        texto_ocr
    )

    placa = limpar_placa(
        placa_sistema
    )

    if not texto:
        return None

    if not placa:
        return None

    if len(texto) < TAMANHO_PLACA:
        return None

    melhor = None

    for inicio in range(
        0,
        len(texto)
        -
        TAMANHO_PLACA
        +
        1,
    ):

        candidato_bruto = texto[
            inicio:
            inicio
            +
            TAMANHO_PLACA
        ]

        candidato_limpo = (
            aplicar_mascara_placa(
                candidato_bruto
            )
        )

        if not candidato_limpo:
            continue

        distancia = distancia_placas(
            placa,
            candidato_limpo,
        )

        match_exato_bruto = (
            candidato_bruto
            ==
            placa
        )

        match_exato_limpo = (
            candidato_limpo
            ==
            placa
        )

        chave = (
            0
            if match_exato_bruto
            else
            1
            if match_exato_limpo
            else
            2,

            distancia,

            inicio,
        )

        if (
            melhor is None
            or
            chave
            <
            melhor["chave"]
        ):

            melhor = {
                "candidato_bruto":
                    candidato_bruto,

                "candidato_limpo":
                    candidato_limpo,

                "distancia":
                    distancia,

                "inicio":
                    inicio,

                "chave":
                    chave,
            }

    return melhor


# ============================================================
# EXTRAÇÃO INTELIGENTE DA PLACA
# ============================================================

def extrair_placa_do_ocr(
    texto_ocr: str,
    placa_sistema: str,
):

    texto_original = str(
        texto_ocr
        or
        ""
    ).strip()

    texto_limpo = limpar_placa(
        texto_original
    )

    placa_sistema = limpar_placa(
        placa_sistema
    )

    if not texto_limpo:

        return {
            "ocr_original":
                texto_original,

            "ocr_limpo":
                "",

            "diferenca":
                999.0,

            "candidato":
                "",

            "inicio":
                None,
        }

    # ========================================================
    # OCR COM 7 CARACTERES
    # ========================================================

    if len(texto_limpo) == TAMANHO_PLACA:

        candidato = aplicar_mascara_placa(
            texto_limpo
        )

        diferenca = distancia_placas(
            placa_sistema,
            candidato,
        )

        return {
            "ocr_original":
                texto_original,

            "ocr_limpo":
                candidato,

            "diferenca":
                diferenca,

            "candidato":
                texto_limpo,

            "inicio":
                0,
        }

    # ========================================================
    # OCR MAIOR QUE 7
    # ========================================================

    melhor = (
        encontrar_melhor_placa_no_texto(
            texto_limpo,
            placa_sistema,
        )
    )

    if melhor is None:

        return {
            "ocr_original":
                texto_original,

            "ocr_limpo":
                texto_limpo,

            "diferenca":
                distancia_placas(
                    placa_sistema,
                    texto_limpo,
                ),

            "candidato":
                texto_limpo,

            "inicio":
                None,
        }

    return {
        "ocr_original":
            texto_original,

        "ocr_limpo":
            melhor[
                "candidato_limpo"
            ],

        "diferenca":
            melhor[
                "distancia"
            ],

        "candidato":
            melhor[
                "candidato_bruto"
            ],

        "inicio":
            melhor[
                "inicio"
            ],
    }


# ============================================================
# JSON
# ============================================================



def ler_json(
    processo,
    nome,
):

    linha = (
        processo.stdout.readline()
    )

    if not linha:

        raise RuntimeError(
            f"Processo {nome} encerrou "
            "antes de responder."
        )

    try:

        return json.loads(
            linha
        )

    except json.JSONDecodeError as erro:

        raise RuntimeError(
            f"Resposta inválida do processo "
            f"{nome}: {linha[:500]}"
        ) from erro


# ============================================================
# ENCERRA
# ============================================================

def encerrar_processo(
    processo,
    nome,
):

    if processo is None:
        return

    if processo.poll() is not None:
        return

    try:

        enviar_json(
            processo,
            {
                "command":
                    "exit"
            },
        )

        processo.wait(
            timeout=10
        )

    except Exception:

        try:

            processo.terminate()

            processo.wait(
                timeout=5
            )

        except Exception:

            try:
                processo.kill()

            except Exception:
                pass


# ============================================================
# INICIA YOLO
# ============================================================

def iniciar_yolo():

    print(
        "\n[INIT] Iniciando YOLO..."
    )

    print(
        f"[YOLO] Python: {PYTHON_YOLO}"
    )

    print(
        f"[YOLO] Script: {DETECTOR_PY}"
    )

    return subprocess.Popen(
        [
            str(PYTHON_YOLO),
            str(DETECTOR_PY),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )


# ============================================================
# INICIA OCR
# ============================================================

def iniciar_ocr():

    print(
        "\n[INIT] Iniciando PaddleOCR..."
    )

    print(
        f"[OCR] Python: {PYTHON_OCR}"
    )

    print(
        f"[OCR] Script: {OCR_PY}"
    )

    return subprocess.Popen(
        [
            str(PYTHON_OCR),
            str(OCR_PY),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.DEVNULL,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )


# ============================================================
# LOCALIZA DRAWING
# ============================================================

def localizar_drawing_da_sheet(
    z: zipfile.ZipFile,
) -> str:

    arquivos = set(
        z.namelist()
    )

    rels_sheet = (
        "xl/worksheets/_rels/"
        "sheet1.xml.rels"
    )

    if rels_sheet not in arquivos:

        raise RuntimeError(
            "sheet1.xml.rels não encontrado."
        )

    sheet_rels = ET.fromstring(
        z.read(
            rels_sheet
        )
    )

    drawing_path = None

    for rel in sheet_rels.findall(
        "rel:Relationship",
        _XML_NS,
    ):

        tipo = (
            rel.get(
                "Type",
                "",
            )
            .lower()
        )

        if "drawing" not in tipo:
            continue

        target = (
            rel.get(
                "Target",
                "",
            )
            .lstrip("./")
        )

        drawing_path = (
            "xl/"
            +
            target.lstrip("/")
        )

        break

    if (
        not drawing_path
        or
        drawing_path not in arquivos
    ):

        raise RuntimeError(
            f"Drawing não encontrado: "
            f"{drawing_path}"
        )

    return drawing_path


# ============================================================
# MAPEIA IMAGENS XML
# ============================================================

def mapear_imagens_xml(
    caminho_excel: str,
) -> dict[int, list[dict]]:

    mapa: dict[
        int,
        list[dict]
    ] = {}

    with zipfile.ZipFile(
        caminho_excel,
        "r",
    ) as z:

        arquivos = set(
            z.namelist()
        )

        drawing_path = (
            localizar_drawing_da_sheet(
                z
            )
        )

        partes = drawing_path.rsplit(
            "/",
            1,
        )

        drawing_rels_path = (
            partes[0]
            +
            "/_rels/"
            +
            partes[1]
            +
            ".rels"
        )

        if drawing_rels_path not in arquivos:

            raise RuntimeError(
                "Relacionamento do drawing "
                f"não encontrado: "
                f"{drawing_rels_path}"
            )

        rels_drawing = ET.fromstring(
            z.read(
                drawing_rels_path
            )
        )

        rid_para_arquivo = {}

        for rel in rels_drawing.findall(
            "rel:Relationship",
            _XML_NS,
        ):

            rid = rel.get(
                "Id"
            )

            target = rel.get(
                "Target",
                "",
            )

            if not rid:
                continue

            rid_para_arquivo[
                rid
            ] = target.split(
                "/"
            )[-1]

        bytes_midia = {}

        for arquivo in arquivos:

            if not arquivo.startswith(
                "xl/media/"
            ):
                continue

            nome = arquivo.split(
                "/"
            )[-1]

            bytes_midia[
                nome
            ] = z.read(
                arquivo
            )

        drawing_root = ET.fromstring(
            z.read(
                drawing_path
            )
        )

        anchors = drawing_root.findall(
            "xdr:twoCellAnchor",
            _XML_NS,
        )

        anchors += drawing_root.findall(
            "xdr:oneCellAnchor",
            _XML_NS,
        )

        for indice, anchor in enumerate(
            anchors,
            start=1,
        ):

            from_elem = anchor.find(
                "xdr:from",
                _XML_NS,
            )

            if from_elem is None:
                continue

            row_elem = from_elem.find(
                "xdr:row",
                _XML_NS,
            )

            if row_elem is None:
                continue

            try:

                row_inicio = (
                    int(
                        row_elem.text
                    )
                    + 1
                )

            except Exception:

                continue

            pic = anchor.find(
                "xdr:pic",
                _XML_NS,
            )

            if pic is None:
                continue

            blip = pic.find(
                ".//a:blip",
                _XML_NS,
            )

            if blip is None:
                continue

            rid = blip.get(
                _R_EMBED
            )

            if not rid:
                continue

            nome_arquivo = (
                rid_para_arquivo.get(
                    rid,
                    "",
                )
            )

            if not nome_arquivo:
                continue

            extensao = os.path.splitext(
                nome_arquivo
            )[1].lower()

            if extensao in {
                ".emf",
                ".wmf",
            }:

                continue

            img_bytes = bytes_midia.get(
                nome_arquivo
            )

            if not img_bytes:
                continue

            mapa.setdefault(
                row_inicio,
                [],
            ).append(
                {
                    "indice_xml":
                        indice,

                    "linha":
                        row_inicio,

                    "nome":
                        nome_arquivo,

                    "dados":
                        bytes(
                            img_bytes
                        ),
                }
            )

    return mapa


# ============================================================
# EXTENSÃO POR BYTES
# ============================================================

def extensao_por_bytes(
    dados: bytes,
) -> str:

    if dados[:3] == b"\xff\xd8\xff":
        return ".jpg"

    if dados[:8] == (
        b"\x89PNG\r\n\x1a\n"
    ):
        return ".png"

    if dados[:4] == b"RIFF":
        return ".webp"

    if dados[:2] == b"BM":
        return ".bmp"

    return ".jpg"


# ============================================================
# CARREGA EXCEL
# ============================================================

def carregar_excel(
    excel_path,
    temp_dir,
):

    print(
        f"\n[EXCEL] Abrindo:\n"
        f"{excel_path}"
    )

    wb = openpyxl.load_workbook(
        excel_path,
        data_only=True,
    )

    if NOME_ABA not in wb.sheetnames:

        raise ValueError(
            f"Aba '{NOME_ABA}' não encontrada.\n"
            f"Abas existentes: "
            f"{wb.sheetnames}"
        )

    sheet = wb[
        NOME_ABA
    ]

    cabecalho = localizar_colunas(
        sheet
    )

    print(
        "[EXCEL] Colunas localizadas:"
    )

    print(
        cabecalho
    )

    faltantes = [
        campo
        for campo in CAMPOS_PASSAGEM
        if campo not in cabecalho
    ]

    if faltantes:

        raise ValueError(
            "Colunas obrigatórias não encontradas:\n"
            +
            "\n".join(
                faltantes
            )
        )

    coluna_id = cabecalho[
        "ID DO SIST. DE ARRECADAÇÃO"
    ]

    coluna_data = cabecalho[
        "DATA"
    ]

    coluna_hora = cabecalho[
        "HORA"
    ]

    coluna_placa = cabecalho[
        "PLACA"
    ]

    coluna_eixos = cabecalho[
        "QUANTIDADE DE EIXOS"
    ]

    linha_passagem = None

    for linha in range(
        LINHA_CABECALHO + 1,
        sheet.max_row + 1,
    ):

        valor_id = sheet.cell(
            row=linha,
            column=coluna_id,
        ).value

        valor_data = sheet.cell(
            row=linha,
            column=coluna_data,
        ).value

        valor_hora = sheet.cell(
            row=linha,
            column=coluna_hora,
        ).value

        valor_placa = sheet.cell(
            row=linha,
            column=coluna_placa,
        ).value

        valor_eixos = sheet.cell(
            row=linha,
            column=coluna_eixos,
        ).value

        if value_vazio(valor_id):
            continue

        if value_vazio(valor_data):
            continue

        if value_vazio(valor_hora):
            continue

        if value_vazio(valor_placa):
            continue

        if value_vazio(valor_eixos):
            continue

        linha_passagem = linha

        break

    if linha_passagem is None:

        raise ValueError(
            f"Nenhuma passagem válida encontrada "
            f"em {excel_path.name}"
        )

    dados_passagem = {}

    for campo in CAMPOS_PASSAGEM:

        coluna = cabecalho[
            campo
        ]

        dados_passagem[
            campo
        ] = sheet.cell(
            row=linha_passagem,
            column=coluna,
        ).value

    id_original = str(
        dados_passagem[
            "ID DO SIST. DE ARRECADAÇÃO"
        ]
    ).strip()

    id_limpo = limpar_placa(
        id_original
    )

    placa = limpar_placa(
        dados_passagem[
            "PLACA"
        ]
    )

    print(
        f"[PASSAGEM] "
        f"Arquivo={excel_path.name}"
    )

    print(
        f"[PASSAGEM] "
        f"Linha={linha_passagem}"
    )

    print(
        f"[PASSAGEM] "
        f"ID={id_limpo}"
    )

    print(
        f"[PASSAGEM] "
        f"Placa={placa}"
    )

    # ========================================================
    # IMAGENS
    # ========================================================

    mapa_imagens = mapear_imagens_xml(
        str(
            excel_path
        )
    )

    imagens_xml = []

    for linha_ancora in sorted(
        mapa_imagens.keys()
    ):

        if linha_ancora < linha_passagem:
            continue

        imagens_xml.extend(
            mapa_imagens[
                linha_ancora
            ]
        )

    if len(imagens_xml) > MAX_IMAGENS_POR_PASSAGEM:

        print(
            f"[AVISO] "
            f"{len(imagens_xml)} imagens encontradas. "
            f"Usando somente as primeiras "
            f"{MAX_IMAGENS_POR_PASSAGEM}."
        )

        imagens_xml = imagens_xml[
            :MAX_IMAGENS_POR_PASSAGEM
        ]

    print(
        f"[EXCEL] "
        f"Imagens da passagem: "
        f"{len(imagens_xml)}"
    )

    # ========================================================
    # PREPARA
    # ========================================================

    imagens = []

    prefixo_arquivo = re.sub(
        r"[^a-zA-Z0-9]",
        "_",
        excel_path.stem,
    )

    for indice, item in enumerate(
        imagens_xml
    ):

        dados_imagem = item[
            "dados"
        ]

        if not dados_imagem:
            continue

        extensao = extensao_por_bytes(
            dados_imagem
        )

        caminho = (
            temp_dir
            /
            (
                f"{prefixo_arquivo}_"
                f"p0_i{indice}"
                f"{extensao}"
            )
        )

        caminho.write_bytes(
            dados_imagem
        )

        nome_original = item[
            "nome"
        ]

        if not nome_original:

            nome_original = (
                f"Imagem_{indice + 1}"
            )

        imagens.append(
            {
                "task_id":
                    (
                        f"{prefixo_arquivo}_"
                        f"p0_i{indice}"
                    ),

                "image_path":
                    str(caminho),

                "indice":
                    indice,

                "indice_original":
                    item[
                        "indice_xml"
                    ],

                "linha":
                    item[
                        "linha"
                    ],

                "nome_original":
                    nome_original,
            }
        )

    print(
        f"[PASSAGEM] "
        f"Imagens preparadas="
        f"{len(imagens)}"
    )

    wb.close()

    return [
        {
            "uid":
                excel_path.stem,

            "arquivo":
                excel_path.name,

            "linha":
                linha_passagem,

            "id":
                id_limpo,

            "placa":
                placa,

            "dados":
                dados_passagem,

            "imagens":
                imagens,
        }
    ]


# ============================================================
# YOLO
# ============================================================

def executar_yolo(
    processo_yolo,
    imagens,
):

    if not imagens:
        return {}

    items = []

    for imagem in imagens:

        caminho = Path(
            imagem[
                "image_path"
            ]
        ).resolve()

        if (
            not caminho.exists()
            or
            caminho.stat().st_size == 0
        ):
            continue

        items.append(
            {
                "task_id":
                    imagem[
                        "task_id"
                    ],

                "image_path":
                    str(caminho),
            }
        )

    if not items:
        return {}

    enviar_json(
        processo_yolo,
        {
            "command":
                "detect_batch_all",

            "images":
                items,
        },
    )

    resposta = ler_json(
        processo_yolo,
        "YOLO",
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        [],
    ):

        task_id = resultado.get(
            "task_id"
        )

        if task_id:

            resultados[
                task_id
            ] = resultado

    return resultados


# ============================================================
# OCR
# ============================================================

def executar_ocr(
    processo_ocr,
    yolo_results,
):

    if not yolo_results:
        return {}

    crops = []

    for (
        task_id_imagem,
        resultado,
    ) in yolo_results.items():

        if not resultado.get(
            "detected",
            False,
        ):
            continue

        detections = resultado.get(
            "detections",
            [],
        )

        if not detections:

            crop_b64 = resultado.get(
                "crop_b64"
            )

            if crop_b64:

                detections = [
                    {
                        "detection_index":
                            1,

                        "crop_b64":
                            crop_b64,
                    }
                ]

        for detection in detections:

            detection_index = detection.get(
                "detection_index",
                1,
            )

            crop_b64 = detection.get(
                "crop_b64"
            )

            if not crop_b64:
                continue

            task_id_ocr = (
                f"{task_id_imagem}"
                f"_d"
                f"{detection_index}"
            )

            crops.append(
                {
                    "task_id":
                        task_id_ocr,

                    "crop_b64":
                        crop_b64,
                }
            )

    if not crops:
        return {}

    enviar_json(
        processo_ocr,
        {
            "command":
                "ocr_batch",

            "images":
                crops,
        },
    )

    resposta = ler_json(
        processo_ocr,
        "OCR",
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        [],
    ):

        task_id = resultado.get(
            "task_id"
        )

        if not task_id:
            continue

        resultados[
            task_id
        ] = resultado

    return resultados


# ============================================================
# PEGA OCR DA IMAGEM
# ============================================================

def obter_ocr_da_imagem(
    imagem,
    ocr_results,
):

    resultados = []

    task_id_base = imagem[
        "task_id"
    ]

    for detection_index in range(
        1,
        MAX_IMAGENS_POR_PASSAGEM + 1,
    ):

        task_id_ocr = (
            f"{task_id_base}"
            f"_d"
            f"{detection_index}"
        )

        resultado = ocr_results.get(
            task_id_ocr
        )

        if resultado is None:
            continue

        resultados.append(
            (
                task_id_ocr,
                resultado,
            )
        )

    return resultados


# ============================================================
# AVALIA OCR
# ============================================================

def avaliar_ocr(
    placa_sistema,
    resultado_ocr,
):

    placa_sistema = limpar_placa(
        placa_sistema
    )

    texto_paddle = resultado_ocr.get(
        "ocr_best",
        "",
    )

    if texto_paddle is None:
        texto_paddle = ""

    texto_paddle = str(
        texto_paddle
    ).strip()

    try:

        confianca = float(
            resultado_ocr.get(
                "confidence_best",
                0.0,
            )
            or
            0.0
        )

    except Exception:

        confianca = 0.0

    extracao = (
        extrair_placa_do_ocr(
            texto_paddle,
            placa_sistema,
        )
    )

    ocr_original = extracao[
        "ocr_original"
    ]

    ocr_limpo = extracao[
        "ocr_limpo"
    ]

    diferenca = extracao[
        "diferenca"
    ]

    candidato = extracao[
        "candidato"
    ]

    if not ocr_limpo:

        return {
            "aprovado":
                False,

            "ocr_bruto":
                ocr_original,

            "ocr_limpo":
                "",

            "confianca":
                confianca,

            "diferenca":
                999.0,

            "diferencas_caracteres":
                "Sem OCR",

            "candidato":
                candidato,
        }

    diferencas = (
        descrever_diferencas(
            placa_sistema,
            ocr_limpo,
        )
    )

    aprovado = (
        diferenca
        <=
        MAX_DIFERENCAS_APROVACAO
    )

    return {
        "aprovado":
            aprovado,

        "ocr_bruto":
            ocr_original,

        "ocr_limpo":
            ocr_limpo,

        "confianca":
            confianca,

        "diferenca":
            round(
                diferenca,
                2,
            ),

        "diferencas_caracteres":
            diferencas,

        "candidato":
            candidato,
    }


# ============================================================
# AVALIA IMAGEM
# ============================================================

def avaliar_imagem(
    passagem,
    imagem,
    yolo_result,
    ocr_results,
):

    if not yolo_result:

        return {
            "aprovado":
                False,

            "ocr_bruto":
                "",

            "ocr_limpo":
                "",

            "confianca":
                0.0,

            "diferenca":
                999.0,

            "diferencas_caracteres":
                "Sem detecção",
        }

    if not yolo_result.get(
        "detected",
        False,
    ):

        return {
            "aprovado":
                False,

            "ocr_bruto":
                "",

            "ocr_limpo":
                "",

            "confianca":
                0.0,

            "diferenca":
                999.0,

            "diferencas_caracteres":
                "Sem detecção",
        }

    ocr_da_imagem = (
        obter_ocr_da_imagem(
            imagem,
            ocr_results,
        )
    )

    if not ocr_da_imagem:

        return {
            "aprovado":
                False,

            "ocr_bruto":
                "",

            "ocr_limpo":
                "",

            "confianca":
                0.0,

            "diferenca":
                999.0,

            "diferencas_caracteres":
                "Sem OCR",
        }

    melhor = None
    melhor_task_id = ""

    for (
        task_id_ocr,
        ocr_result,
    ) in ocr_da_imagem:

        avaliacao = avaliar_ocr(
            passagem[
                "placa"
            ],
            ocr_result,
        )

        if avaliacao.get(
            "aprovado",
            False,
        ):

            avaliacao[
                "_task_id_ocr"
            ] = task_id_ocr

            return avaliacao

        if (
            melhor is None
            or
            avaliacao[
                "diferenca"
            ]
            <
            melhor[
                "diferenca"
            ]
        ):

            melhor = avaliacao

            melhor_task_id = (
                task_id_ocr
            )

    if melhor is not None:

        melhor[
            "_task_id_ocr"
        ] = melhor_task_id

        return melhor

    return {
        "aprovado":
            False,

        "ocr_bruto":
            "",

        "ocr_limpo":
            "",

        "confianca":
            0.0,

        "diferenca":
            999.0,

        "diferencas_caracteres":
            "Sem OCR",
    }


# ============================================================
# RESULTADO SEM LEITURA
# ============================================================

def resultado_sem_leitura():

    return {
        "ocr_bruto":
            "",

        "ocr_limpo":
            "",

        "confianca":
            0.0,

        "diferenca":
            999.0,

        "diferencas_caracteres":
            "Nenhuma leitura OCR",
    }


# ============================================================
# PROCESSAMENTO
# ============================================================

def processar_pasta(
    pasta,
):

    pasta = Path(
        pasta
    ).resolve()

    arquivos = sorted(
        [
            arquivo
            for arquivo in pasta.rglob(
                "*.xlsx"
            )
            if (
                not arquivo.name.startswith(
                    "~$"
                )
                and
                arquivo.stem
                !=
                Path(
                    NOME_SAIDA
                ).stem
                and
                not arquivo.stem.endswith(
                    "_validado"
                )
            )
        ]
    )

    if not arquivos:

        print(
            "[ERRO] Nenhum XLSX encontrado."
        )

        return

    # ========================================================
    # STATUS GERAL
    # ========================================================

    print(
        "\n"
        + "=" * 75
    )

    print(
        "VALIDADOR DE FOTOS"
    )

    print(
        "=" * 75
    )

    print(
        f"Pasta: {pasta}"
    )

    print(
        f"Arquivos encontrados: "
        f"{len(arquivos)}"
    )

    print(
        "=" * 75
    )

    # ========================================================
    # CHECK
    # ========================================================

    print(
        "\n[CHECK] Verificando arquivos..."
    )

    for arquivo in (
        PYTHON_YOLO,
        PYTHON_OCR,
        DETECTOR_PY,
        OCR_PY,
    ):

        if not Path(
            arquivo
        ).exists():

            raise FileNotFoundError(
                f"Arquivo não encontrado:\n"
                f"{arquivo}"
            )

        print(
            f"[OK] {arquivo}"
        )

    temp_dir = Path(
        tempfile.mkdtemp(
            prefix=
            "validador_fotos_"
        )
    )

    processo_yolo = None
    processo_ocr = None

    try:

        # ====================================================
        # CARREGA ARQUIVOS
        # ====================================================

        todas_passagens = []

        arquivos_ok = 0
        arquivos_erro = 0

        print(
            "\n[LEITURA] Carregando arquivos..."
        )

        for numero, arquivo in enumerate(
            arquivos,
            start=1,
        ):

            try:

                passagens = carregar_excel(
                    arquivo,
                    temp_dir,
                )

                if len(passagens) != 1:

                    raise RuntimeError(
                        f"O arquivo "
                        f"{arquivo.name} "
                        f"retornou "
                        f"{len(passagens)} "
                        f"passagens. "
                        f"O esperado é 1."
                    )

                todas_passagens.extend(
                    passagens
                )

                arquivos_ok += 1

            except Exception as erro:

                arquivos_erro += 1

                print(
                    f"\n[ERRO AO LER EXCEL] "
                    f"{arquivo.name}: "
                    f"{type(erro).__name__}: "
                    f"{erro}"
                )

        total_passagens = len(
            todas_passagens
        )

        total_imagens = sum(
            len(
                passagem[
                    "imagens"
                ]
            )
            for passagem in (
                todas_passagens
            )
        )

        print(
            "\n"
            + "=" * 75
        )

        print(
            "DADOS CARREGADOS"
        )

        print(
            "=" * 75
        )

        print(
            f"Arquivos lidos: "
            f"{arquivos_ok}"
        )

        print(
            f"Arquivos com erro: "
            f"{arquivos_erro}"
        )

        print(
            f"Passagens: "
            f"{total_passagens}"
        )

        print(
            f"Imagens: "
            f"{total_imagens}"
        )

        print(
            "=" * 75
        )

        if not todas_passagens:

            print(
                "[ERRO] Nenhuma passagem encontrada."
            )

            return

        # ====================================================
        # INICIA MODELOS
        # ====================================================

        processo_yolo = iniciar_yolo()

        processo_ocr = iniciar_ocr()

        print(
            "\n[INIT] Aguardando carregamento dos modelos..."
        )

        time.sleep(
            15
        )

        print(
            "[INIT] Modelos carregados. "
            "Iniciando análise..."
        )

        # ====================================================
        # CONTROLE
        # ====================================================

        ids_aprovados = set()
        ids_esgotados = set()

        proxima_imagem = {
            passagem[
                "uid"
            ]: 0
            for passagem
            in todas_passagens
        }

        melhor_resultado = {}

        qtd_yolo_por_id = {}
        qtd_ocr_por_id = {}

        total_yolo = 0
        total_detectadas = 0
        total_ocr = 0

        # ====================================================
        # BARRA ÚNICA
        # ====================================================

        barra = tqdm(
            total=
                total_passagens,

            desc=
                "YOLO + OCR",

            unit=
                "ID",

            dynamic_ncols=
                True,

            leave=
                True,
        )

        concluidos_anteriores = 0

        try:

            while True:

                # ==========================================
                # UMA IMAGEM POR PASSAGEM
                # ==========================================

                lote_rodada = []

                for passagem in (
                    todas_passagens
                ):

                    uid = passagem[
                        "uid"
                    ]

                    if uid in ids_aprovados:
                        continue

                    if uid in ids_esgotados:
                        continue

                    indice = (
                        proxima_imagem[
                            uid
                        ]
                    )

                    imagens = passagem[
                        "imagens"
                    ]

                    if not imagens:

                        ids_esgotados.add(
                            uid
                        )

                        continue

                    if (
                        indice
                        >=
                        len(
                            imagens
                        )
                    ):

                        ids_esgotados.add(
                            uid
                        )

                        continue

                    imagem = imagens[
                        indice
                    ]

                    proxima_imagem[
                        uid
                    ] = (
                        indice + 1
                    )

                    lote_rodada.append(
                        (
                            passagem,
                            imagem,
                        )
                    )

                if not lote_rodada:
                    break

                # ==========================================
                # LOTES DE 64
                # ==========================================

                for inicio_lote in range(
                    0,
                    len(
                        lote_rodada
                    ),
                    BATCH_SIZE,
                ):

                    lote = lote_rodada[
                        inicio_lote:
                        inicio_lote
                        +
                        BATCH_SIZE
                    ]

                    imagens_lote = [
                        imagem
                        for (
                            _,
                            imagem,
                        ) in lote
                    ]

                    # ======================================
                    # YOLO
                    # ======================================

                    yolo_resultados = (
                        executar_yolo(
                            processo_yolo,
                            imagens_lote,
                        )
                    )

                    total_yolo += len(
                        imagens_lote
                    )

                    total_detectadas += sum(
                        1
                        for resultado
                        in (
                            yolo_resultados.values()
                        )
                        if resultado.get(
                            "detected",
                            False,
                        )
                    )

                    # ======================================
                    # CONTAGEM YOLO
                    # ======================================

                    for (
                        passagem,
                        imagem,
                    ) in lote:

                        uid = passagem[
                            "uid"
                        ]

                        qtd_yolo_por_id[
                            uid
                        ] = (
                            qtd_yolo_por_id.get(
                                uid,
                                0,
                            )
                            + 1
                        )

                    # ======================================
                    # OCR
                    # ======================================

                    ocr_resultados = (
                        executar_ocr(
                            processo_ocr,
                            yolo_resultados,
                        )
                    )

                    total_ocr += len(
                        ocr_resultados
                    )

                    # ======================================
                    # CONTAGEM OCR
                    # ======================================

                    for (
                        passagem,
                        imagem,
                    ) in lote:

                        uid = passagem[
                            "uid"
                        ]

                        task_base = imagem[
                            "task_id"
                        ]

                        quantidade = 0

                        for task_id_ocr in (
                            ocr_resultados
                        ):

                            if task_id_ocr.startswith(
                                f"{task_base}_d"
                            ):

                                quantidade += 1

                        qtd_ocr_por_id[
                            uid
                        ] = (
                            qtd_ocr_por_id.get(
                                uid,
                                0,
                            )
                            + quantidade
                        )

                    # ======================================
                    # AVALIA
                    # ======================================

                    for (
                        passagem,
                        imagem,
                    ) in lote:

                        uid = passagem[
                            "uid"
                        ]

                        if uid in ids_aprovados:
                            continue

                        task_id = imagem[
                            "task_id"
                        ]

                        yolo_result = (
                            yolo_resultados.get(
                                task_id,
                                {},
                            )
                        )

                        avaliacao = (
                            avaliar_imagem(
                                passagem,
                                imagem,
                                yolo_result,
                                ocr_resultados,
                            )
                        )

                        diferenca_atual = float(
                            avaliacao.get(
                                "diferenca",
                                999.0,
                            )
                            or
                            999.0
                        )

                        ocr_atual = limpar_placa(
                            avaliacao.get(
                                "ocr_limpo",
                                "",
                            )
                        )

                        # ==================================
                        # APROVOU
                        # ==================================

                        if avaliacao.get(
                            "aprovado",
                            False,
                        ):

                            melhor_resultado[
                                uid
                            ] = {

                                **avaliacao,

                                "imagem":
                                    imagem[
                                        "indice"
                                    ],

                                "nome_imagem":
                                    imagem[
                                        "nome_original"
                                    ],
                            }

                            ids_aprovados.add(
                                uid
                            )

                            continue

                        # ==================================
                        # MELHOR RESULTADO
                        # ==================================

                        anterior = (
                            melhor_resultado.get(
                                uid
                            )
                        )

                        if anterior is None:

                            melhor_resultado[
                                uid
                            ] = {

                                **avaliacao,

                                "imagem":
                                    imagem[
                                        "indice"
                                    ],

                                "nome_imagem":
                                    imagem[
                                        "nome_original"
                                    ],
                            }

                            continue

                        diferenca_anterior = float(
                            anterior.get(
                                "diferenca",
                                999.0,
                            )
                            or
                            999.0
                        )

                        ocr_anterior = limpar_placa(
                            anterior.get(
                                "ocr_limpo",
                                "",
                            )
                        )

                        try:

                            confianca_atual = float(
                                avaliacao.get(
                                    "confianca",
                                    0.0,
                                )
                                or
                                0.0
                            )

                        except Exception:

                            confianca_atual = 0.0

                        try:

                            confianca_anterior = float(
                                anterior.get(
                                    "confianca",
                                    0.0,
                                )
                                or
                                0.0
                            )

                        except Exception:

                            confianca_anterior = 0.0

                        deve_substituir = False

                        if (
                            diferenca_atual
                            <
                            diferenca_anterior
                        ):

                            deve_substituir = True

                        elif (
                            diferenca_atual
                            ==
                            diferenca_anterior
                            and
                            ocr_atual
                            and
                            not ocr_anterior
                        ):

                            deve_substituir = True

                        elif (
                            diferenca_atual
                            ==
                            diferenca_anterior
                            and
                            ocr_atual
                            and
                            ocr_anterior
                            and
                            confianca_atual
                            >
                            confianca_anterior
                        ):

                            deve_substituir = True

                        if deve_substituir:

                            melhor_resultado[
                                uid
                            ] = {

                                **avaliacao,

                                "imagem":
                                    imagem[
                                        "indice"
                                    ],

                                "nome_imagem":
                                    imagem[
                                        "nome_original"
                                    ],
                            }

                # ==========================================
                # ESGOTADOS
                # ==========================================

                for passagem in todas_passagens:

                    uid = passagem[
                        "uid"
                    ]

                    if uid in ids_aprovados:
                        continue

                    if uid in ids_esgotados:
                        continue

                    if (
                        proxima_imagem[
                            uid
                        ]
                        >=
                        len(
                            passagem[
                                "imagens"
                            ]
                        )
                    ):

                        ids_esgotados.add(
                            uid
                        )

                # ==========================================
                # PROGRESSO DA ANÁLISE
                # ==========================================

                concluidos = (
                    len(
                        ids_aprovados
                    )
                    +
                    len(
                        ids_esgotados
                    )
                )

                delta = (
                    concluidos
                    -
                    concluidos_anteriores
                )

                if delta > 0:

                    barra.update(
                        delta
                    )

                    concluidos_anteriores = (
                        concluidos
                    )

                barra.set_postfix(
                    aprovados=
                        len(
                            ids_aprovados
                        ),

                    revisar=
                        len(
                            ids_esgotados
                        ),

                    pendentes=
                        max(
                            0,
                            total_passagens
                            -
                            concluidos,
                        ),

                    ocr=
                        total_ocr,
                )

        finally:

            barra.close()

        # ====================================================
        # RESULTADO
        # ====================================================

        registros = []

        for passagem in (
            todas_passagens
        ):

            uid = passagem[
                "uid"
            ]

            melhor = (
                melhor_resultado.get(
                    uid
                )
            )

            if melhor is None:

                melhor = (
                    resultado_sem_leitura()
                )

            resultado_final = (
                "Aprovado"
                if (
                    uid in ids_aprovados
                    or
                    melhor.get(
                        "aprovado",
                        False,
                    )
                )
                else
                "Revisar"
            )

            linha = {}

            linha[
                "Arquivo Excel"
            ] = passagem[
                "arquivo"
            ]

            for campo in CAMPOS_PASSAGEM:

                linha[
                    campo
                ] = passagem[
                    "dados"
                ].get(
                    campo,
                    "",
                )

            linha[
                "OCR Bruto"
            ] = melhor.get(
                "ocr_bruto",
                "",
            )

            linha[
                "OCR Limpo"
            ] = melhor.get(
                "ocr_limpo",
                "",
            )

            linha[
                "OCR Confidence"
            ] = melhor.get(
                "confianca",
                0.0,
            )

            linha[
                "Diferença OCR"
            ] = melhor.get(
                "diferenca",
                999.0,
            )

            linha[
                "Diferenças de caracteres"
            ] = melhor.get(
                "diferencas_caracteres",
                "Nenhuma leitura OCR",
            )

            linha[
                "Resultado"
            ] = resultado_final

            if "imagem" in melhor:

                linha[
                    "Imagem Usada"
                ] = (
                    f"Imagem "
                    f"{int(melhor['imagem']) + 1}"
                )

            else:

                linha[
                    "Imagem Usada"
                ] = ""

            linha[
                "Qtd Imagens"
            ] = len(
                passagem[
                    "imagens"
                ]
            )

            linha[
                "Qtd YOLO"
            ] = qtd_yolo_por_id.get(
                uid,
                0,
            )

            linha[
                "Qtd OCR"
            ] = qtd_ocr_por_id.get(
                uid,
                0,
            )

            registros.append(
                linha
            )

        # ====================================================
        # DATAFRAME
        # ====================================================

        df = pd.DataFrame(
            registros,
            columns=
                CAMPOS_SAIDA,
        )

        # ====================================================
        # ARQUIVO DE SAÍDA
        # ====================================================

        saida = caminho_resultado("via_colinas", NOME_SAIDA)

        if saida.exists():
            saida.unlink()

        workbook = (
            openpyxl.Workbook()
        )

        ws = (
            workbook.active
        )

        ws.title = (
            "Resultado"
        )

        # ====================================================
        # CABEÇALHO
        # ====================================================

        for coluna, nome in enumerate(
            CAMPOS_SAIDA,
            start=1,
        ):

            celula = ws.cell(
                row=1,
                column=coluna,
                value=nome,
            )

            celula.font = Font(
                bold=True
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        ws.row_dimensions[
            1
        ].height = 35

        # ====================================================
        # DADOS
        # ====================================================

        for (
            linha_excel,
            (
                _,
                registro,
            ),
        ) in enumerate(
            df.iterrows(),
            start=2,
        ):

            for coluna, nome in enumerate(
                CAMPOS_SAIDA,
                start=1,
            ):

                valor = registro[
                    nome
                ]

                if pd.isna(
                    valor
                ):

                    valor = ""

                celula = ws.cell(
                    row=linha_excel,
                    column=coluna,
                    value=valor,
                )

                celula.alignment = Alignment(
                    vertical="top",
                    wrap_text=(
                        nome
                        ==
                        "Diferenças de caracteres"
                    ),
                )

        # ====================================================
        # FREEZE
        # ====================================================

        ws.freeze_panes = "A2"

        # ====================================================
        # FILTRO
        # ====================================================

        ws.auto_filter.ref = (
            ws.dimensions
        )

        # ====================================================
        # LARGURAS
        # ====================================================

        larguras = {

            "Arquivo Excel":
                45,

            "ID DO SIST. DE ARRECADAÇÃO":
                25,

            "DATA":
                14,

            "HORA":
                12,

            "SP":
                12,

            "KM+M":
                12,

            "PLACA":
                18,

            "QUANTIDADE DE EIXOS":
                20,

            "OCR Bruto":
                40,

            "OCR Limpo":
                25,

            "OCR Confidence":
                18,

            "Diferença OCR":
                16,

            "Diferenças de caracteres":
                55,

            "Resultado":
                16,

            "Imagem Usada":
                16,

            "Qtd Imagens":
                14,

            "Qtd YOLO":
                12,

            "Qtd OCR":
                12,
        }

        cabecalhos = {
            ws.cell(
                row=1,
                column=coluna,
            ).value:
                coluna

            for coluna in range(
                1,
                ws.max_column + 1,
            )
        }

        for nome, largura in (
            larguras.items()
        ):

            coluna = cabecalhos.get(
                nome
            )

            if coluna is None:
                continue

            ws.column_dimensions[
                get_column_letter(
                    coluna
                )
            ].width = largura

        # ====================================================
        # FORMATO CONFIDENCE
        # ====================================================

        coluna = cabecalhos.get(
            "OCR Confidence"
        )

        if coluna:

            for linha in range(
                2,
                ws.max_row + 1,
            ):

                ws.cell(
                    row=linha,
                    column=coluna,
                ).number_format = (
                    "0.0000"
                )

        # ====================================================
        # FORMATO DIFERENÇA
        # ====================================================

        coluna = cabecalhos.get(
            "Diferença OCR"
        )

        if coluna:

            for linha in range(
                2,
                ws.max_row + 1,
            ):

                ws.cell(
                    row=linha,
                    column=coluna,
                ).number_format = (
                    "0.00"
                )

        # ====================================================
        # CORES
        # ====================================================

        coluna = cabecalhos.get(
            "Resultado"
        )

        if coluna:

            for linha in range(
                2,
                ws.max_row + 1,
            ):

                celula = ws.cell(
                    row=linha,
                    column=coluna,
                )

                if (
                    celula.value
                    ==
                    "Aprovado"
                ):

                    celula.fill = PatternFill(
                        fill_type=
                            "solid",

                        fgColor=
                            "C6EFCE",
                    )

                    celula.font = Font(
                        bold=True,
                        color=
                            "006100",
                    )

                else:

                    celula.fill = PatternFill(
                        fill_type=
                            "solid",

                        fgColor=
                            "FFC7CE",
                    )

                    celula.font = Font(
                        bold=True,
                        color=
                            "9C0006",
                    )

                celula.alignment = Alignment(
                    horizontal=
                        "center",

                    vertical=
                        "center",
                )

        # ====================================================
        # DIVISÓRIAS
        # ====================================================

        coluna_ocr_inicio = (
            len(
                CAMPOS_PASSAGEM
            )
            + 2
        )

        coluna_comparacao_inicio = (
            len(
                CAMPOS_PASSAGEM
            )
            +
            len(
                CAMPOS_OCR
            )
            + 2
        )

        linha_divisoria = Side(
            style="medium",
            color="808080",
        )

        for linha in range(
            1,
            ws.max_row + 1,
        ):

            for coluna in (
                coluna_ocr_inicio,
                coluna_comparacao_inicio,
            ):

                celula = ws.cell(
                    row=linha,
                    column=coluna,
                )

                celula.border = Border(
                    left=
                        celula.border.left,

                    right=
                        linha_divisoria,

                    top=
                        celula.border.top,

                    bottom=
                        celula.border.bottom,
                )

        # ====================================================
        # SALVA
        # ====================================================

        workbook.save(
            saida
        )

        # ====================================================
        # RESUMO FINAL
        # ====================================================

        aprovados = len(
            ids_aprovados
        )

        revisar = (
            total_passagens
            -
            aprovados
        )

        print(
            "\n"
            + "=" * 75
        )

        print(
            "RESULTADO FINAL"
        )

        print(
            "=" * 75
        )

        print(
            f"Arquivos encontrados: "
            f"{len(arquivos)}"
        )

        print(
            f"Arquivos lidos: "
            f"{arquivos_ok}"
        )

        print(
            f"Arquivos com erro: "
            f"{arquivos_erro}"
        )

        print(
            f"Passagens: "
            f"{total_passagens}"
        )

        print(
            f"Aprovados: "
            f"{aprovados}"
        )

        print(
            f"Revisar: "
            f"{revisar}"
        )

        print(
            f"Imagens existentes: "
            f"{total_imagens}"
        )

        print(
            f"Imagens enviadas ao YOLO: "
            f"{total_yolo}"
        )

        print(
            f"Detecções YOLO: "
            f"{total_detectadas}"
        )

        print(
            f"OCR executados: "
            f"{total_ocr}"
        )

        print(
            f"Saída única: "
            f"{saida}"
        )

        print(
            "=" * 75
        )

        return df

    finally:

        encerrar_processo(
            processo_yolo,
            "YOLO",
        )

        encerrar_processo(
            processo_ocr,
            "OCR",
        )

        shutil.rmtree(
            temp_dir,
            ignore_errors=True,
        )


# ============================================================
# MAIN
# ============================================================

def main():

    if len(sys.argv) < 2:

        print(
            'Uso: python validador_fotos.py '
            '"C:\\caminho\\da\\pasta"'
        )

        sys.exit(1)

    pasta = Path(
        sys.argv[1]
    ).resolve()

    if not pasta.exists():

        print(
            f"[ERRO] Pasta não encontrada:\n"
            f"{pasta}"
        )

        sys.exit(1)

    if not pasta.is_dir():

        print(
            f"[ERRO] O caminho não é uma pasta:\n"
            f"{pasta}"
        )

        sys.exit(1)

    processar_pasta(
        pasta
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()
