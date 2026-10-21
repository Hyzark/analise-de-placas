from __future__ import annotations

import sys
import json
import re
import subprocess
import tempfile
import shutil
import time
from pathlib import Path

import openpyxl
import pandas as pd


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BATCH_SIZE = 64
MAX_DIFERENCAS_APROVACAO = 3

TAMANHO_PLACA = 7

# Máscara Mercosul:
# LLLDLDD
#
# 0 = letra
# 1 = letra
# 2 = letra
# 3 = dígito
# 4 = letra
# 5 = dígito
# 6 = dígito
#
MASCARA_PLACA = "LLLDLDD"

BASE_DIR = Path(__file__).resolve().parent

PYTHON_YOLO = (
    BASE_DIR
    / ".venv_yolo"
    / "Scripts"
    / "python.exe"
)

PYTHON_OCR = (
    BASE_DIR
    / ".venv_ocr"
    / "Scripts"
    / "python.exe"
)

DETECTOR_PY = BASE_DIR / "detector_process.py"
OCR_PY = BASE_DIR / "ocr_process.py"


# ============================================================
# NORMALIZAÇÃO DE NOMES DE COLUNAS
# ============================================================

def normalizar_nome_coluna(nome):

    nome = str(
        nome or ""
    ).strip().upper()

    nome = (
        nome
        .replace("Á", "A")
        .replace("À", "A")
        .replace("Ã", "A")
        .replace("Â", "A")
        .replace("É", "E")
        .replace("Ê", "E")
        .replace("Í", "I")
        .replace("Ó", "O")
        .replace("Õ", "O")
        .replace("Ô", "O")
        .replace("Ú", "U")
        .replace("Ç", "C")
    )

    nome = re.sub(
        r"[^A-Z0-9]+",
        "_",
        nome,
    )

    return nome.strip("_")


# ============================================================
# LIMPA TEXTO
# ============================================================

def limpar_placa(texto):

    return re.sub(
        r"[^A-Z0-9]",
        "",
        str(texto or "").upper(),
    )


# ============================================================
# CONFUNDÍVEIS
# ============================================================

CONFUNDIVEIS = {
    ("0", "O"),
    ("O", "0"),

    ("1", "I"),
    ("I", "1"),

    ("1", "L"),
    ("L", "1"),

    ("T", "I"),
    ("I", "T"),

    ("2", "Z"),
    ("Z", "2"),

    ("Z", "7"),
    ("7", "Z"),

    ("4", "A"),
    ("A", "4"),

    ("5", "S"),
    ("S", "5"),

    ("6", "G"),
    ("G", "6"),

    ("8", "B"),
    ("B", "8"),

    ("B", "E"),
    ("E", "B"),

    ("D", "0"),
    ("0", "D"),

    ("D", "O"),
    ("O", "D"),
}


# ============================================================
# MAPA PARA POSIÇÕES DE LETRAS
# ============================================================

MAPA_LETRAS = {
    "0": "O",
    "1": "I",
    "2": "Z",
    "4": "A",
    "5": "S",
    "6": "G",
    "7": "Z",
    "8": "B",
}


# ============================================================
# MAPA PARA POSIÇÕES DE DÍGITOS
# ============================================================

MAPA_DIGITOS = {
    "O": "0",
    "Q": "0",
    "D": "0",

    "I": "1",
    "L": "1",
    "T": "7",

    "Z": "2",

    "A": "4",

    "S": "5",

    "G": "6",

    "B": "8",
}


# ============================================================
# APLICA MÁSCARA LLLDLDD
# ============================================================

def aplicar_mascara_placa(
    candidato,
):

    candidato = limpar_placa(
        candidato
    )

    if len(candidato) != TAMANHO_PLACA:
        return ""

    resultado = []

    for indice, char in enumerate(
        candidato
    ):

        # ====================================================
        # POSIÇÃO 5 — NUNCA ALTERAR
        # ====================================================

        if indice == 4:

            resultado.append(
                char
            )

            continue

        tipo = MASCARA_PLACA[
            indice
        ]

        # ====================================================
        # POSIÇÃO DE LETRA
        # ====================================================

        if tipo == "L":

            if char.isalpha():

                resultado.append(
                    char
                )

            elif char in MAPA_LETRAS:

                resultado.append(
                    MAPA_LETRAS[
                        char
                    ]
                )

            else:

                resultado.append(
                    char
                )

        # ====================================================
        # POSIÇÃO DE DÍGITO
        # ====================================================

        else:

            if char.isdigit():

                resultado.append(
                    char
                )

            elif char in MAPA_DIGITOS:

                resultado.append(
                    MAPA_DIGITOS[
                        char
                    ]
                )

            else:

                resultado.append(
                    char
                )

    return "".join(
        resultado
    )


# ============================================================
# LEVENSHTEIN
# ============================================================

def distancia_placas(
    a,
    b,
):

    a = limpar_placa(
        a
    )

    b = limpar_placa(
        b
    )

    if not a or not b:
        return 999.0

    anterior = list(
        range(
            len(b) + 1
        )
    )

    for i, ca in enumerate(
        a,
        start=1,
    ):

        atual = [i]

        for j, cb in enumerate(
            b,
            start=1,
        ):

            if ca == cb:

                custo = 0.0

            elif (
                ca,
                cb,
            ) in CONFUNDIVEIS:

                custo = 0.35

            else:

                custo = 1.0

            atual.append(
                min(
                    anterior[j] + 1,
                    atual[j - 1] + 1,
                    anterior[j - 1] + custo,
                )
            )

        anterior = atual

    return anterior[-1]


# ============================================================
# DESCREVE DIFERENÇAS
# ============================================================

def descrever_diferencas(
    placa_esperada,
    placa_ocr,
):

    esperado = limpar_placa(
        placa_esperada
    )

    recebido = limpar_placa(
        placa_ocr
    )

    if not esperado or not recebido:
        return "Sem comparação"

    linhas = (
        len(esperado)
        + 1
    )

    colunas = (
        len(recebido)
        + 1
    )

    matriz = [
        [0.0 for _ in range(colunas)]
        for _ in range(linhas)
    ]

    operacoes = [
        [None for _ in range(colunas)]
        for _ in range(linhas)
    ]

    for i in range(linhas):

        matriz[i][0] = float(i)

    for j in range(colunas):

        matriz[0][j] = float(j)

    for i in range(
        1,
        linhas,
    ):

        for j in range(
            1,
            colunas,
        ):

            ca = esperado[
                i - 1
            ]

            cb = recebido[
                j - 1
            ]

            if ca == cb:

                custo = 0.0

            elif (
                ca,
                cb,
            ) in CONFUNDIVEIS:

                custo = 0.35

            else:

                custo = 1.0

            substituicao = (
                matriz[
                    i - 1
                ][
                    j - 1
                ]
                + custo
            )

            insercao = (
                matriz[i][j - 1]
                + 1.0
            )

            remocao = (
                matriz[i - 1][j]
                + 1.0
            )

            menor = min(
                substituicao,
                insercao,
                remocao,
            )

            matriz[i][j] = menor

            if menor == substituicao:

                if custo == 0:

                    operacoes[i][j] = (
                        "igual",
                        ca,
                        cb,
                    )

                else:

                    operacoes[i][j] = (
                        "substituicao",
                        ca,
                        cb,
                    )

            elif menor == insercao:

                operacoes[i][j] = (
                    "insercao",
                    "",
                    cb,
                )

            else:

                operacoes[i][j] = (
                    "remocao",
                    ca,
                    "",
                )

    # ========================================================
    # BACKTRACE
    # ========================================================

    i = len(esperado)
    j = len(recebido)

    diferencas = []

    while (
        i > 0
        or
        j > 0
    ):

        if (
            i > 0
            and
            j > 0
        ):

            tipo = operacoes[
                i
            ][
                j
            ][0]

        elif i > 0:

            tipo = "remocao"

        else:

            tipo = "insercao"

        if tipo == "igual":

            i -= 1
            j -= 1

        elif tipo == "substituicao":

            diferencas.append(
                (
                    i,
                    recebido[j - 1],
                    esperado[i - 1],
                )
            )

            i -= 1
            j -= 1

        elif tipo == "insercao":

            diferencas.append(
                (
                    j,
                    recebido[j - 1],
                    "—",
                )
            )

            j -= 1

        else:

            diferencas.append(
                (
                    i,
                    "—",
                    esperado[i - 1],
                )
            )

            i -= 1

    diferencas.reverse()

    if not diferencas:
        return "Nenhuma"

    resultado = []

    for (
        posicao,
        ocr,
        sistema,
    ) in diferencas:

        if sistema == "—":

            resultado.append(
                f"Posição {posicao}: "
                f"OCR={ocr} "
                f"(caractere extra)"
            )

        elif ocr == "—":

            resultado.append(
                f"Posição {posicao}: "
                f"OCR ausente "
                f"(esperado={sistema})"
            )

        else:

            resultado.append(
                f"Posição {posicao}: "
                f"OCR={ocr} "
                f"≠ Sistema={sistema}"
            )

    return "; ".join(
        resultado
    )


# ============================================================
# PROCURA A MELHOR SEQUÊNCIA DE 7 CARACTERES
# ============================================================

def encontrar_melhor_placa_no_texto(
    texto_ocr,
    placa_sistema,
):

    texto = limpar_placa(
        texto_ocr
    )

    placa = limpar_placa(
        placa_sistema
    )

    if not texto:
        return None

    if not placa:
        return None

    # ========================================================
    # TEXTO MENOR QUE 7
    # ========================================================

    if len(texto) < TAMANHO_PLACA:

        candidato = aplicar_mascara_placa(
            texto
        )

        if not candidato:
            return None

        distancia = distancia_placas(
            placa,
            candidato,
        )

        return {
            "candidato_bruto":
                texto,

            "candidato_limpo":
                candidato,

            "distancia":
                distancia,

            "inicio":
                0,

            "match_exato_bruto":
                texto == placa,

            "match_exato_limpo":
                candidato == placa,
        }

    # ========================================================
    # TESTA TODAS AS JANELAS DE 7
    # ========================================================

    melhor = None

    for inicio in range(
        0,
        len(texto)
        -
        TAMANHO_PLACA
        +
        1,
    ):

        candidato_bruto = texto[
            inicio:
            inicio
            +
            TAMANHO_PLACA
        ]

        candidato_limpo = (
            aplicar_mascara_placa(
                candidato_bruto
            )
        )

        if not candidato_limpo:
            continue

        distancia = distancia_placas(
            placa,
            candidato_limpo,
        )

        match_exato_bruto = (
            candidato_bruto
            ==
            placa
        )

        match_exato_limpo = (
            candidato_limpo
            ==
            placa
        )

        # ====================================================
        # PRIORIDADE
        #
        # 0 = bruto exatamente igual
        # 1 = máscara exatamente igual
        # 2 = melhor distância
        # ====================================================

        chave = (
            0
            if match_exato_bruto
            else
            1
            if match_exato_limpo
            else
            2,

            distancia,

            inicio,
        )

        if (
            melhor is None
            or
            chave
            <
            melhor["chave"]
        ):

            melhor = {
                "candidato_bruto":
                    candidato_bruto,

                "candidato_limpo":
                    candidato_limpo,

                "distancia":
                    distancia,

                "inicio":
                    inicio,

                "match_exato_bruto":
                    match_exato_bruto,

                "match_exato_limpo":
                    match_exato_limpo,

                "chave":
                    chave,
            }

    return melhor


# ============================================================
# EXTRAI A PLACA DO TEXTO OCR
# ============================================================

def extrair_placa_do_ocr(
    texto_ocr,
    placa_sistema,
):

    texto_original = str(
        texto_ocr or ""
    ).strip()

    texto_limpo = limpar_placa(
        texto_original
    )

    placa_sistema = limpar_placa(
        placa_sistema
    )

    if not texto_limpo:

        return {
            "ocr_original":
                texto_original,

            "ocr_limpo":
                "",

            "diferenca":
                999.0,

            "candidato":
                "",

            "inicio":
                None,
        }

    # ========================================================
    # EXATAMENTE 7 CARACTERES
    # ========================================================

    if len(texto_limpo) == TAMANHO_PLACA:

        candidato = aplicar_mascara_placa(
            texto_limpo
        )

        diferenca = distancia_placas(
            placa_sistema,
            candidato,
        )

        return {
            "ocr_original":
                texto_original,

            "ocr_limpo":
                candidato,

            "diferenca":
                diferenca,

            "candidato":
                texto_limpo,

            "inicio":
                0,
        }

    # ========================================================
    # MAIS DE 7 OU MENOS DE 7
    # ========================================================

    melhor = (
        encontrar_melhor_placa_no_texto(
            texto_limpo,
            placa_sistema,
        )
    )

    if melhor is None:

        return {
            "ocr_original":
                texto_original,

            "ocr_limpo":
                texto_limpo,

            "diferenca":
                distancia_placas(
                    placa_sistema,
                    texto_limpo,
                ),

            "candidato":
                texto_limpo,

            "inicio":
                None,
        }

    return {
        "ocr_original":
            texto_original,

        "ocr_limpo":
            melhor[
                "candidato_limpo"
            ],

        "diferenca":
            melhor[
                "distancia"
            ],

        "candidato":
            melhor[
                "candidato_bruto"
            ],

        "inicio":
            melhor[
                "inicio"
            ],
    }


# ============================================================
# BARRA DE PROGRESSO
# ============================================================

def mostrar_progresso(
    atual,
    total,
    aprovados,
    sem_deteccao,
    inicio_tempo,
):

    if total <= 0:

        porcentagem = 100.0

    else:

        porcentagem = (
            atual / total
        ) * 100

    largura = 45

    preenchido = int(
        largura
        * porcentagem
        / 100
    )

    barra = (
        "█" * preenchido
        +
        "░" * (
            largura
            -
            preenchido
        )
    )

    tempo_decorrido = (
        time.time()
        -
        inicio_tempo
    )

    if atual > 0:

        media = (
            tempo_decorrido
            /
            atual
        )

        restante = (
            total
            -
            atual
        )

        eta_segundos = (
            media
            *
            restante
        )

    else:

        eta_segundos = 0

    eta_segundos = int(
        eta_segundos
    )

    horas = (
        eta_segundos
        //
        3600
    )

    minutos = (
        eta_segundos
        %
        3600
    ) // 60

    segundos = (
        eta_segundos
        %
        60
    )

    eta = (
        f"{horas:02d}:"
        f"{minutos:02d}:"
        f"{segundos:02d}"
    )

    texto = (
        f"\rYOLO + OCR "
        f"[{barra}] "
        f"{porcentagem:6.2f}% "
        f"| {atual}/{total} "
        f"| Aprovados: {aprovados} "
        f"| Sem det.: {sem_deteccao} "
        f"| ETA: {eta}"
    )

    print(
        texto,
        end="",
        flush=True,
    )


# ============================================================
# COMUNICAÇÃO
# ============================================================

def enviar_json(
    processo,
    dados,
):

    if processo.poll() is not None:

        raise RuntimeError(
            f"Processo já encerrou. "
            f"Código: {processo.returncode}"
        )

    processo.stdin.write(
        json.dumps(
            dados,
            ensure_ascii=False,
        )
        + "\n"
    )

    processo.stdin.flush()


def ler_json(
    processo,
    nome,
):

    linha = (
        processo.stdout.readline()
    )

    if not linha:

        raise RuntimeError(
            f"O processo {nome} "
            f"encerrou antes de responder."
        )

    return json.loads(
        linha
    )


def encerrar_processo(
    processo,
    nome,
):

    if processo is None:
        return

    if processo.poll() is not None:
        return

    try:

        enviar_json(
            processo,
            {
                "command":
                    "exit"
            },
        )

        processo.wait(
            timeout=10
        )

    except Exception as e:

        print(
            f"\n[{nome}] "
            f"Erro ao encerrar: {e}"
        )

        try:

            processo.terminate()

            processo.wait(
                timeout=5
            )

        except Exception:

            try:
                processo.kill()

            except Exception:
                pass


# ============================================================
# INICIA YOLO
# ============================================================

def iniciar_yolo():

    print(
        "[INIT] Iniciando YOLO..."
    )

    return subprocess.Popen(
        [
            str(PYTHON_YOLO),
            str(DETECTOR_PY),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=None,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )


# ============================================================
# INICIA OCR
# ============================================================

def iniciar_ocr():

    print(
        "[INIT] Iniciando PaddleOCR..."
    )

    return subprocess.Popen(
        [
            str(PYTHON_OCR),
            str(OCR_PY),
        ],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=None,
        text=True,
        bufsize=1,
        cwd=str(BASE_DIR),
    )


# ============================================================
# LOCALIZA COLUNAS
# ============================================================

def localizar_colunas(sheet):

    cabecalho = {}

    for col in range(
        1,
        sheet.max_column + 1,
    ):

        valor = sheet.cell(
            row=1,
            column=col,
        ).value

        if valor is None:
            continue

        nome = normalizar_nome_coluna(
            valor
        )

        cabecalho[
            nome
        ] = col

    return cabecalho


def encontrar_coluna(
    cabecalho,
    nomes,
):

    for nome in nomes:

        nome_normalizado = (
            normalizar_nome_coluna(
                nome
            )
        )

        if (
            nome_normalizado
            in cabecalho
        ):

            return cabecalho[
                nome_normalizado
            ]

    return None


# ============================================================
# EXTRAI IMAGEM
# ============================================================

def extrair_bytes_imagem(img):

    try:

        return img._data()

    except Exception:

        return None


# ============================================================
# EXTENSÃO
# ============================================================

def descobrir_extensao(
    img,
    dados,
):

    formato = getattr(
        img,
        "format",
        None,
    )

    if formato:

        formato = str(
            formato
        ).lower()

        if formato in (
            "jpg",
            "jpeg",
        ):

            return ".jpg"

        if formato == "png":
            return ".png"

        if formato == "bmp":
            return ".bmp"

        if formato == "gif":
            return ".gif"

        if formato == "tiff":
            return ".tiff"

    if dados[:3] == b"\xff\xd8\xff":
        return ".jpg"

    if dados[:8] == (
        b"\x89PNG\r\n\x1a\n"
    ):
        return ".png"

    if dados[:4] == b"RIFF":
        return ".webp"

    return ".png"


# ============================================================
# LÊ EXCEL
# ============================================================

def carregar_excel(
    excel_path,
    temp_dir,
):

    print(
        f"\n[EXCEL] Abrindo:"
        f"\n{excel_path}"
    )

    wb = openpyxl.load_workbook(
        excel_path,
        data_only=True,
    )

    sheet = wb.active

    # ========================================================
    # COLUNAS
    # ========================================================

    cabecalho = localizar_colunas(
        sheet
    )

    coluna_id = encontrar_coluna(
        cabecalho,
        [
            "ID_PASSAGEM",
            "ID PASSAGEM",
            "ID",
            "PASSAGEM",
        ],
    )

    coluna_placa = encontrar_coluna(
        cabecalho,
        [
            "PLACA",
            "PLACA_VEICULO",
            "PLACA VEICULO",
        ],
    )

    if coluna_id is None:

        raise ValueError(
            "Não encontrei a coluna "
            "'ID_PASSAGEM' na planilha.\n"
            f"Colunas encontradas: "
            f"{list(cabecalho.keys())}"
        )

    if coluna_placa is None:

        raise ValueError(
            "Não encontrei a coluna "
            "'PLACA' na planilha.\n"
            f"Colunas encontradas: "
            f"{list(cabecalho.keys())}"
        )

    print(
        f"[EXCEL] Coluna ID: "
        f"{coluna_id}"
    )

    print(
        f"[EXCEL] Coluna PLACA: "
        f"{coluna_placa}"
    )

    # ========================================================
    # IMAGENS POR LINHA
    # ========================================================

    imagens_por_linha = {}

    for img in sheet._images:

        try:

            linha = (
                img.anchor._from.row
                + 1
            )

        except Exception:

            continue

        imagens_por_linha.setdefault(
            linha,
            [],
        ).append(
            img
        )

    print(
        f"[EXCEL] Total de imagens: "
        f"{len(sheet._images)}"
    )

    # ========================================================
    # PASSAGENS
    # ========================================================

    passagens = []

    for linha in range(
        2,
        sheet.max_row + 1,
    ):

        id_passagem = sheet.cell(
            row=linha,
            column=coluna_id,
        ).value

        placa = sheet.cell(
            row=linha,
            column=coluna_placa,
        ).value

        if (
            id_passagem is None
            or
            str(
                id_passagem
            ).strip()
            == ""
        ):

            continue

        imagens = []

        imagens_linha = (
            imagens_por_linha.get(
                linha,
                [],
            )
        )

        for img_idx, img in enumerate(
            imagens_linha
        ):

            dados = (
                extrair_bytes_imagem(
                    img
                )
            )

            if not dados:
                continue

            ext = descobrir_extensao(
                img,
                dados,
            )

            arquivo = (
                temp_dir
                /
                (
                    f"linha_{linha}"
                    f"_img_{img_idx}"
                    f"{ext}"
                )
            )

            arquivo.write_bytes(
                dados
            )

            imagens.append(
                {
                    "task_id":
                        (
                            f"l{linha}"
                            f"_i{img_idx}"
                        ),

                    "image_path":
                        str(
                            arquivo
                        ),

                    "indice":
                        img_idx,
                }
            )

        passagens.append(
            {
                "linha":
                    linha,

                "id_passagem":
                    id_passagem,

                "placa":
                    placa,

                "imagens":
                    imagens,
            }
        )

    print(
        f"[EXCEL] Passagens reais: "
        f"{len(passagens)}"
    )

    wb.close()

    return (
        passagens,
        cabecalho,
    )


# ============================================================
# YOLO
# ============================================================

def executar_yolo_batch(
    processo_yolo,
    imagens,
):

    if not imagens:
        return {}

    items = []

    for imagem in imagens:

        items.append(
            {
                "task_id":
                    imagem[
                        "task_id"
                    ],

                "image_path":
                    imagem[
                        "image_path"
                    ],
            }
        )

    enviar_json(
        processo_yolo,
        {
            "command":
                "detect_batch",

            "images":
                items,
        },
    )

    resposta = ler_json(
        processo_yolo,
        "YOLO",
    )

    resultados = {}

    for resultado in resposta.get(
        "results",
        [],
    ):

        task_id = resultado.get(
            "task_id"
        )

        if task_id:

            resultados[
                task_id
            ] = resultado

    return resultados


# ============================================================
# OCR DE UMA IMAGEM
# ============================================================

def executar_ocr_imagem(
    processo_ocr,
    task_id,
    crop_b64,
):

    if not crop_b64:
        return {}

    enviar_json(
        processo_ocr,
        {
            "command":
                "ocr_batch",

            "images": [
                {
                    "task_id":
                        task_id,

                    "crop_b64":
                        crop_b64,
                }
            ],
        },
    )

    resposta = ler_json(
        processo_ocr,
        "OCR",
    )

    resultados = resposta.get(
        "results",
        [],
    )

    if not resultados:
        return {}

    return resultados[0]


# ============================================================
# OCR POR PASSAGEM
# ============================================================

def processar_ocr_por_passagem(
    passagem,
    todos_yolo,
    processo_ocr,
):

    placa = limpar_placa(
        passagem[
            "placa"
        ]
    )

    melhor_task = None
    melhor_resultado = None

    melhor_distancia = 999.0

    melhor_placa_limpa = ""

    qtd_yolo = 0
    qtd_ocr = 0

    # ========================================================
    # PERCORRE AS IMAGENS DA PASSAGEM
    # ========================================================

    for imagem in passagem[
        "imagens"
    ]:

        task_id = imagem[
            "task_id"
        ]

        yolo_result = (
            todos_yolo.get(
                task_id,
                {},
            )
        )

        # ====================================================
        # YOLO NÃO DETECTOU
        # ====================================================

        if not yolo_result.get(
            "detected",
            False,
        ):

            continue

        qtd_yolo += 1

        crop_b64 = (
            yolo_result.get(
                "crop_b64"
            )
        )

        if not crop_b64:

            detections = (
                yolo_result.get(
                    "detections",
                    [],
                )
            )

            if detections:

                crop_b64 = detections[
                    0
                ].get(
                    "crop_b64"
                )

        if not crop_b64:

            continue

        # ====================================================
        # OCR
        # ====================================================

        resultado_ocr = (
            executar_ocr_imagem(
                processo_ocr,
                task_id,
                crop_b64,
            )
        )

        qtd_ocr += 1

        texto_ocr = str(
            resultado_ocr.get(
                "ocr_best",
                ""
            )
            or
            ""
        ).strip()

        # ====================================================
        # SEM TEXTO
        # ====================================================

        if not limpar_placa(
            texto_ocr
        ):

            continue

        # ====================================================
        # NOVA LIMPEZA INTELIGENTE
        # ====================================================

        extracao = (
            extrair_placa_do_ocr(
                texto_ocr,
                placa,
            )
        )

        placa_ocr_limpa = (
            extracao[
                "ocr_limpo"
            ]
        )

        distancia = (
            extracao[
                "diferenca"
            ]
        )

        # ====================================================
        # DEBUG
        # ====================================================

        print(
            "\n[OCR]"
        )

        print(
            f"ID: "
            f"{passagem['id_passagem']}"
        )

        print(
            f"Placa sistema: "
            f"{placa}"
        )

        print(
            f"OCR Paddle: "
            f"{texto_ocr}"
        )

        print(
            f"OCR limpo: "
            f"{placa_ocr_limpa}"
        )

        print(
            f"Distância: "
            f"{distancia}"
        )

        if (
            texto_ocr
            !=
            placa_ocr_limpa
        ):

            print(
                f"[LIMPEZA] "
                f"{texto_ocr} "
                f"→ "
                f"{placa_ocr_limpa}"
            )

        # ====================================================
        # GUARDA MELHOR RESULTADO
        #
        # Aqui já guardamos a placa LIMPA.
        # ====================================================

        if (
            melhor_resultado is None
            or
            distancia
            <
            melhor_distancia
        ):

            melhor_distancia = (
                distancia
            )

            melhor_task = (
                task_id
            )

            melhor_resultado = (
                resultado_ocr
            )

            melhor_placa_limpa = (
                placa_ocr_limpa
            )

        # ====================================================
        # MATCH
        # ====================================================

        if (
            distancia
            <=
            MAX_DIFERENCAS_APROVACAO
        ):

            print(
                f"[MATCH] "
                f"{passagem['id_passagem']} "
                f"→ "
                f"{placa_ocr_limpa} "
                f"→ dif="
                f"{distancia}"
            )

            return {
                "task_id":
                    task_id,

                "resultado":
                    resultado_ocr,

                # IMPORTANTE:
                # agora é a placa limpa,
                # não o texto bruto do Paddle.
                "ocr":
                    placa_ocr_limpa,

                "ocr_bruto":
                    texto_ocr,

                "distancia":
                    round(
                        distancia,
                        2,
                    ),

                "status":
                    "Aprovado",

                "qtd_yolo":
                    qtd_yolo,

                "qtd_ocr":
                    qtd_ocr,

                "candidato":
                    extracao.get(
                        "candidato",
                        "",
                    ),
            }

    # ========================================================
    # TERMINOU SEM MATCH
    # ========================================================

    if melhor_resultado is not None:

        return {
            "task_id":
                melhor_task,

            "resultado":
                melhor_resultado,

            "ocr":
                melhor_placa_limpa,

            "ocr_bruto":
                melhor_resultado.get(
                    "ocr_best",
                    "",
                ),

            "distancia":
                round(
                    melhor_distancia,
                    2,
                ),

            "status":
                "Revisar",

            "qtd_yolo":
                qtd_yolo,

            "qtd_ocr":
                qtd_ocr,

            "candidato":
                "",
        }

    # ========================================================
    # NENHUM OCR
    # ========================================================

    return {
        "task_id":
            None,

        "resultado":
            {},

        "ocr":
            "",

        "ocr_bruto":
            "",

        "distancia":
            999.0,

        "status":
            "Sem detecção",

        "qtd_yolo":
            qtd_yolo,

        "qtd_ocr":
            qtd_ocr,

        "candidato":
            "",
    }


# ============================================================
# PROCESSA UMA PLANILHA
# ============================================================

def processar_excel(
    excel_path,
    processo_yolo,
    processo_ocr,
):

    temp_dir = Path(
        tempfile.mkdtemp(
            prefix=
            "validador_excel_"
        )
    )

    try:

        passagens, cabecalho = (
            carregar_excel(
                excel_path,
                temp_dir,
            )
        )

        # ====================================================
        # TODAS AS IMAGENS
        # ====================================================

        todas_imagens = []

        mapa_passagem = {}

        for passagem in passagens:

            linha = passagem[
                "linha"
            ]

            for imagem in (
                passagem[
                    "imagens"
                ]
            ):

                todas_imagens.append(
                    imagem
                )

                mapa_passagem[
                    imagem[
                        "task_id"
                    ]
                ] = linha

        total_imagens = len(
            todas_imagens
        )

        # ====================================================
        # YOLO
        # ====================================================

        todos_yolo = {}

        total_lotes = (
            (
                total_imagens
                +
                BATCH_SIZE
                -
                1
            )
            //
            BATCH_SIZE
        )

        for lote_idx in range(
            total_lotes
        ):

            inicio = (
                lote_idx
                *
                BATCH_SIZE
            )

            fim = min(
                inicio
                +
                BATCH_SIZE,
                total_imagens,
            )

            lote = todas_imagens[
                inicio:fim
            ]

            yolo_results = (
                executar_yolo_batch(
                    processo_yolo,
                    lote,
                )
            )

            todos_yolo.update(
                yolo_results
            )

        # ====================================================
        # OCR POR PASSAGEM
        # ====================================================

        resultados_ocr_passagens = {}

        inicio_tempo = time.time()

        total_passagens = len(
            passagens
        )

        aprovados = 0
        sem_deteccao = 0

        for indice, passagem in enumerate(
            passagens,
            start=1,
        ):

            linha = passagem[
                "linha"
            ]

            resultado = (
                processar_ocr_por_passagem(
                    passagem,
                    todos_yolo,
                    processo_ocr,
                )
            )

            resultados_ocr_passagens[
                linha
            ] = resultado

            # ------------------------------------------------
            # Estatísticas
            # ------------------------------------------------

            if (
                resultado["status"]
                ==
                "Aprovado"
            ):

                aprovados += 1

            elif (
                resultado["status"]
                ==
                "Sem detecção"
            ):

                sem_deteccao += 1

            mostrar_progresso(
                indice,
                total_passagens,
                aprovados,
                sem_deteccao,
                inicio_tempo,
            )

        print()

        # ====================================================
        # RESULTADO FINAL
        # ====================================================

        resultados_finais = []

        for passagem in passagens:

            linha = passagem[
                "linha"
            ]

            placa = passagem[
                "placa"
            ]

            resultado_passagem = (
                resultados_ocr_passagens.get(
                    linha,
                    {
                        "task_id":
                            None,

                        "resultado":
                            {},

                        "ocr":
                            "",

                        "ocr_bruto":
                            "",

                        "distancia":
                            999.0,

                        "status":
                            "Sem detecção",

                        "qtd_yolo":
                            0,

                        "qtd_ocr":
                            0,
                    },
                )
            )

            resultado_ocr = (
                resultado_passagem[
                    "resultado"
                ]
            )

            task_id = (
                resultado_passagem[
                    "task_id"
                ]
            )

            yolo_result = (
                todos_yolo.get(
                    task_id,
                    {},
                )
            )

            qtd_imagens = len(
                passagem[
                    "imagens"
                ]
            )

            qtd_yolo = (
                resultado_passagem[
                    "qtd_yolo"
                ]
            )

            qtd_ocr = (
                resultado_passagem[
                    "qtd_ocr"
                ]
            )

            resultados_finais.append(
                {
                    "Linha Excel":
                        linha,

                    "ID_PASSAGEM":
                        passagem[
                            "id_passagem"
                        ],

                    "DATA_HORA":
                        None,

                    "PRACA_PEDAGIO":
                        None,

                    "SENTIDO":
                        None,

                    "PISTA":
                        None,

                    "MODALIDADE_PAGAMENTO":
                        None,

                    "PLACA":
                        placa,

                    "CATEGORIA":
                        None,

                    "DISPOSITIVO TAG":
                        None,

                    "CATEGORIA_COBRADA":
                        None,

                    "EIXOS_SUSPENSOS":
                        None,

                    "Qtd Imagens":
                        qtd_imagens,

                    "Qtd YOLO":
                        qtd_yolo,

                    "Qtd OCR":
                        qtd_ocr,

                    # =================================================
                    # PLACA OCR = LIMPA
                    # =================================================

                    "Placa OCR":
                        resultado_passagem[
                            "ocr"
                        ],

                    # =================================================
                    # OCR BRUTO
                    # =================================================

                    "OCR Bruto":
                        resultado_passagem.get(
                            "ocr_bruto",
                            "",
                        ),

                    "Distancia OCR":
                        resultado_passagem[
                            "distancia"
                        ],

                    "Status OCR":
                        resultado_passagem[
                            "status"
                        ],

                    "Imagem Usada":
                        task_id or "",

                    "YOLO Confidence":
                        yolo_result.get(
                            "confidence",
                            0,
                        ),

                    # =========================================
                    # CUBIC
                    # =========================================

                    "OCR Cubic":
                        resultado_ocr.get(
                            "ocr_cubic",
                            "",
                        ),

                    "Conf Cubic":
                        resultado_ocr.get(
                            "confidence_cubic",
                            0,
                        ),

                    "OCR Cubic Auxiliar":
                        resultado_ocr.get(
                            "ocr_cubic_auxiliar",
                            "",
                        ),

                    # =========================================
                    # CLAHE
                    # =========================================

                    "OCR CLAHE":
                        resultado_ocr.get(
                            "ocr_clahe",
                            "",
                        ),

                    "Conf CLAHE":
                        resultado_ocr.get(
                            "confidence_clahe",
                            0,
                        ),

                    "OCR CLAHE Auxiliar":
                        resultado_ocr.get(
                            "ocr_clahe_auxiliar",
                            "",
                        ),

                    # =========================================
                    # NORMAL
                    # =========================================

                    "OCR Best Original":
                        resultado_ocr.get(
                            "ocr_best_original",
                            "",
                        ),

                    "OCR Best Normal":
                        resultado_ocr.get(
                            "ocr_best_normal",
                            "",
                        ),

                    "Conf Best Normal":
                        resultado_ocr.get(
                            "confidence_best_normal",
                            0,
                        ),

                    # =========================================
                    # SR
                    # =========================================

                    "SR Usada":
                        resultado_ocr.get(
                            "sr_used",
                            False,
                        ),

                    "SR OCR":
                        resultado_ocr.get(
                            "sr_ocr",
                            "",
                        ),

                    "SR Confidence":
                        resultado_ocr.get(
                            "sr_confidence",
                            0,
                        ),

                    "SR Variant":
                        resultado_ocr.get(
                            "sr_variant",
                            "",
                        ),

                    "Resultado Graças SR":
                        resultado_ocr.get(
                            "resultado_gracas_sr",
                            False,
                        ),

                    # =========================================
                    # FINAL PADDLE
                    # =========================================

                    "OCR Best Final":
                        resultado_ocr.get(
                            "ocr_best",
                            "",
                        ),

                    "Conf Best Final":
                        resultado_ocr.get(
                            "confidence_best",
                            0,
                        ),

                    "Best Variant":
                        resultado_ocr.get(
                            "best_variant",
                            "",
                        ),

                    "Auxiliar Alterou":
                        resultado_ocr.get(
                            "auxiliar_alterou",
                            False,
                        ),
                }
            )

        return resultados_finais

    finally:

        shutil.rmtree(
            temp_dir,
            ignore_errors=True,
        )


# ============================================================
# PROCESSA PASTA
# ============================================================

def processar_pasta(
    pasta,
):

    pasta = Path(
        pasta
    ).resolve()

    arquivos = sorted(
        [
            p
            for p in pasta.rglob(
                "*.xlsx"
            )
            if (
                not p.name.startswith(
                    "~$"
                )
                and
                not p.stem.endswith(
                    "_validado"
                )
            )
        ]
    )

    if not arquivos:

        print(
            "[ERRO] Nenhum Excel encontrado."
        )

        return

    print(
        "\n"
        + "=" * 75
    )

    print(
        "VALIDADOR EXCEL"
    )

    print(
        "=" * 75
    )

    print(
        f"Pasta: {pasta}"
    )

    print(
        f"Arquivos: {len(arquivos)}"
    )

    print(
        "=" * 75
    )

    # ========================================================
    # DEPENDÊNCIAS
    # ========================================================

    for arquivo in [
        PYTHON_YOLO,
        PYTHON_OCR,
        DETECTOR_PY,
        OCR_PY,
    ]:

        if not arquivo.exists():

            raise FileNotFoundError(
                f"Arquivo não encontrado:\n"
                f"{arquivo}"
            )

    # ========================================================
    # PROCESSOS
    # ========================================================

    processo_yolo = iniciar_yolo()
    processo_ocr = iniciar_ocr()

    total_passagens = 0
    total_aprovados = 0
    total_revisar = 0
    total_sem_deteccao = 0

    total_imagens = 0
    total_yolo = 0
    total_ocr = 0

    arquivos_ok = 0
    arquivos_erro = 0

    tempo_inicio_total = time.time()

    try:

        print(
            "\n[INIT] "
            "Aguardando modelos..."
        )

        time.sleep(
            15
        )

        # ====================================================
        # CADA EXCEL
        # ====================================================

        for arquivo_idx, excel_path in enumerate(
            arquivos,
            start=1,
        ):

            print(
                "\n"
                + "#" * 75
            )

            print(
                f"ARQUIVO "
                f"{arquivo_idx}/"
                f"{len(arquivos)}"
            )

            print(
                excel_path
            )

            print(
                "#" * 75
            )

            try:

                resultados = (
                    processar_excel(
                        excel_path,
                        processo_yolo,
                        processo_ocr,
                    )
                )

                df = pd.DataFrame(
                    resultados
                )

                saida = (
                    excel_path.parent
                    /
                    (
                        excel_path.stem
                        +
                        "_validado.xlsx"
                    )
                )

                # =================================================
                # DADOS ORIGINAIS
                # =================================================

                wb_original = (
                    openpyxl.load_workbook(
                        excel_path,
                        data_only=True,
                    )
                )

                sheet_original = (
                    wb_original.active
                )

                cabecalho_original = (
                    localizar_colunas(
                        sheet_original
                    )
                )

                mapa_colunas = {
                    "DATA_HORA": [
                        "DATA_HORA",
                        "DATA HORA",
                    ],

                    "PRACA_PEDAGIO": [
                        "PRACA_PEDAGIO",
                        "PRACA PEDAGIO",
                    ],

                    "SENTIDO": [
                        "SENTIDO",
                    ],

                    "PISTA": [
                        "PISTA",
                    ],

                    "MODALIDADE_PAGAMENTO": [
                        "MODALIDADE_PAGAMENTO",
                        "MODALIDADE PAGAMENTO",
                    ],

                    "CATEGORIA": [
                        "CATEGORIA",
                    ],

                    "DISPOSITIVO TAG": [
                        "DISPOSITIVO TAG",
                        "DISPOSITIVO_TAG",
                    ],

                    "CATEGORIA_COBRADA": [
                        "CATEGORIA_COBRADA",
                        "CATEGORIA COBRADA",
                    ],

                    "EIXOS_SUSPENSOS": [
                        "EIXOS_SUSPENSOS",
                        "EIXOS SUSPENSOS",
                    ],
                }

                for campo, nomes in (
                    mapa_colunas.items()
                ):

                    coluna = encontrar_coluna(
                        cabecalho_original,
                        nomes,
                    )

                    if coluna is None:
                        continue

                    df[campo] = [
                        sheet_original.cell(
                            row=int(
                                linha
                            ),
                            column=coluna,
                        ).value
                        for linha in df[
                            "Linha Excel"
                        ]
                    ]

                wb_original.close()

                # =================================================
                # REORDENA COLUNAS
                # =================================================

                colunas_preferenciais = [
                    "Linha Excel",
                    "ID_PASSAGEM",
                    "DATA_HORA",
                    "PRACA_PEDAGIO",
                    "SENTIDO",
                    "PISTA",
                    "MODALIDADE_PAGAMENTO",
                    "PLACA",
                    "CATEGORIA",
                    "DISPOSITIVO TAG",
                    "CATEGORIA_COBRADA",
                    "EIXOS_SUSPENSOS",
                    "Qtd Imagens",
                    "Qtd YOLO",
                    "Qtd OCR",
                    "Placa OCR",
                    "OCR Bruto",
                    "Distancia OCR",
                    "Status OCR",
                    "Imagem Usada",
                    "YOLO Confidence",
                    "OCR Cubic",
                    "Conf Cubic",
                    "OCR Cubic Auxiliar",
                    "OCR CLAHE",
                    "Conf CLAHE",
                    "OCR CLAHE Auxiliar",
                    "OCR Best Original",
                    "OCR Best Normal",
                    "Conf Best Normal",
                    "SR Usada",
                    "SR OCR",
                    "SR Confidence",
                    "SR Variant",
                    "Resultado Graças SR",
                    "OCR Best Final",
                    "Conf Best Final",
                    "Best Variant",
                    "Auxiliar Alterou",
                ]

                colunas_existentes = [
                    c
                    for c in (
                        colunas_preferenciais
                    )
                    if c in df.columns
                ]

                outras_colunas = [
                    c
                    for c in df.columns
                    if c not in colunas_existentes
                ]

                df = df[
                    colunas_existentes
                    +
                    outras_colunas
                ]

                # =================================================
                # SALVA
                # =================================================

                df.to_excel(
                    saida,
                    index=False,
                )

                qtd = len(
                    df
                )

                aprovados = int(
                    (
                        df[
                            "Status OCR"
                        ]
                        ==
                        "Aprovado"
                    ).sum()
                )

                sem_deteccao = int(
                    (
                        df[
                            "Status OCR"
                        ]
                        ==
                        "Sem detecção"
                    ).sum()
                )

                revisar = int(
                    (
                        df[
                            "Status OCR"
                        ]
                        ==
                        "Revisar"
                    ).sum()
                )

                imagens = int(
                    df[
                        "Qtd Imagens"
                    ]
                    .fillna(0)
                    .sum()
                )

                yolo = int(
                    df[
                        "Qtd YOLO"
                    ]
                    .fillna(0)
                    .sum()
                )

                ocr = int(
                    df[
                        "Qtd OCR"
                    ]
                    .fillna(0)
                    .sum()
                )

                total_passagens += qtd
                total_aprovados += aprovados
                total_sem_deteccao += (
                    sem_deteccao
                )
                total_revisar += (
                    revisar
                )

                total_imagens += imagens
                total_yolo += yolo
                total_ocr += ocr

                arquivos_ok += 1

                print(
                    "\n[OK]"
                )

                print(
                    f"Passagens: {qtd}"
                )

                print(
                    f"Aprovados: "
                    f"{aprovados}"
                )

                print(
                    f"Sem detecção: "
                    f"{sem_deteccao}"
                )

                print(
                    f"Revisar: "
                    f"{revisar}"
                )

                print(
                    f"Saída: {saida}"
                )

            except Exception as e:

                arquivos_erro += 1

                print(
                    "\n[ERRO] "
                    f"{type(e).__name__}: "
                    f"{e}"
                )

    finally:

        encerrar_processo(
            processo_yolo,
            "YOLO",
        )

        encerrar_processo(
            processo_ocr,
            "OCR",
        )

    # ========================================================
    # RESUMO
    # ========================================================

    tempo_total = (
        time.time()
        -
        tempo_inicio_total
    )

    horas = int(
        tempo_total
        //
        3600
    )

    minutos = int(
        (
            tempo_total
            %
            3600
        )
        //
        60
    )

    segundos = int(
        tempo_total
        %
        60
    )

    if total_passagens > 0:

        taxa_aprovacao = (
            total_aprovados
            /
            total_passagens
        ) * 100

    else:

        taxa_aprovacao = 0.0

    print(
        "\n"
        + "=" * 75
    )

    print(
        "RESUMO FINAL"
    )

    print(
        "=" * 75
    )

    print(
        f"Arquivos...............: "
        f"{len(arquivos)}"
    )

    print(
        f"Processados............: "
        f"{arquivos_ok}"
    )

    print(
        f"Com erro...............: "
        f"{arquivos_erro}"
    )

    print(
        f"Passagens..............: "
        f"{total_passagens}"
    )

    print(
        f"Aprovados..............: "
        f"{total_aprovados}"
    )

    print(
        f"Sem detecção...........: "
        f"{total_sem_deteccao}"
    )

    print(
        f"Revisar................: "
        f"{total_revisar}"
    )

    print(
        f"Taxa de aprovação......: "
        f"{taxa_aprovacao:.2f}%"
    )

    print(
        f"Imagens analisadas.....: "
        f"{total_imagens}"
    )

    print(
        f"YOLO detectou..........: "
        f"{total_yolo}"
    )

    print(
        f"OCR executados.........: "
        f"{total_ocr}"
    )

    print(
        f"Tempo total............: "
        f"{horas:02d}:"
        f"{minutos:02d}:"
        f"{segundos:02d}"
    )

    print(
        "=" * 75
    )


# ============================================================
# MAIN
# ============================================================

def main():

    if len(sys.argv) < 2:

        print(
            'Uso: python validador_excel.py '
            '"C:\\caminho\\da\\pasta"'
        )

        sys.exit(1)

    pasta = Path(
        sys.argv[1]
    ).resolve()

    if not pasta.exists():

        print(
            f"[ERRO] Pasta não encontrada:\n"
            f"{pasta}"
        )

        sys.exit(1)

    if not pasta.is_dir():

        print(
            f"[ERRO] O caminho não é uma pasta:\n"
            f"{pasta}"
        )

        sys.exit(1)

    processar_pasta(
        pasta
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    main()